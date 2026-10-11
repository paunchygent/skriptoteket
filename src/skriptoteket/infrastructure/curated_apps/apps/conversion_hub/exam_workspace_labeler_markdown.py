"""Faithful DOCX markdown for the item-type labeler (D14); the converter chunker stays unchanged.

Each chunk keeps the converter's anchor, kind and text. Only ``marked_text`` changes: it
renders the block as the teacher sees it in Word, with heading level, indentation, Word
bullets, bold as ``**...**`` and italic as ``<i>...</i>`` (the converter's existing cue
for inline facit). Code adds no text and removes none.
"""

from __future__ import annotations

import io
from collections.abc import Iterator
from dataclasses import replace

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.styles.style import CharacterStyle
from docx.table import Table, _Cell
from docx.text.hyperlink import Hyperlink
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from skriptoteket.domain.curated_apps.exam_workspace.docx_chunks import DocxChunk
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_docx_chunker import (  # noqa: E501
    PythonDocxExamChunker,
)

# Content markdown cannot carry faithfully; such a block needs an image of that part.
_IMAGE_ONLY = tuple(
    qn(tag)
    for tag in ("w:drawing", "w:pict", "w:object", "w:txbxContent", "w:gridSpan", "w:vMerge")
)
_TWIPS_PER_SPACE = 180  # 9 pt per rendered space


class LabelerSourceImageRequired(ValueError):
    """The source block needs a rendered image rather than a text approximation."""

    def __init__(self, anchor: str) -> None:
        self.anchor = anchor
        super().__init__(f"Source image required for {anchor}")


def _style_chain(style: CharacterStyle | None) -> Iterator[CharacterStyle]:
    """Yield a style and its bases; Word ignores a basedOn cycle, so stop at a repeat."""
    seen: set[str | None] = set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        yield style
        style = style.base_style


def _font_flag(run: Run, paragraph: Paragraph, name: str) -> bool:
    """Resolve bold/italic as Word displays it: run, then character style, then paragraph style."""
    direct = getattr(run.font, name)
    if direct is not None:
        return bool(direct)
    for style in (*_style_chain(run.style), *_style_chain(paragraph.style)):
        flag = getattr(style.font, name)
        if flag is not None:
            return bool(flag)
    return False


def _emphasis(paragraph: Paragraph) -> str:
    stretches: list[tuple[tuple[bool, bool], str]] = []
    for content in paragraph.iter_inner_content():
        for run in content.runs if isinstance(content, Hyperlink) else (content,):
            flags = (_font_flag(run, paragraph, "bold"), _font_flag(run, paragraph, "italic"))
            if stretches and stretches[-1][0] == flags:
                stretches[-1] = (flags, stretches[-1][1] + run.text)
            else:
                stretches.append((flags, run.text))
    rendered: list[str] = []
    for (bold, italic), text in stretches:
        core = text.strip()
        if core and (bold or italic):
            start = text.index(core)
            marked = f"<i>{core}</i>" if italic else core
            marked = f"**{marked}**" if bold else marked
            text = text[:start] + marked + text[start + len(core) :]
        rendered.append(text)
    return "".join(rendered)


def _child(element: BaseOxmlElement | None, tag: str) -> BaseOxmlElement | None:
    return None if element is None else element.find(qn(tag))


def _val(element: BaseOxmlElement | None, attribute: str = "w:val") -> str | None:
    return None if element is None else element.get(qn(attribute))


def _paragraph_property(paragraph: Paragraph, tag: str) -> BaseOxmlElement | None:
    """Find a paragraph property directly or through the paragraph style chain."""
    found = _child(_child(paragraph._p, "w:pPr"), tag)
    for style in _style_chain(paragraph.style):
        if found is not None:
            break
        found = _child(_child(style.element, "w:pPr"), tag)
    return found


def _indent_spaces(indentation: BaseOxmlElement | None) -> int | None:
    left = _val(indentation, "w:left") or _val(indentation, "w:start")
    if left is None:
        return None
    hanging = int(_val(indentation, "w:hanging") or 0)
    first_line = int(_val(indentation, "w:firstLine") or 0)
    return max(0, round((int(left) - hanging + first_line) / _TWIPS_PER_SPACE))


def _list_level(paragraph: Paragraph, document: DocxDocument) -> BaseOxmlElement | None:
    """Return the numbering level definition that Word renders for this paragraph."""
    properties = _paragraph_property(paragraph, "w:numPr")
    num_id = _val(_child(properties, "w:numId"))
    if num_id is None or num_id == "0":
        return None
    level = _val(_child(properties, "w:ilvl")) or "0"
    numbering: BaseOxmlElement = document.part.numbering_part.element
    instance = next(n for n in numbering.findall(qn("w:num")) if _val(n, "w:numId") == num_id)
    abstract_id = _val(_child(instance, "w:abstractNumId"))
    abstract: BaseOxmlElement = next(
        a
        for a in numbering.findall(qn("w:abstractNum"))
        if _val(a, "w:abstractNumId") == abstract_id
    )
    found: BaseOxmlElement = next(
        lvl for lvl in abstract.findall(qn("w:lvl")) if _val(lvl, "w:ilvl") == level
    )
    return found


def _paragraph(paragraph: Paragraph, document: DocxDocument) -> str:
    marker = ""
    level = _list_level(paragraph, document)
    if level is not None:
        number_format = _val(_child(level, "w:numFmt"))
        if number_format == "bullet":
            marker = "- "
        elif number_format != "none":
            # A Word-generated number is not in the text; render that block as an image.
            raise LabelerSourceImageRequired("")
    direct = _child(_child(paragraph._p, "w:pPr"), "w:ind")
    spaces = _indent_spaces(direct)
    if spaces is None and level is not None:
        spaces = _indent_spaces(_child(_child(level, "w:pPr"), "w:ind"))
    if spaces is None:
        spaces = _indent_spaces(_paragraph_property(paragraph, "w:ind"))
    style_name = paragraph.style.name if paragraph.style is not None else ""
    heading = ""
    if style_name == "Title":
        heading = "# "
    elif style_name.startswith("Heading ") and style_name[8:].isdigit():
        heading = "#" * min(6, int(style_name[8:])) + " "
    return " " * (spaces or 0) + marker + heading + _emphasis(paragraph)


def _table(table: Table, document: DocxDocument) -> str:
    rows: list[str] = []
    for row in table.rows:
        cells: list[str] = []
        for element in row._tr.tc_lst:
            cell = _Cell(element, table)
            if cell.tables:
                raise LabelerSourceImageRequired("")
            cells.append("<br>".join(_paragraph(p, document) for p in cell.paragraphs))
        rows.append("| " + " | ".join(cells) + " |")
    # The GFM separator row carries table structure, not source text.
    width = len(table.rows[0]._tr.tc_lst)
    rows.insert(1, "| " + " | ".join("---" for _ in range(width)) + " |")
    return "\n".join(rows)


class PythonDocxLabelerMarkdown:
    """Project converter chunks to faithful labeler markdown from the same DOCX bytes."""

    def chunk(self, *, content: bytes) -> tuple[DocxChunk, ...]:
        source = PythonDocxExamChunker().chunk(content=content)
        by_anchor = {chunk.anchor: chunk for chunk in source}
        document = Document(io.BytesIO(content))
        projected: list[DocxChunk] = []
        for ordinal, element in enumerate(document.element.body, start=1):
            anchor = f"body:{ordinal:04d}"
            chunk = by_anchor.get(anchor)
            if chunk is None:
                continue
            if any(next(element.iter(tag), None) is not None for tag in _IMAGE_ONLY):
                raise LabelerSourceImageRequired(anchor)
            try:
                if isinstance(element, CT_P):
                    markdown = _paragraph(Paragraph(element, document), document)
                else:
                    assert isinstance(element, CT_Tbl)
                    markdown = _table(Table(element, document), document)
            except LabelerSourceImageRequired as error:
                raise LabelerSourceImageRequired(anchor) from error
            projected.append(replace(chunk, marked_text=markdown))
        return tuple(projected)
