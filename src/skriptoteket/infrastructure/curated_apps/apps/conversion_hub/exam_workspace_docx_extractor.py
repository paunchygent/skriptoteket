"""python-docx adapter for the deterministic exam workspace extractor.

Maps a DOCX upload to the neutral block stream consumed by the pure domain
rules in ``docx_extraction``. OCR/scanned input has no path here (deferred
by ST-SKRIPT-39-04 S5); unsupported constructs are surfaced as notes, not
dropped silently.
"""

import hashlib
import io
import zipfile
from uuid import UUID

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from skriptoteket.domain.curated_apps.exam_workspace.docx_extraction import (
    DocxBlockKind,
    DocxExtractionResult,
    DocxSourceBlock,
    extract_native_exam_from_docx_blocks,
)
from skriptoteket.domain.errors import DomainError, ErrorCode

_STYLE_KIND_PREFIXES: tuple[tuple[str, DocxBlockKind], ...] = (
    ("Heading 1", DocxBlockKind.TITLE),
    ("Title", DocxBlockKind.TITLE),
    ("Heading 2", DocxBlockKind.SECTION),
    ("Heading 3", DocxBlockKind.HEADING),
    ("Heading", DocxBlockKind.HEADING),
    ("List", DocxBlockKind.LIST),
    ("Compact", DocxBlockKind.LIST),
)


def _block_kind(style_name: str) -> DocxBlockKind:
    for prefix, kind in _STYLE_KIND_PREFIXES:
        if style_name.startswith(prefix):
            return kind
    return DocxBlockKind.BODY


class PythonDocxExamExtractor:
    """Deterministic DOCX extraction via python-docx."""

    def extract(self, *, document_id: UUID, filename: str, content: bytes) -> DocxExtractionResult:
        try:
            parsed = Document(io.BytesIO(content))
        except (
            PackageNotFoundError,
            zipfile.BadZipFile,
            KeyError,
            ValueError,
        ) as error:
            raise DomainError(
                code=ErrorCode.VALIDATION_ERROR,
                message="Filen kunde inte läsas som ett Word-dokument (.docx).",
                details={"filename": filename},
            ) from error

        blocks: list[DocxSourceBlock] = []
        index = 0
        for paragraph in parsed.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style_name = paragraph.style.name if paragraph.style is not None else ""
            blocks.append(DocxSourceBlock(kind=_block_kind(style_name), text=text, index=index))
            index += 1

        notes: list[str] = []
        table_count = len(parsed.tables)
        if table_count:
            for table in parsed.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    line = " | ".join(cell for cell in cells if cell)
                    if line:
                        blocks.append(
                            DocxSourceBlock(kind=DocxBlockKind.BODY, text=line, index=index)
                        )
                        index += 1
            notes.append(f"tables_flattened: {table_count}")
        image_count = len(parsed.inline_shapes)
        if image_count:
            notes.append(f"inline_images_not_imported: {image_count}")

        if not blocks:
            raise DomainError(
                code=ErrorCode.VALIDATION_ERROR,
                message="Word-dokumentet saknar läsbart innehåll.",
                details={"filename": filename},
            )

        return extract_native_exam_from_docx_blocks(
            blocks=tuple(blocks),
            document_id=document_id,
            source_filename=filename,
            source_sha256=hashlib.sha256(content).hexdigest(),
            notes=tuple(notes),
        )
