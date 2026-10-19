"""Pictures the chunker cannot carry as PNG or JPEG never vanish silently (TASK-SKRIPT-39-05-09)."""

import io
from uuid import uuid4

import pytest
from docx import Document
from docx.document import Document as DocxDocument
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml.ns import qn
from PIL import Image

from skriptoteket.domain.curated_apps.exam_workspace.docx_chunks import (
    DocxImage,
    DocxMissingImage,
    asset_id_for,
)
from skriptoteket.domain.curated_apps.exam_workspace.docx_source_review import review_source_item
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_body import (
    NativeTextSegment,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_document import (
    NativeExamDocument,
    NativeExamDocumentOrigin,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_export_rules import (
    native_export_blockers,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_vocabulary import (
    NativeItemReviewState,
    NativeParseOrigin,
)
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_docx_chunker import (  # noqa: E501
    PythonDocxExamChunker,
)
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_labeler_block_image import (  # noqa: E501
    DocxBlockImageRenderer,
)

pytestmark = pytest.mark.unit


def _png() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (24, 12), (10, 200, 30)).save(buffer, format="PNG")
    return buffer.getvalue()


def _animated_gif() -> bytes:
    first = Image.new("RGB", (5, 4), (220, 20, 20))
    second = Image.new("RGB", (5, 4), (20, 20, 220))
    buffer = io.BytesIO()
    first.save(buffer, format="GIF", save_all=True, append_images=[second], duration=100, loop=0)
    return buffer.getvalue()


def _transparent_gif() -> bytes:
    image = Image.new("P", (2, 1), 0)
    image.putpalette([255, 0, 0, 0, 0, 0] + [0, 0, 0] * 254)
    image.putpixel((1, 0), 1)
    buffer = io.BytesIO()
    image.save(buffer, format="GIF", transparency=0)
    return buffer.getvalue()


def _save(document: DocxDocument) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _chunk(document: DocxDocument):
    return PythonDocxExamChunker().chunk(content=_save(document))


def _replace_picture_part(document: DocxDocument, media_type: str, extension: str) -> None:
    """Point the document's only picture at an embedded part of another media type."""
    paragraph = document.element.body.findall(qn("w:p"))[-1]
    blip = next(paragraph.iter(qn("a:blip")))
    part = Part(PackURI(f"/word/media/figur.{extension}"), media_type, b"\x01\x00\x00\x00", None)
    blip.set(qn("r:embed"), document.part.relate_to(part, RELATIONSHIP_TYPE.IMAGE))


def _unsupported_document(media_type: str, extension: str) -> DocxDocument:
    document = Document()
    document.add_picture(io.BytesIO(_png()))
    _replace_picture_part(document, media_type, extension)
    return document


def test_a_gif_picture_becomes_its_first_frame_as_png() -> None:
    document = Document()
    document.add_picture(io.BytesIO(_animated_gif()))

    (chunk,) = _chunk(document)

    (image,) = chunk.images
    assert image.media_type == "image/png"
    assert image.payload.startswith(b"\x89PNG\r\n\x1a\n")
    assert (image.width, image.height) == (5, 4)
    assert image.asset_id == asset_id_for(image.payload)
    assert chunk.missing_images == ()
    with Image.open(io.BytesIO(image.payload)) as decoded:
        assert decoded.size == (5, 4)
        assert decoded.convert("RGB").getpixel((0, 0)) == (220, 20, 20)


def test_a_transparent_gif_keeps_its_transparency_in_png() -> None:
    document = Document()
    document.add_picture(io.BytesIO(_transparent_gif()))

    (chunk,) = _chunk(document)
    (image,) = chunk.images

    assert image.media_type == "image/png"
    with Image.open(io.BytesIO(image.payload)) as decoded:
        assert decoded.mode == "RGBA"
        assert decoded.getpixel((0, 0)) == (255, 0, 0, 0)
        assert decoded.getpixel((1, 0)) == (0, 0, 0, 255)


def test_valid_png_and_jpeg_payloads_stay_byte_exact() -> None:
    png = _png()
    jpeg_buffer = io.BytesIO()
    Image.new("RGB", (40, 20), (10, 20, 30)).save(jpeg_buffer, format="JPEG")
    jpeg = jpeg_buffer.getvalue()
    document = Document()
    document.add_picture(io.BytesIO(png))
    document.paragraphs[0].add_run(" och ")
    document.paragraphs[0].add_run().add_picture(io.BytesIO(jpeg))

    (chunk,) = _chunk(document)

    assert [image.payload for image in chunk.images] == [png, jpeg]
    assert [image.media_type for image in chunk.images] == ["image/png", "image/jpeg"]
    assert chunk.missing_images == ()


@pytest.mark.parametrize(
    ("media_type", "extension"),
    [("image/x-emf", "emf"), ("image/x-wmf", "wmf")],
)
def test_an_unrenderable_picture_keeps_a_named_descriptor_and_its_chunk(
    media_type: str, extension: str
) -> None:
    document = _unsupported_document(media_type, extension)

    (chunk,) = _chunk(document)

    assert chunk.kind == "drawing"
    assert chunk.text == ""
    assert chunk.images == ()
    assert chunk.missing_images == (
        DocxMissingImage(reference=f"/word/media/figur.{extension}", media_type=media_type),
    )


def test_a_picture_without_a_part_keeps_the_relationship_id_as_its_name() -> None:
    document = Document()
    document.add_picture(io.BytesIO(_png()))
    paragraph = document.element.body.findall(qn("w:p"))[-1]
    next(paragraph.iter(qn("a:blip"))).set(qn("r:embed"), "rIdMissing")

    (chunk,) = _chunk(document)

    assert chunk.images == ()
    assert chunk.missing_images == (
        DocxMissingImage(reference="rIdMissing", media_type="application/octet-stream"),
    )


def test_a_picture_that_is_only_linked_keeps_a_named_descriptor() -> None:
    document = Document()
    document.add_picture(io.BytesIO(_png()))
    paragraph = document.element.body.findall(qn("w:p"))[-1]
    blip = next(paragraph.iter(qn("a:blip")))
    blip.attrib.pop(qn("r:embed"))
    blip.set(qn("r:link"), "rIdLinked")

    (chunk,) = _chunk(document)

    assert chunk.images == ()
    assert chunk.missing_images == (
        DocxMissingImage(reference="rIdLinked", media_type="application/octet-stream"),
    )


def test_a_gif_pillow_cannot_decode_keeps_a_named_descriptor() -> None:
    document = Document()
    document.add_picture(io.BytesIO(_png()))
    _replace_picture_part(document, "image/gif", "trasig.gif")

    (chunk,) = _chunk(document)

    assert chunk.images == ()
    assert chunk.missing_images == (
        DocxMissingImage(reference="/word/media/figur.trasig.gif", media_type="image/gif"),
    )


def test_an_image_only_block_is_rendered_whole_when_the_renderer_can(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    # The block renderer rejects EMF today, so a faithful whole-block render is proven with
    # a synthetic success; the real unsupported outcome is the review test below.
    rendered = _png()
    document = _unsupported_document("image/x-emf", "emf")
    calls: list[str] = []

    def fake_render(self: DocxBlockImageRenderer, element, *, anchor: str) -> bytes:
        calls.append(anchor)
        return rendered

    monkeypatch.setattr(DocxBlockImageRenderer, "render", fake_render)

    (chunk,) = _chunk(document)

    assert chunk.missing_images == ()
    assert chunk.images == (
        DocxImage(
            payload=rendered,
            media_type="image/png",
            offset=0,
            marked_offset=0,
            width=24,
            height=12,
        ),
    )
    assert calls == ["body:0001"]


def test_an_inline_picture_with_surrounding_text_is_never_rendered_whole(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Före ")
    paragraph.add_run().add_picture(io.BytesIO(_png()))
    paragraph.add_run(" efter")
    _replace_picture_part(document, "image/x-emf", "emf")

    def fake_render(self: DocxBlockImageRenderer, element, *, anchor: str) -> bytes:
        raise AssertionError("an inline picture must not be rendered whole")

    monkeypatch.setattr(DocxBlockImageRenderer, "render", fake_render)

    (chunk,) = _chunk(document)

    assert chunk.images == ()
    assert chunk.missing_images == (
        DocxMissingImage(
            reference="/word/media/figur.emf", media_type="image/x-emf", offset=5, marked_offset=5
        ),
    )


def test_an_unavailable_render_leaves_review_state_and_blocks_export() -> None:
    document = _unsupported_document("image/x-emf", "emf")
    (chunk,) = _chunk(document)

    item = review_source_item(
        (chunk,), "Frågan kunde inte tolkas. Kontrollera källtexten manuellt."
    )

    assert item.review.state is NativeItemReviewState.REVIEW_REQUIRED
    assert item.review.parse_origin is NativeParseOrigin.SOURCE_FALLBACK
    assert item.review.reasons == (
        "Frågan kunde inte tolkas. Kontrollera källtexten manuellt.",
        "Bilden /word/media/figur.emf (image/x-emf) kan inte visas i provet. "
        "Granska frågan manuellt.",
    )
    (paragraph,) = item.body
    assert paragraph.segments == (
        NativeTextSegment(text="[Bild saknas: /word/media/figur.emf (image/x-emf)]", marks=()),
    )
    blockers = native_export_blockers(
        NativeExamDocument(
            document_id=uuid4(),
            revision=1,
            title="Prov",
            items=(item,),
            assets=(),
            origin=NativeExamDocumentOrigin(kind="created"),
        )
    )
    assert "unbuilt_question" in {blocker.reason for blocker in blockers}
