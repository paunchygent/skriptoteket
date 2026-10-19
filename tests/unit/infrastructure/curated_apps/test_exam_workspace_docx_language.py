"""The student DOCX and the facit print document print their fixed words in Swedish (2026-10-18).

Each corpus exam is converted from its recorded provider rounds. No exam is named, and
English exams get the same Swedish words. Only lines the export generates are checked.

The student DOCX (W2) prints headings and inline points and no answer key at all, so its
generated lines and the absence of every answer-key heading and key row are checked. The
facit print document (W1) adds the answer-key lines and criteria after each keyed question;
those lines, with the facit headings, are checked in Swedish too. Exam text
copied from the source may contain any word.
"""

import io
import re

import pytest
from docx import Document

from skriptoteket.domain.curated_apps.exam_conversion.exam_print_house_template import (
    ExamPrintVariant,
)
from skriptoteket.domain.curated_apps.exam_workspace.exam_language_labels import (
    SWEDISH_EXAM_LABELS,
)
from skriptoteket.domain.curated_apps.exam_workspace.exam_print_document import (
    PrintAnswerKey,
    PrintHeading,
    build_print_document,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_document import (
    NativeExamDocument,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_vocabulary import (
    NativeExamItemKind,
)
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_docx_writer import (  # noqa: E501
    ExamWorkspaceDocxWriter,
)
from tests.unit.domain.curated_apps.exam_workspace.exam_print_document_test_support import (
    entry_blocks,
)
from tests.unit.scripts.exam_conversion_corpus_support import RECORDED_STEMS, replayed

_ENGLISH = {"Question", "Gap", "Points", "missing", "Answer key", "Correct answer", "Pair"}
# A generated answer-key row begins with one of the Swedish key labels, never exam text.
_KEY_ROW = re.compile(r"^(Rätt svar|Lucka \d+|Par): ")
# A question-text paragraph carrying its points inline, such as "… prövas? (1 p)".
_INLINE_POINTS = re.compile(r"\(\d[\d,]* p\)$")
_CHOICE_KINDS = (
    NativeExamItemKind.SINGLE_CHOICE,
    NativeExamItemKind.MULTIPLE_RESPONSE,
)


def _docx_generated_lines(stem: str) -> list[str]:
    """Headings and inline points of the student DOCX export."""
    document, _ = replayed(stem)
    parsed = Document(io.BytesIO(ExamWorkspaceDocxWriter().build_docx_bytes(document)))
    lines: list[str] = []
    for paragraph in parsed.paragraphs:
        style = paragraph.style.name if paragraph.style is not None else ""
        text = paragraph.text
        if style in {"Heading 2", "Heading 3"}:
            lines.append(text)
        elif match := _INLINE_POINTS.search(text):
            lines.append(match.group(0))
    return lines


def _facit_generated_lines(stem: str) -> tuple[list[str], bool]:
    """The facit document's generated lines, and whether the exam has a keyed item.

    Generated lines are every PrintAnswerKey line and the facit headings.
    A line the exam itself prints (a body line) is not generated export wording.
    """
    document, _ = replayed(stem)
    print_document = build_print_document(document, variant=ExamPrintVariant.FACIT)
    lines: list[str] = []
    for entry in print_document.entries:
        for block in entry_blocks(entry):
            if isinstance(block, PrintAnswerKey):
                lines.extend(block.lines)
            elif isinstance(block, PrintHeading):
                lines.append(block.text)
    printed = {line for item in document.items for line in item.plain_text_lines()}
    return [line for line in lines if line not in printed], _has_keyed_item(document)


def _has_keyed_item(document: NativeExamDocument) -> bool:
    """Whether any question carries an answer the facit prints as a key or criteria."""
    for item in document.items:
        if item.kind is NativeExamItemKind.FREE_TEXT:
            if item.criteria:
                return True
        elif item.kind in _CHOICE_KINDS:
            if item.answer_key.correct_choice_ids:
                return True
        elif item.kind in (NativeExamItemKind.GAP_FILL, NativeExamItemKind.INLINE_CHOICE):
            if any(gap.accepted_values or gap.correct_option_id for gap in item.gaps):
                return True
        elif item.kind is NativeExamItemKind.ORDERING:
            return True
        elif item.kind is NativeExamItemKind.MATCHING:
            if item.answer_key.match_pairs:
                return True
    return False


def _words_of(text: str) -> set[str]:
    return set(re.findall(r"\w+", text))


@pytest.mark.parametrize("stem", RECORDED_STEMS)
def test_student_docx_generated_lines_hold_no_english_fixed_word(stem: str) -> None:
    document, _ = replayed(stem)
    labels = SWEDISH_EXAM_LABELS
    assert labels.question == "Fråga"
    forbidden_words = {word for phrase in _ENGLISH for word in phrase.split()}
    # Lines the exam itself prints are not generated export wording.
    printed = {line for item in document.items for line in item.plain_text_lines()}
    generated = [line for line in _docx_generated_lines(stem) if line not in printed]
    assert generated
    for line in generated:
        foreign = _words_of(line) & forbidden_words
        assert not foreign, f"{stem}: {line!r} holds {sorted(foreign)}"


@pytest.mark.parametrize("stem", RECORDED_STEMS)
def test_student_docx_owns_every_inline_points_and_no_answer_key(stem: str) -> None:
    document, _ = replayed(stem)
    parsed = Document(io.BytesIO(ExamWorkspaceDocxWriter().build_docx_bytes(document)))
    texts = [paragraph.text for paragraph in parsed.paragraphs]
    for item in document.items:
        if item.points is None or not float(item.points).is_integer():
            continue
        assert any(text.endswith(f" ({int(item.points)} p)") for text in texts), (
            f"{stem}: no inline points for item {item.item_id}"
        )
    key_headings = [
        paragraph.text
        for paragraph in parsed.paragraphs
        if paragraph.style is not None and paragraph.style.name == "Heading 2"
    ]
    assert key_headings == [], f"{stem}: student DOCX prints answer-key headings {key_headings!r}"
    key_rows = [text for text in texts if _KEY_ROW.match(text)]
    assert key_rows == [], f"{stem}: student DOCX prints answer-key rows {key_rows!r}"


@pytest.mark.parametrize("stem", RECORDED_STEMS)
def test_facit_key_lines_hold_no_english_fixed_word(stem: str) -> None:
    generated, keyed = _facit_generated_lines(stem)
    if keyed:
        assert generated, f"{stem}: the facit prints no generated key lines"
    forbidden_words = {word for phrase in _ENGLISH for word in phrase.split()}
    for line in generated:
        foreign = _words_of(line) & forbidden_words
        assert not foreign, f"{stem}: {line!r} holds {sorted(foreign)}"
