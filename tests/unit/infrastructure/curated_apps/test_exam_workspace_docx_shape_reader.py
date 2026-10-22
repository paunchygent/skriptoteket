"""Challenge the Word shape reader (TASK-SKRIPT-39-04-22 unit 1) with real and synthetic DrawingML.

The two fixture drawings are real Word output from ``metodprov_med_bilder.docx``; every other
case is hand-written DrawingML read through a python-docx document, so the theme, the style
chain and the document defaults are the ones python-docx exposes. Each rule family of D2 has an
unsupported case: the reader must raise ``UnsupportedShapeFigure`` and never return a partial
figure.
"""

from __future__ import annotations

from pathlib import Path
from typing import TypedDict

import pytest
from docx import Document
from docx.document import Document as DocxDocument
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.oxml.xmlchemy import BaseOxmlElement

from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_docx_shape_figure import (  # noqa: E501
    SUPPORTED_PRESETS,
    Arrowhead,
    ArrowheadKind,
    FigureShape,
    LabelAnchor,
    LabelParagraph,
    LabelRun,
    ShapeLabel,
    ShapeOutline,
    UnsupportedShapeFigure,
)
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_docx_shape_reader import (  # noqa: E501
    read_shape_figure,
)

pytestmark = pytest.mark.unit

_FIXTURE = (
    Path(__file__).resolve().parents[3]
    / "fixtures"
    / "exam_conversion"
    / "pdf_corpus"
    / "metodprov_med_bilder.docx"
)
_NS = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
    'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:wpg="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup" '
    'xmlns:wpc="http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas" '
    'xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main"'
)
_URI = {
    "wps": "http://schemas.microsoft.com/office/word/2010/wordprocessingShape",
    "wpg": "http://schemas.microsoft.com/office/word/2010/wordprocessingGroup",
    "wpc": "http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas",
}
_BLACK_LINE = '<a:ln w="19050"><a:solidFill><a:srgbClr val="000000"/></a:solidFill></a:ln>'


def _theme(accent1: str = "4472C4", effect_third: str = "<a:effectLst/>") -> bytes:
    """A theme in the Office 2013 layout: the colour scheme and the three style lists."""
    return (
        '<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" name="T">'
        "<a:themeElements>"
        '<a:clrScheme name="T">'
        '<a:dk1><a:sysClr val="windowText" lastClr="000000"/></a:dk1>'
        '<a:lt1><a:sysClr val="window" lastClr="FFFFFF"/></a:lt1>'
        '<a:dk2><a:srgbClr val="44546A"/></a:dk2>'
        '<a:lt2><a:srgbClr val="E7E6E6"/></a:lt2>'
        f'<a:accent1><a:srgbClr val="{accent1}"/></a:accent1>'
        '<a:accent2><a:srgbClr val="ED7D31"/></a:accent2>'
        "</a:clrScheme>"
        '<a:fmtScheme name="T">'
        '<a:fillStyleLst><a:solidFill><a:schemeClr val="phClr"/></a:solidFill>'
        '<a:gradFill><a:gsLst><a:gs pos="0"><a:schemeClr val="phClr"/></a:gs>'
        "</a:gsLst></a:gradFill>"
        '<a:solidFill><a:srgbClr val="123456"/></a:solidFill></a:fillStyleLst>'
        "<a:lnStyleLst>"
        '<a:ln w="6350"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill>'
        '<a:prstDash val="solid"/></a:ln>'
        '<a:ln w="12700"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill>'
        '<a:prstDash val="solid"/></a:ln>'
        '<a:ln w="19050"><a:gradFill><a:gsLst/></a:gradFill></a:ln>'
        "</a:lnStyleLst>"
        "<a:effectStyleLst>"
        "<a:effectStyle><a:effectLst/></a:effectStyle>"
        "<a:effectStyle><a:effectLst/></a:effectStyle>"
        f"<a:effectStyle>{effect_third}</a:effectStyle>"
        "</a:effectStyleLst>"
        "</a:fmtScheme>"
        "</a:themeElements></a:theme>"
    ).encode()


def _document(theme: bytes | None = None, *, drop_theme: bool = False) -> DocxDocument:
    document = Document()
    if drop_theme:
        rel_id = next(
            key for key, rel in document.part.rels.items() if rel.reltype == RELATIONSHIP_TYPE.THEME
        )
        document.part.rels.pop(rel_id)
    elif theme is not None:
        document.part.part_related_by(RELATIONSHIP_TYPE.THEME)._blob = theme
    return document


def _drawing(
    graphic: str, *, uri: str = "wps", cx: int = 1000, cy: int = 500, anchor: bool = False
) -> BaseOxmlElement:
    container = "wp:anchor" if anchor else "wp:inline"
    data_uri = _URI.get(uri, uri)
    return parse_xml(
        f"<w:drawing {_NS}><{container}>"
        f'<wp:extent cx="{cx}" cy="{cy}"/><wp:docPr id="1" name="d"/>'
        f'<a:graphic><a:graphicData uri="{data_uri}">{graphic}</a:graphicData></a:graphic>'
        f"</{container}></w:drawing>"
    )


def _wsp(
    *,
    prst: str = "rect",
    off: tuple[int, int] = (0, 0),
    ext: tuple[int, int] = (1000, 500),
    xfrm_attrs: str = "",
    geometry: str | None = None,
    guides: str = "",
    fill: str = "<a:noFill/>",
    line: str = _BLACK_LINE,
    sp_extra: str = "",
    style: str = "",
    txbx: str = "",
    body: str = "<wps:bodyPr/>",
    connector: bool = False,
) -> str:
    geometry = geometry or f'<a:prstGeom prst="{prst}"><a:avLst>{guides}</a:avLst></a:prstGeom>'
    non_visual = "<wps:cNvCnPr/>" if connector else "<wps:cNvSpPr/>"
    return (
        f'<wps:wsp><wps:cNvPr id="2" name="s"/>{non_visual}<wps:spPr>'
        f'<a:xfrm {xfrm_attrs}><a:off x="{off[0]}" y="{off[1]}"/>'
        f'<a:ext cx="{ext[0]}" cy="{ext[1]}"/></a:xfrm>'
        f"{geometry}{fill}{line}{sp_extra}</wps:spPr>{style}{txbx}{body}</wps:wsp>"
    )


def _group(
    children: str,
    *,
    tag: str = "wpg:wgp",
    off: tuple[int, int] = (0, 0),
    ext: tuple[int, int] = (1000, 500),
    ch_off: tuple[int, int] = (0, 0),
    ch_ext: tuple[int, int] = (1000, 500),
    xfrm_attrs: str = "",
) -> str:
    return (
        f'<{tag}><wpg:cNvPr id="3" name="g"/><wpg:cNvGrpSpPr/><wpg:grpSpPr>'
        f'<a:xfrm {xfrm_attrs}><a:off x="{off[0]}" y="{off[1]}"/>'
        f'<a:ext cx="{ext[0]}" cy="{ext[1]}"/>'
        f'<a:chOff x="{ch_off[0]}" y="{ch_off[1]}"/>'
        f'<a:chExt cx="{ch_ext[0]}" cy="{ch_ext[1]}"/></a:xfrm></wpg:grpSpPr>'
        f"{children}</{tag}>"
    )


def _canvas(children: str, *, background: str = "<wpc:bg><a:noFill/></wpc:bg>") -> str:
    return (
        f"<wpc:wpc>{background}<wpc:whole><a:ln><a:noFill/></a:ln></wpc:whole>{children}</wpc:wpc>"
    )


class _LabelXml(TypedDict):
    txbx: str
    body: str


def _label(*paragraphs: str, body: str = "<wps:bodyPr/>") -> _LabelXml:
    content = "".join(paragraphs)
    return {"txbx": f"<wps:txbx><w:txbxContent>{content}</w:txbxContent></wps:txbx>", "body": body}


def _para(text: str, *, jc: str | None = None, rpr: str = "", ppr: str = "") -> str:
    jc_xml = f'<w:jc w:val="{jc}"/>' if jc else ""
    ppr_xml = f"<w:pPr>{ppr}{jc_xml}</w:pPr>" if ppr or jc else ""
    rpr_xml = f"<w:rPr>{rpr}</w:rPr>" if rpr else ""
    return f"<w:p>{ppr_xml}<w:r>{rpr_xml}<w:t>{text}</w:t></w:r></w:p>"


def _read_one(shape_xml: str, *, document: DocxDocument | None = None) -> FigureShape:
    figure = read_shape_figure(_drawing(shape_xml), document or _document())
    assert len(figure.shapes) == 1
    return figure.shapes[0]


def _rejected(drawing: BaseOxmlElement, needle: str, document: DocxDocument | None = None) -> None:
    with pytest.raises(UnsupportedShapeFigure) as caught:
        read_shape_figure(drawing, document or _document())
    assert needle in caught.value.reason


def _rejected_shape(shape_xml: str, needle: str, document: DocxDocument | None = None) -> None:
    _rejected(_drawing(shape_xml), needle, document)


def _fixture_drawing(doc_pr_id: str) -> tuple[DocxDocument, BaseOxmlElement]:
    document = Document(str(_FIXTURE))
    for drawing in document.element.body.iter(qn("w:drawing")):
        doc_pr = drawing.find(f".//{qn('wp:docPr')}")
        if doc_pr is not None and doc_pr.get("id") == doc_pr_id:
            return document, drawing
    raise AssertionError(f"drawing {doc_pr_id} missing from the fixture")


def _centered(text: str, anchor: LabelAnchor) -> ShapeLabel:
    return ShapeLabel(
        paragraphs=(LabelParagraph((LabelRun(text, 10.0, "000000"),), "center"),),
        inset_left_emu=36000,
        inset_top_emu=36000,
        inset_right_emu=36000,
        inset_bottom_emu=36000,
        anchor=anchor,
    )


_BLACK_19050 = ShapeOutline(color="000000", width_emu=19050)


# --- real Word output ------------------------------------------------------------------------


def test_population_group_reads_the_two_ellipses_exactly() -> None:
    document, drawing = _fixture_drawing("9001")

    figure = read_shape_figure(drawing, document)

    assert (figure.width_emu, figure.height_emu) == (4320000, 1655999)
    assert figure.shapes == (
        FigureShape(
            preset="ellipse",
            x_emu=0,
            y_emu=0,
            width_emu=4320000,
            height_emu=1655999,
            fill=None,
            outline=_BLACK_19050,
            label=_centered("Svenska gymnasieelever (populationen)", "t"),
        ),
        FigureShape(
            preset="ellipse",
            x_emu=2232000,
            y_emu=684000,
            width_emu=1800000,
            height_emu=827999,
            fill="DDEBF7",
            outline=_BLACK_19050,
            label=_centered("Elever på ett program (urvalet)", "ctr"),
        ),
    )
    assert not figure.plain_lines


def test_selection_group_reads_boxes_and_arrows_in_paint_order() -> None:
    document, drawing = _fixture_drawing("9002")

    figure = read_shape_figure(drawing, document)

    assert (figure.width_emu, figure.height_emu) == (5040000, 864000)
    expected = [
        ("rect", 0, 1296000, 864000, None, "Alla elever på skolan"),
        ("rightArrow", 1404000, 432000, 288000, "404040", None),
        ("rect", 1908000, 1296000, 864000, None, "Elever som är i biblioteket"),
        ("rightArrow", 3311999, 432000, 288000, "404040", None),
        ("rect", 3744000, 1296000, 864000, None, "Svarar på frågan"),
    ]
    assert [
        (s.preset, s.x_emu, s.width_emu, s.height_emu, s.fill, _text(s)) for s in figure.shapes
    ] == expected
    assert [s.y_emu for s in figure.shapes] == [0, 288000, 0, 288000, 0]
    assert all(s.outline == _BLACK_19050 for s in figure.shapes)
    boxes = [s for s in figure.shapes if s.preset == "rect"]
    assert all(s.label == _centered(_text(s) or "", "ctr") for s in boxes)
    assert all(s.label is None for s in figure.shapes if s.preset == "rightArrow")


def _text(shape: FigureShape) -> str | None:
    if shape.label is None:
        return None
    return "".join(run.text for p in shape.label.paragraphs for run in p.runs)


def test_a_picture_drawing_of_the_fixture_is_outside_d2() -> None:
    document, drawing = _fixture_drawing("9206")

    _rejected(drawing, "picture", document)


# --- containers ------------------------------------------------------------------------------


def test_single_shape_with_ext_equal_to_the_extent_fills_the_box_whatever_its_offset() -> None:
    drawing = _drawing(_wsp(off=(777, 888), ext=(4000, 2000)), cx=4000, cy=2000)

    figure = read_shape_figure(drawing, _document())

    assert (figure.width_emu, figure.height_emu) == (4000, 2000)
    shape = figure.shapes[0]
    assert (shape.x_emu, shape.y_emu, shape.width_emu, shape.height_emu) == (0, 0, 4000, 2000)


def test_rotated_single_shape_keeps_its_own_ext_centred_in_the_extent_box() -> None:
    # Word sizes wp:extent to the rotated bounding box: a 1000000 square turned 45 degrees.
    drawing = _drawing(
        _wsp(xfrm_attrs='rot="2700000"', off=(777, 888), ext=(1000000, 1000000)),
        cx=1414214,
        cy=1414214,
    )

    figure = read_shape_figure(drawing, _document())

    assert (figure.width_emu, figure.height_emu) == (1414214, 1414214)
    shape = figure.shapes[0]
    assert (shape.x_emu, shape.y_emu, shape.width_emu, shape.height_emu) == (
        207107,
        207107,
        1000000,
        1000000,
    )
    assert shape.rotation_deg == 45.0


def test_single_shape_ext_centres_independently_in_x_and_y() -> None:
    drawing = _drawing(_wsp(ext=(3000, 500)), cx=4000, cy=2000)

    shape = read_shape_figure(drawing, _document()).shapes[0]

    assert (shape.x_emu, shape.y_emu, shape.width_emu, shape.height_emu) == (500, 750, 3000, 500)


def test_an_anchored_drawing_reads_like_an_inline_one() -> None:
    figure = read_shape_figure(_drawing(_wsp(), anchor=True, cx=900, cy=300), _document())

    assert (figure.width_emu, figure.height_emu) == (900, 300)


def test_group_children_scale_from_child_space_to_the_figure() -> None:
    child = _wsp(off=(1500, 2500), ext=(1000, 500))
    group = _group(child, ext=(4000, 1000), ch_off=(1000, 2000), ch_ext=(2000, 500))

    figure = read_shape_figure(_drawing(group, uri="wpg", cx=4000, cy=1000), _document())

    shape = figure.shapes[0]
    assert (shape.x_emu, shape.y_emu, shape.width_emu, shape.height_emu) == (1000, 1000, 2000, 1000)


def test_a_top_level_group_maps_onto_the_wp_extent_not_its_own_offset() -> None:
    child = _wsp(off=(0, 0), ext=(100, 100))
    group = _group(child, off=(5000, 6000), ext=(200, 200), ch_off=(0, 0), ch_ext=(100, 100))

    figure = read_shape_figure(_drawing(group, uri="wpg", cx=200, cy=200), _document())

    shape = figure.shapes[0]
    assert (shape.x_emu, shape.y_emu, shape.width_emu, shape.height_emu) == (0, 0, 200, 200)


def test_nested_group_composes_inside_the_parent_mapping() -> None:
    inner = _group(
        _wsp(off=(10, 10), ext=(20, 20)),
        tag="wpg:grpSp",
        off=(100, 200),
        ext=(60, 60),
        ch_off=(0, 0),
        ch_ext=(30, 30),
    )
    outer = _group(inner, ext=(2000, 1000), ch_off=(0, 0), ch_ext=(1000, 500))

    figure = read_shape_figure(_drawing(outer, uri="wpg", cx=2000, cy=1000), _document())

    # child (10,10)-(30,30) -> nested space x2: (20,20) size 40 -> shifted to (120,220)
    # -> outer scale x2: (240, 440) size 80.
    shape = figure.shapes[0]
    assert (shape.x_emu, shape.y_emu, shape.width_emu, shape.height_emu) == (240, 440, 80, 80)


def test_group_with_shapes_and_nested_group_keeps_paint_order() -> None:
    first = _wsp(prst="rect", off=(0, 0), ext=(10, 10))
    nested = _group(
        _wsp(prst="ellipse", off=(0, 0), ext=(10, 10)),
        tag="wpg:grpSp",
        ext=(10, 10),
        ch_ext=(10, 10),
    )
    last = _wsp(prst="triangle", off=(0, 0), ext=(10, 10))
    group = _group(first + nested + last)

    figure = read_shape_figure(_drawing(group, uri="wpg"), _document())

    assert [s.preset for s in figure.shapes] == ["rect", "ellipse", "triangle"]


def test_canvas_children_keep_their_emu_coordinates() -> None:
    canvas = _canvas(_wsp(off=(300, 400), ext=(500, 600)))

    figure = read_shape_figure(_drawing(canvas, uri="wpc", cx=2000, cy=2000), _document())

    shape = figure.shapes[0]
    assert (shape.x_emu, shape.y_emu, shape.width_emu, shape.height_emu) == (300, 400, 500, 600)


def test_canvas_with_a_group_applies_the_group_mapping() -> None:
    inner = _group(
        _wsp(off=(0, 0), ext=(10, 10)),
        tag="wpg:grpSp",
        off=(100, 100),
        ext=(50, 50),
        ch_ext=(10, 10),
    )

    figure = read_shape_figure(_drawing(_canvas(inner), uri="wpc"), _document())

    shape = figure.shapes[0]
    assert (shape.x_emu, shape.y_emu, shape.width_emu, shape.height_emu) == (100, 100, 50, 50)


# --- geometry --------------------------------------------------------------------------------


@pytest.mark.parametrize("preset", sorted(SUPPORTED_PRESETS))
def test_every_supported_preset_is_read(preset: str) -> None:
    connector = preset in {"line", "straightConnector1"}

    shape = _read_one(_wsp(prst=preset, connector=connector))

    assert shape.preset == preset


def test_rotation_flips_and_adjustments_are_read() -> None:
    shape = _read_one(
        _wsp(
            prst="roundRect",
            xfrm_attrs='rot="5400000" flipH="1" flipV="true"',
            guides='<a:gd name="adj" fmla="val 25000"/>',
        )
    )

    assert shape.rotation_deg == 90.0
    assert shape.flip_h is True
    assert shape.flip_v is True
    assert shape.adjustments == (("adj", 25000),)


def test_rotation_is_normalised_into_a_turn() -> None:
    assert _read_one(_wsp(xfrm_attrs='rot="-5400000"')).rotation_deg == 270.0
    assert _read_one(_wsp(xfrm_attrs='rot="21600000"')).rotation_deg == 0.0


def test_rotated_shape_in_an_evenly_scaled_group_keeps_its_rotation() -> None:
    group = _group(
        _wsp(xfrm_attrs='rot="2700000"', ext=(10, 10)), ext=(2000, 1000), ch_ext=(2000, 1000)
    )

    figure = read_shape_figure(_drawing(group, uri="wpg", cx=2000, cy=1000), _document())

    assert figure.shapes[0].rotation_deg == 45.0


# --- fill and outline ------------------------------------------------------------------------


def test_no_fill_element_and_no_style_means_no_fill() -> None:
    shape = _read_one(_wsp(fill=""))

    assert shape.fill is None


def test_solid_fill_colour_is_upper_case_hex() -> None:
    shape = _read_one(_wsp(fill='<a:solidFill><a:srgbClr val="ddebf7"/></a:solidFill>'))

    assert shape.fill == "DDEBF7"


def test_outline_defaults_to_9525_and_reads_dash_and_arrowheads() -> None:
    line = (
        '<a:ln><a:solidFill><a:srgbClr val="FF0000"/></a:solidFill><a:prstDash val="sysDash"/>'
        '<a:round/><a:headEnd type="oval"/><a:tailEnd type="triangle" w="lg" len="sm"/></a:ln>'
    )

    shape = _read_one(_wsp(line=line, prst="straightConnector1", connector=True))

    assert shape.outline == ShapeOutline(
        color="FF0000",
        width_emu=9525,
        dash="sysDash",
        head=Arrowhead("oval"),
        tail=Arrowhead("triangle", width="lg", length="sm"),
    )


@pytest.mark.parametrize("kind", ["triangle", "arrow", "stealth", "oval", "diamond"])
def test_each_arrowhead_kind_is_read(kind: ArrowheadKind) -> None:
    line = (
        '<a:ln w="12700"><a:solidFill><a:srgbClr val="000000"/></a:solidFill>'
        f'<a:tailEnd type="{kind}"/></a:ln>'
    )

    shape = _read_one(_wsp(line=line))

    assert shape.outline is not None
    assert shape.outline.tail == Arrowhead(kind)


def test_arrowhead_type_none_is_no_arrowhead() -> None:
    line = (
        '<a:ln><a:solidFill><a:srgbClr val="000000"/></a:solidFill>'
        '<a:headEnd type="none"/><a:tailEnd/></a:ln>'
    )

    shape = _read_one(_wsp(line=line))

    assert shape.outline is not None
    assert (shape.outline.head, shape.outline.tail) == (None, None)


def test_no_fill_outline_is_no_outline() -> None:
    assert _read_one(_wsp(line="<a:ln><a:noFill/></a:ln>")).outline is None


def test_no_outline_element_and_no_style_is_no_outline() -> None:
    assert _read_one(_wsp(line="")).outline is None


# --- colours and the theme -------------------------------------------------------------------


def _scheme_fill(inner: str, scheme: str = "accent1") -> str:
    return f'<a:solidFill><a:schemeClr val="{scheme}">{inner}</a:schemeClr></a:solidFill>'


@pytest.mark.parametrize(
    ("inner", "expected"),
    [
        ("", "4472C4"),
        ('<a:lumMod val="60000"/><a:lumOff val="40000"/>', "8FAADC"),
        ('<a:lumMod val="40000"/><a:lumOff val="60000"/>', "B4C7E7"),
        ('<a:lumMod val="75000"/>', "2F5597"),
        ('<a:shade val="50000"/>', "2F528F"),
        ('<a:alpha val="100000"/>', "4472C4"),
    ],
)
def test_scheme_colour_resolves_through_the_theme_with_modifiers(inner: str, expected: str) -> None:
    shape = _read_one(_wsp(fill=_scheme_fill(inner)), document=_document(_theme()))

    assert shape.fill == expected


def test_tint_lightens_towards_white_and_full_tint_is_the_colour() -> None:
    document = _document(_theme())

    full = _read_one(_wsp(fill=_scheme_fill('<a:tint val="100000"/>')), document=document)
    none = _read_one(_wsp(fill=_scheme_fill('<a:tint val="0"/>')), document=document)

    assert full.fill == "4472C4"
    assert none.fill == "FFFFFF"


@pytest.mark.parametrize(
    ("scheme", "expected"),
    [
        ("tx1", "000000"),
        ("bg1", "FFFFFF"),
        ("tx2", "44546A"),
        ("bg2", "E7E6E6"),
        ("accent2", "ED7D31"),
    ],
)
def test_scheme_names_map_to_the_theme_slots(scheme: str, expected: str) -> None:
    shape = _read_one(_wsp(fill=_scheme_fill("", scheme)), document=_document(_theme()))

    assert shape.fill == expected


def test_sys_and_preset_colours() -> None:
    sys_fill = '<a:solidFill><a:sysClr val="windowText" lastClr="1A2B3C"/></a:solidFill>'
    white = '<a:solidFill><a:prstClr val="white"/></a:solidFill>'
    black = '<a:solidFill><a:prstClr val="black"/></a:solidFill>'

    assert _read_one(_wsp(fill=sys_fill)).fill == "1A2B3C"
    assert _read_one(_wsp(fill=white)).fill == "FFFFFF"
    assert _read_one(_wsp(fill=black)).fill == "000000"


def test_a_concrete_colour_never_reads_the_theme() -> None:
    shape = _read_one(
        _wsp(fill='<a:solidFill><a:srgbClr val="010203"/></a:solidFill>'),
        document=_document(drop_theme=True),
    )

    assert shape.fill == "010203"


_OFFICE_STYLE = (
    "<wps:style>"
    '<a:lnRef idx="{ln}"><a:schemeClr val="accent1"><a:shade val="50000"/></a:schemeClr></a:lnRef>'
    '<a:fillRef idx="{fill}"><a:schemeClr val="accent1"/></a:fillRef>'
    '<a:effectRef idx="{effect}"><a:scrgbClr r="0" g="0" b="0"/></a:effectRef>'
    '<a:fontRef idx="minor"><a:schemeClr val="lt1"/></a:fontRef>'
    "</wps:style>"
)


def _style(ln: int = 2, fill: int = 1, effect: int = 0) -> str:
    return _OFFICE_STYLE.format(ln=ln, fill=fill, effect=effect)


def test_style_fill_and_line_references_resolve_through_the_theme() -> None:
    shape = _read_one(_wsp(fill="", line="", style=_style()), document=_document(_theme()))

    assert shape.fill == "4472C4"
    assert shape.outline == ShapeOutline(color="2F528F", width_emu=12700)


def test_style_index_zero_means_no_fill_and_no_outline() -> None:
    shape = _read_one(
        _wsp(fill="", line="", style=_style(ln=0, fill=0)), document=_document(_theme())
    )

    assert (shape.fill, shape.outline) == (None, None)


def test_a_shape_that_sets_its_own_fill_ignores_the_style_fill() -> None:
    shape = _read_one(
        _wsp(fill='<a:solidFill><a:srgbClr val="00FF00"/></a:solidFill>', style=_style()),
        document=_document(_theme()),
    )

    assert shape.fill == "00FF00"


def test_style_line_takes_the_theme_width_unless_the_shape_sets_one() -> None:
    explicit = '<a:ln w="3175"/>'
    shape = _read_one(_wsp(fill="", line=explicit, style=_style()), document=_document(_theme()))

    assert shape.outline == ShapeOutline(color="2F528F", width_emu=3175)


def test_style_fill_may_be_a_concrete_theme_colour() -> None:
    shape = _read_one(
        _wsp(fill="", line="", style=_style(ln=0, fill=3)), document=_document(_theme())
    )

    assert shape.fill == "123456"


def test_a_style_without_a_theme_is_unsupported() -> None:
    _rejected_shape(_wsp(fill="", style=_style()), "theme", _document(drop_theme=True))


def test_a_scheme_colour_without_a_theme_is_unsupported() -> None:
    _rejected_shape(_wsp(fill=_scheme_fill("")), "theme", _document(drop_theme=True))


def test_style_effect_reference_to_an_empty_effect_list_is_fine() -> None:
    shape = _read_one(_wsp(style=_style(ln=0, fill=0, effect=2)), document=_document(_theme()))

    assert shape.fill is None


# --- labels ----------------------------------------------------------------------------------


def test_label_runs_marks_size_colour_alignment_and_breaks() -> None:
    content = (
        '<w:p><w:pPr><w:jc w:val="end"/></w:pPr>'
        '<w:r><w:rPr><w:b/><w:i w:val="1"/><w:u w:val="single"/><w:sz w:val="28"/>'
        '<w:color w:val="ff0000"/></w:rPr><w:t>Hej</w:t><w:tab/><w:t>du</w:t><w:br/>'
        "<w:t>nu</w:t></w:r>"
        '<w:r><w:rPr><w:b w:val="0"/><w:u w:val="none"/></w:rPr>'
        '<w:t xml:space="preserve"> x</w:t></w:r>'
        "</w:p>"
    )

    shape = _read_one(_wsp(**_label(content)))

    assert shape.label == ShapeLabel(
        paragraphs=(
            LabelParagraph(
                runs=(
                    LabelRun("Hej du\nnu", 14.0, "FF0000", bold=True, italic=True, underline=True),
                    LabelRun(" x", 11.0, "000000"),
                ),
                align="right",
            ),
        ),
    )


@pytest.mark.parametrize(
    ("jc", "expected"),
    [("left", "left"), ("start", "left"), ("center", "center"), ("right", "right"),
     ("end", "right"), ("both", "justify"), (None, "left")],
)  # fmt: skip
def test_paragraph_alignment_values(jc: str | None, expected: str) -> None:
    shape = _read_one(_wsp(**_label(_para("t", jc=jc))))

    assert shape.label is not None
    assert shape.label.paragraphs[0].align == expected


def test_paragraph_style_alignment_and_size_apply_when_the_paragraph_sets_none() -> None:
    document = _document()
    style = document.styles.add_style("BoxText", 1)
    style.element.get_or_add_pPr().append(parse_xml(f'<w:jc {_NS} w:val="center"/>'))
    style.element.get_or_add_rPr().append(parse_xml(f'<w:sz {_NS} w:val="36"/>'))

    shape = _read_one(
        _wsp(**_label(_para("t", ppr='<w:pStyle w:val="BoxText"/>'))), document=document
    )

    assert shape.label is not None
    assert shape.label.paragraphs[0].align == "center"
    assert shape.label.paragraphs[0].runs[0].size_pt == 18.0


def test_a_style_chain_supplies_bold_from_the_paragraph_style() -> None:
    document = _document()
    style = document.styles.add_style("Strong box", 1)
    style.element.get_or_add_rPr().append(parse_xml(f"<w:b {_NS}/>"))

    shape = _read_one(
        _wsp(**_label(_para("t", ppr='<w:pStyle w:val="Strongbox"/>'))), document=document
    )

    assert shape.label is not None
    assert shape.label.paragraphs[0].runs[0].bold is True


def test_size_falls_back_to_the_document_default() -> None:
    shape = _read_one(_wsp(**_label(_para("t"))))

    assert shape.label is not None
    assert shape.label.paragraphs[0].runs[0].size_pt == 11.0  # python-docx template: sz 22


def test_size_is_ten_points_when_nothing_sets_one() -> None:
    document = _document()
    defaults = document.styles.element.find(qn("w:docDefaults"))
    assert defaults is not None
    document.styles.element.remove(defaults)

    shape = _read_one(_wsp(**_label(_para("t"))), document=document)

    assert shape.label is not None
    assert shape.label.paragraphs[0].runs[0].size_pt == 10.0


def test_auto_colour_takes_the_style_font_colour_else_black() -> None:
    styled = _read_one(
        _wsp(style=_style(ln=0, fill=0), **_label(_para("t", rpr='<w:color w:val="auto"/>'))),
        document=_document(_theme()),
    )
    plain = _read_one(_wsp(**_label(_para("t"))))

    assert styled.label is not None and plain.label is not None
    assert styled.label.paragraphs[0].runs[0].color == "FFFFFF"
    assert plain.label.paragraphs[0].runs[0].color == "000000"


def test_body_properties_read_insets_and_anchor() -> None:
    body = '<wps:bodyPr lIns="1" tIns="2" rIns="3" bIns="4" anchor="b" wrap="square" rtlCol="0"/>'

    shape = _read_one(_wsp(**_label(_para("t"), body=body)))

    assert shape.label is not None
    assert (
        shape.label.inset_left_emu,
        shape.label.inset_top_emu,
        shape.label.inset_right_emu,
        shape.label.inset_bottom_emu,
        shape.label.anchor,
    ) == (1, 2, 3, 4, "b")


def test_body_properties_default_to_word_insets_and_top_anchor() -> None:
    shape = _read_one(_wsp(**_label(_para("t"))))

    assert shape.label is not None
    assert (
        shape.label.inset_left_emu,
        shape.label.inset_top_emu,
        shape.label.inset_right_emu,
        shape.label.inset_bottom_emu,
        shape.label.anchor,
    ) == (91440, 45720, 91440, 45720, "t")


def test_two_paragraphs_keep_their_own_alignment_and_an_empty_one_stays() -> None:
    shape = _read_one(_wsp(**_label(_para("a", jc="center"), "<w:p/>", _para("b", jc="right"))))

    assert shape.label is not None
    assert [(p.align, [r.text for r in p.runs]) for p in shape.label.paragraphs] == [
        ("center", ["a"]),
        ("left", []),
        ("right", ["b"]),
    ]


def test_a_label_without_any_text_is_no_label() -> None:
    assert _read_one(_wsp(**_label("<w:p/>", "<w:p><w:r><w:t></w:t></w:r></w:p>"))).label is None


def test_bookmarks_and_proof_marks_around_runs_are_ignored() -> None:
    content = (
        '<w:p><w:bookmarkStart w:id="0" w:name="b"/><w:proofErr w:type="spellStart"/>'
        '<w:r><w:t>ok</w:t></w:r><w:bookmarkEnd w:id="0"/></w:p>'
    )

    shape = _read_one(_wsp(**_label(content)))

    assert _text(shape) == "ok"


# --- unsupported: one case per rule family ---------------------------------------------------


def test_custom_geometry_is_unsupported() -> None:
    _rejected_shape(_wsp(geometry="<a:custGeom><a:pathLst/></a:custGeom>"), "custGeom")


def test_an_unlisted_preset_is_unsupported() -> None:
    _rejected_shape(_wsp(prst="star5"), "star5")


def test_a_guide_formula_other_than_val_is_unsupported() -> None:
    _rejected_shape(_wsp(guides='<a:gd name="adj" fmla="*/ w 1 2"/>'), "fmla")


@pytest.mark.parametrize(
    ("fill", "needle"),
    [
        ("<a:gradFill><a:gsLst/></a:gradFill>", "gradFill"),
        ('<a:blipFill><a:blip r:embed="rId9"/></a:blipFill>', "blipFill"),
        ('<a:pattFill prst="pct5"/>', "pattFill"),
        ("<a:grpFill/>", "grpFill"),
    ],
)
def test_non_solid_fills_are_unsupported(fill: str, needle: str) -> None:
    _rejected_shape(_wsp(fill=fill), needle)


def test_a_gradient_outline_is_unsupported() -> None:
    _rejected_shape(_wsp(line="<a:ln><a:gradFill><a:gsLst/></a:gradFill></a:ln>"), "gradFill")


def test_a_custom_dash_is_unsupported() -> None:
    line = (
        '<a:ln><a:solidFill><a:srgbClr val="000000"/></a:solidFill>'
        '<a:custDash><a:ds d="1" sp="1"/></a:custDash></a:ln>'
    )
    _rejected_shape(_wsp(line=line), "custDash")


def test_an_unknown_dash_is_unsupported() -> None:
    line = (
        '<a:ln><a:solidFill><a:srgbClr val="000000"/></a:solidFill><a:prstDash val="zig"/></a:ln>'
    )
    _rejected_shape(_wsp(line=line), "zig")


def test_an_unknown_arrowhead_type_is_unsupported() -> None:
    line = (
        '<a:ln><a:solidFill><a:srgbClr val="000000"/></a:solidFill>'
        '<a:tailEnd type="swoosh"/></a:ln>'
    )
    _rejected_shape(_wsp(line=line), "swoosh")


@pytest.mark.parametrize(
    "extra",
    [
        '<a:effectLst><a:outerShdw blurRad="1"/></a:effectLst>',
        "<a:effectDag><a:cont/></a:effectDag>",
        '<a:scene3d><a:camera prst="orthographicFront"/></a:scene3d>',
        '<a:sp3d extrusionH="100"/>',
    ],
)
def test_spPr_effects_are_unsupported(extra: str) -> None:
    _rejected_shape(_wsp(sp_extra=extra), "effect" if "ffect" in extra else "3d")


def test_an_empty_effect_list_is_no_effect() -> None:
    assert _read_one(_wsp(sp_extra="<a:effectLst/>")).preset == "rect"


def test_a_style_effect_reference_to_a_shadow_is_unsupported() -> None:
    theme = _theme(effect_third='<a:effectLst><a:outerShdw blurRad="40000"/></a:effectLst>')

    _rejected_shape(_wsp(style=_style(ln=0, fill=0, effect=3)), "effect", _document(theme))


def test_a_gradient_theme_fill_is_unsupported() -> None:
    _rejected_shape(_wsp(fill="", style=_style(fill=2)), "gradFill", _document(_theme()))


def test_a_non_solid_theme_line_is_unsupported() -> None:
    _rejected_shape(_wsp(line="", style=_style(ln=3)), "gradFill", _document(_theme()))


def test_alpha_below_full_is_unsupported() -> None:
    fill = '<a:solidFill><a:srgbClr val="000000"><a:alpha val="50000"/></a:srgbClr></a:solidFill>'
    _rejected_shape(_wsp(fill=fill), "alpha")


def test_a_colour_modifier_outside_the_list_is_unsupported() -> None:
    fill = '<a:solidFill><a:srgbClr val="808080"><a:satMod val="50000"/></a:srgbClr></a:solidFill>'
    _rejected_shape(_wsp(fill=fill), "satMod")


def test_a_preset_colour_other_than_black_and_white_is_unsupported() -> None:
    _rejected_shape(_wsp(fill='<a:solidFill><a:prstClr val="red"/></a:solidFill>'), "red")


def test_a_linked_text_box_is_unsupported() -> None:
    _rejected_shape(
        _wsp(sp_extra="").replace("</wps:wsp>", "<wps:linkedTxbx id='1' seq='0'/></wps:wsp>"),
        "linkedTxbx",
    )


def test_a_picture_inside_a_group_is_unsupported() -> None:
    picture = '<pic:pic><pic:blipFill><a:blip r:embed="rId9"/></pic:blipFill></pic:pic>'
    group = _group(_wsp() + picture)

    _rejected(_drawing(group, uri="wpg"), "pic")


def test_a_graphic_frame_inside_a_canvas_is_unsupported() -> None:
    _rejected(_drawing(_canvas("<wpg:graphicFrame/>"), uri="wpc"), "graphicFrame")


@pytest.mark.parametrize(
    "uri", ["http://schemas.openxmlformats.org/drawingml/2006/chart",
            "http://schemas.openxmlformats.org/drawingml/2006/diagram",
            "http://schemas.openxmlformats.org/drawingml/2006/picture"],
)  # fmt: skip
def test_a_graphic_that_is_not_a_shape_container_is_unsupported(uri: str) -> None:
    _rejected(_drawing("<a:ext/>", uri=uri), uri.rsplit("/", 1)[1])


def test_a_group_with_zero_child_extent_is_unsupported() -> None:
    _rejected(_drawing(_group(_wsp(), ch_ext=(0, 500)), uri="wpg"), "chExt")


def test_a_shape_turned_off_the_axes_in_an_unevenly_scaled_group_is_unsupported() -> None:
    group = _group(_wsp(xfrm_attrs='rot="2700000"'), ext=(2000, 500), ch_ext=(1000, 500))

    _rejected(_drawing(group, uri="wpg", cx=2000, cy=500), "scaled")


def test_a_canvas_background_with_a_fill_is_unsupported() -> None:
    canvas = _canvas(
        _wsp(), background='<wpc:bg><a:solidFill><a:srgbClr val="FFFF00"/></a:solidFill></wpc:bg>'
    )

    _rejected(_drawing(canvas, uri="wpc"), "bg")


def test_a_canvas_whole_with_an_outline_is_unsupported() -> None:
    canvas = _canvas(_wsp()).replace(
        "<a:ln><a:noFill/></a:ln>",
        '<a:ln><a:solidFill><a:srgbClr val="000000"/></a:solidFill></a:ln>',
    )

    _rejected(_drawing(canvas, uri="wpc"), "whole")


def test_a_canvas_without_background_elements_is_fine() -> None:
    canvas = f"<wpc:wpc>{_wsp()}</wpc:wpc>"

    assert len(read_shape_figure(_drawing(canvas, uri="wpc"), _document()).shapes) == 1


def test_a_hyperlink_in_a_label_is_unsupported() -> None:
    content = '<w:p><w:hyperlink w:anchor="x"><w:r><w:t>l</w:t></w:r></w:hyperlink></w:p>'
    _rejected_shape(_wsp(**_label(content)), "hyperlink")


def test_a_table_in_a_label_is_unsupported() -> None:
    _rejected_shape(_wsp(**_label("<w:tbl/>")), "tbl")


@pytest.mark.parametrize(
    ("run_child", "name"),
    [
        ('<w:fldChar w:fldCharType="begin"/>', "w:fldChar"),
        ("<w:instrText>PAGE</w:instrText>", "w:instrText"),
        ("<w:drawing/>", "w:drawing"),
        ("<w:pict/>", "w:pict"),
        ("<w:object/>", "w:object"),
        ('<w:footnoteReference w:id="1"/>', "w:footnoteReference"),
        ('<w:endnoteReference w:id="1"/>', "w:endnoteReference"),
        ('<w:commentReference w:id="1"/>', "w:commentReference"),
    ],
)
def test_a_run_with_a_field_object_or_reference_is_unsupported(run_child: str, name: str) -> None:
    content = f"<w:p><w:r><w:t>a</w:t>{run_child}</w:r></w:p>"
    _rejected_shape(_wsp(**_label(content)), name)


def test_a_numbered_label_paragraph_is_unsupported() -> None:
    ppr = '<w:numPr><w:ilvl w:val="0"/><w:numId w:val="3"/></w:numPr>'
    _rejected_shape(_wsp(**_label(_para("a", ppr=ppr))), "numPr")


@pytest.mark.parametrize("anchor", ["just", "dist"])
def test_an_anchor_other_than_top_centre_bottom_is_unsupported(anchor: str) -> None:
    _rejected_shape(_wsp(**_label(_para("t"), body=f'<wps:bodyPr anchor="{anchor}"/>')), anchor)


@pytest.mark.parametrize("attrs", ['vert="vert"', 'rot="5400000"'])
def test_vertical_or_rotated_text_is_unsupported(attrs: str) -> None:
    _rejected_shape(_wsp(**_label(_para("t"), body=f"<wps:bodyPr {attrs}/>")), "bodyPr")


def test_a_group_without_shapes_is_unsupported() -> None:
    _rejected(_drawing(_group(""), uri="wpg"), "no shapes")


def test_a_drawing_without_a_container_is_unsupported() -> None:
    _rejected(parse_xml(f"<w:drawing {_NS}/>"), "wp:inline")


def test_a_drawing_with_an_empty_or_negative_extent_is_unsupported() -> None:
    _rejected(_drawing(_wsp(), cx=0, cy=0), "extent")
    _rejected(_drawing(_wsp(), cx=-5), "extent")


def test_a_plain_horizontal_line_has_a_zero_height_box_and_is_plain_lines() -> None:
    line = _wsp(prst="line", ext=(4000, 0), connector=True)

    figure = read_shape_figure(_drawing(line, cx=4000, cy=0), _document())

    assert (figure.width_emu, figure.height_emu) == (4000, 0)
    assert figure.plain_lines


def test_a_shape_without_geometry_is_unsupported() -> None:
    _rejected_shape(_wsp(geometry=" "), "geometry")


def test_two_wsp_in_a_single_shape_graphic_is_unsupported() -> None:
    _rejected(_drawing(_wsp() + _wsp()), "one")


def test_the_theme_is_parsed_without_resolving_entities() -> None:
    hostile = (
        b'<?xml version="1.0"?><!DOCTYPE t [<!ENTITY x SYSTEM "file:///etc/passwd">]>'
        b'<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        b'<a:themeElements><a:clrScheme><a:accent1><a:srgbClr val="&x;"/></a:accent1>'
        b"</a:clrScheme></a:themeElements></a:theme>"
    )

    _rejected_shape(_wsp(fill=_scheme_fill("")), "theme", _document(hostile))


@pytest.mark.parametrize("mark", ["strike", "highlight", "vanish", "vertAlign"])
def test_a_run_mark_the_label_model_lacks_is_unsupported(mark: str) -> None:
    val = {"highlight": ' w:val="yellow"', "vertAlign": ' w:val="superscript"'}.get(mark, "")
    _rejected_shape(
        _wsp(**_label(_para("t", rpr=f"<w:{mark}{val}/>"))),
        f"w:{mark}",
    )


def test_a_mark_switched_off_is_not_a_mark() -> None:
    rpr = '<w:strike w:val="0"/><w:vertAlign w:val="baseline"/><w:highlight w:val="none"/>'

    assert _text(_read_one(_wsp(**_label(_para("t", rpr=rpr))))) == "t"


def test_word_saved_text_box_markup_is_read() -> None:
    """Word writes a text warp of "none" into bodyPr and the resolved hex beside a theme colour."""
    body = (
        '<wps:bodyPr rot="0" vert="horz" wrap="square" lIns="91440" tIns="45720" rIns="91440" '
        'bIns="45720" anchor="ctr" anchorCtr="0"><a:prstTxWarp prst="textNoShape"><a:avLst/>'
        "</a:prstTxWarp><a:noAutofit/></wps:bodyPr>"
    )
    rpr = '<w:color w:val="FFFFFF" w:themeColor="background1"/>'

    shape = _read_one(_wsp(**_label(_para("t", rpr=rpr), body=body)))

    assert shape.label is not None
    assert (shape.label.paragraphs[0].runs[0].color, shape.label.anchor) == ("FFFFFF", "ctr")


def test_a_shrinking_text_fit_is_unsupported() -> None:
    body = '<wps:bodyPr><a:normAutofit fontScale="92500" lnSpcReduction="10000"/></wps:bodyPr>'

    _rejected_shape(_wsp(**_label(_para("t"), body=body)), "normAutofit")


# --- extension lists ---------------------------------------------------------------------------

# Word writes the VML fallback's fill and line into a:extLst; nothing draws them, so the reader
# ignores the list. The hidden colours differ from the visible ones so a leak would show.
_EXT_LST = (
    '<a:extLst><a:ext uri="{909E8E84-426E-40DD-AFC4-6F175D3DCCD1}"><a14:hiddenFill>'
    '<a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill></a14:hiddenFill></a:ext>'
    '<a:ext uri="{91240B29-F687-4F45-9708-019B960494DF}"><a14:hiddenLine w="38100">'
    '<a:solidFill><a:srgbClr val="FF0000"/></a:solidFill></a14:hiddenLine></a:ext></a:extLst>'
)


def test_an_extension_list_in_the_shape_properties_is_ignored() -> None:
    assert _read_one(_wsp(sp_extra=_EXT_LST)) == _read_one(_wsp())


def test_an_extension_list_in_the_outline_is_ignored() -> None:
    line = _BLACK_LINE.replace("</a:ln>", f"{_EXT_LST}</a:ln>")

    assert _read_one(_wsp(line=line)) == _read_one(_wsp())


def test_an_extension_list_in_a_theme_line_style_is_ignored() -> None:
    plain = _theme()
    extended = plain.replace(b"</a:ln>", _EXT_LST.encode() + b"</a:ln>", 1).replace(
        b"<a:theme ", b'<a:theme xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main" '
    )
    assert extended != plain
    shape = _wsp(line="", style=_style(ln=1, fill=0))

    assert _read_one(shape, document=_document(extended)) == _read_one(
        shape, document=_document(plain)
    )


def test_an_extension_list_in_the_group_properties_is_ignored() -> None:
    plain = _group(_wsp())
    extended = plain.replace("</a:xfrm></wpg:grpSpPr>", f"</a:xfrm>{_EXT_LST}</wpg:grpSpPr>")
    assert extended != plain

    assert read_shape_figure(_drawing(extended, uri="wpg"), _document()) == read_shape_figure(
        _drawing(plain, uri="wpg"), _document()
    )


def test_an_extension_list_in_the_body_properties_is_ignored() -> None:
    plain = _wsp(**_label(_para("t"), body="<wps:bodyPr/>"))
    extended = _wsp(**_label(_para("t"), body=f"<wps:bodyPr>{_EXT_LST}</wps:bodyPr>"))

    assert _read_one(extended) == _read_one(plain)


def test_an_extension_list_in_the_canvas_layers_is_ignored() -> None:
    plain = _canvas(_wsp())
    extended = plain.replace(
        "<a:ln><a:noFill/></a:ln>", f"<a:ln><a:noFill/>{_EXT_LST}</a:ln>{_EXT_LST}"
    ).replace("<wpc:bg><a:noFill/></wpc:bg>", f"<wpc:bg><a:noFill/>{_EXT_LST}</wpc:bg>")
    assert extended.count("a:extLst>") == 6

    assert read_shape_figure(_drawing(extended, uri="wpc"), _document()) == read_shape_figure(
        _drawing(plain, uri="wpc"), _document()
    )
