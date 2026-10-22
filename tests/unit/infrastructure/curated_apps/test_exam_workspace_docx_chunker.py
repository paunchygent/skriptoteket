"""Source-preservation proof for the DOCX chunking adapter."""

import io
import zipfile
from dataclasses import replace
from pathlib import Path
from uuid import uuid4

import pytest
from docx import Document
from docx.document import Document as DocxDocument
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import parse_xml
from lxml import etree  # type: ignore[import-untyped]  # transitive of python-docx, no stubs

from skriptoteket.domain.curated_apps.exam_workspace.docx_build_assets import asset_manifest
from skriptoteket.domain.curated_apps.exam_workspace.docx_chunks import DocxChunk
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_body import (
    NativeAssetSegment,
    NativeParagraph,
    NativeTextSegment,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_document import (
    NativeExamDocument,
    NativeExamDocumentOrigin,
)
from skriptoteket.domain.errors import DomainError, ErrorCode
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_docx_chunker import (  # noqa: E501
    PythonDocxExamChunker,
)
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_docx_writer import (  # noqa: E501
    ExamWorkspaceDocxWriter,
)
from tests.unit.domain.curated_apps.exam_workspace.native_exam_test_builders import (
    make_free_text_item,
)

_QUIZ = Path(
    "tests/fixtures/exam_conversion/docx_corpus/chapter_4_poetry_and_lyrics_quiz_med_facit.docx"
)


def _bytes(document: DocxDocument) -> bytes:
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _text_projection(chunks: tuple[DocxChunk, ...]) -> tuple[DocxChunk, ...]:
    """The chunks as the text projection: text, marked text and pictures, no formatting."""
    return tuple(replace(chunk, runs=(), paragraphs=()) for chunk in chunks)


def test_quiz_preserves_all_paragraphs_and_soft_break_options() -> None:
    content = _QUIZ.read_bytes()
    chunks = PythonDocxExamChunker().chunk(content=content)

    assert len(chunks) == 109
    assert tuple(chunk.text for chunk in chunks) == tuple(
        paragraph.text for paragraph in Document(io.BytesIO(content)).paragraphs
    )
    assert (chunks[4].anchor, chunks[4].kind, chunks[4].text, chunks[4].marked_text) == (
        "body:0007",
        "paragraph",
        "A. deep\nB. strange\nC. dignified\nD. thrifty",
        "A. deep\nB. strange\nC. dignified\nD. thrifty",
    )
    assert chunks[5].text == "Correct answer: A. deep"
    assert chunks[5].marked_text == "Correct answer: <i>A. deep</i>"
    assert chunks[-1].text == "Correct answer: alliteration"


def test_body_order_anchors_and_whitespace_survive_tables_and_empty_paragraphs() -> None:
    document = Document()
    document.add_heading("  A heading  ", level=2)
    document.add_paragraph("")
    document.add_paragraph("  First\nsecond\tline  ")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "  Cell\nsoft break  "
    table.cell(0, 0).add_paragraph("Another paragraph")
    table.cell(0, 1).text = "Right"
    table.cell(1, 0).text = "Bottom"
    document.add_paragraph("After the table")
    second_table = document.add_table(rows=1, cols=1)
    second_table.cell(0, 0).text = "Second table"
    content = _bytes(document)

    first = PythonDocxExamChunker().chunk(content=content)

    assert _text_projection(first) == (
        DocxChunk("body:0001", "heading", "  A heading  ", "  A heading  "),
        DocxChunk("body:0003", "paragraph", "  First\nsecond\tline  ", "  First\nsecond\tline  "),
        DocxChunk(
            "body:0004",
            "table",
            "  Cell\nsoft break  \nAnother paragraph | Right\nBottom | ",
            "  Cell\nsoft break  \nAnother paragraph | Right\nBottom | ",
        ),
        DocxChunk("body:0005", "paragraph", "After the table", "After the table"),
        DocxChunk("body:0006", "table", "Second table", "Second table"),
    )
    assert PythonDocxExamChunker().chunk(content=content) == first


def test_merged_and_nested_cells_preserve_text_once_in_order() -> None:
    document = Document()
    table = document.add_table(rows=1, cols=2)
    cell = table.cell(0, 0).merge(table.cell(0, 1))
    cell.text = "Before"
    nested = cell.add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested\nlines"
    cell.paragraphs[-1].text = "After"

    assert _text_projection(PythonDocxExamChunker().chunk(content=_bytes(document))) == (
        DocxChunk(
            "body:0001", "table", "Before\nNested\nlines\nAfter", "Before\nNested\nlines\nAfter"
        ),
    )


def test_course_exam_marks_teacher_italics_only_in_model_text() -> None:
    content = _QUIZ.with_name("svenska_2_kursavslutande_prov_med_facit.docx").read_bytes()
    chunks = {chunk.anchor: chunk for chunk in PythonDocxExamChunker().chunk(content=content)}

    assert chunks["body:0086"].text == (
        "a) De väntar inte. är en huvudsats\n"
        "b) I bisatsen att de inte väntar står satsadverbialet före det finita verbet."
    )
    assert chunks["body:0086"].marked_text == (
        "a) <i>De väntar inte.</i> är en <i>huvudsats</i>\n"
        "b) I bisatsen <i>att de inte väntar</i> står satsadverbialet <i>före</i> "
        "det finita verbet."
    )
    assert chunks["body:0063"].marked_text.startswith(
        "a) I ordet lekfullheten är <i>lek</i> <i>R</i>"
    )


def test_italic_runs_in_paragraphs_and_tables_are_marked() -> None:
    document = Document()
    paragraph = document.add_paragraph("Svar: ")
    paragraph.add_run("två ").italic = True
    paragraph.add_run("ord").italic = True
    paragraph.add_run(" slut")
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].add_run("Cell ")
    cell.paragraphs[0].add_run("svar").italic = True

    chunks = PythonDocxExamChunker().chunk(content=_bytes(document))

    assert [(c.text, c.marked_text) for c in chunks] == [
        ("Svar: två ord slut", "Svar: <i>två ord</i> slut"),
        ("Cell svar", "Cell <i>svar</i>"),
    ]


@pytest.mark.parametrize("content", [b"", b"not a docx", b"PK\x03\x04truncated"])
def test_invalid_docx_reports_swedish_domain_error(content: bytes) -> None:
    with pytest.raises(DomainError) as exc_info:
        PythonDocxExamChunker().chunk(content=content)

    assert exc_info.value.code is ErrorCode.VALIDATION_ERROR
    assert exc_info.value.message == "Filen kunde inte läsas som ett Word-dokument (.docx)."


def test_malformed_document_xml_reports_domain_error() -> None:
    stream = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(_bytes(Document()))) as source:
        with zipfile.ZipFile(stream, "w") as target:
            for entry in source.infolist():
                target.writestr(
                    entry,
                    b"<broken" if entry.filename == "word/document.xml" else source.read(entry),
                )

    with pytest.raises(DomainError) as exc_info:
        PythonDocxExamChunker().chunk(content=stream.getvalue())

    assert exc_info.value.code is ErrorCode.VALIDATION_ERROR
    assert "Word-dokument" in exc_info.value.message


@pytest.mark.parametrize("layout", ["empty", "whitespace", "table"])
def test_empty_docx_reports_swedish_domain_error(layout: str) -> None:
    document = Document()
    if layout == "whitespace":
        document.add_paragraph(" \t\n ")
    elif layout == "table":
        document.add_table(rows=2, cols=2)

    with pytest.raises(DomainError) as exc_info:
        PythonDocxExamChunker().chunk(content=_bytes(document))

    assert exc_info.value.code is ErrorCode.VALIDATION_ERROR
    assert exc_info.value.message == "Word-dokumentet saknar läsbart innehåll."


# --- drawn figures: Word shapes become pictures (TASK-SKRIPT-39-04-22 D1-D4) ---------------

_PNG_SIGNATURE = b"\x89PNG\r\n\x1a\n"
_NS = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:v="urn:schemas-microsoft-com:vml" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:wpg="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup" '
    'xmlns:wpc="http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas"'
)
_URI = {
    "wps": "http://schemas.microsoft.com/office/word/2010/wordprocessingShape",
    "wpg": "http://schemas.microsoft.com/office/word/2010/wordprocessingGroup",
    "wpc": "http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas",
}
_LINE = '<a:ln w="19050"><a:solidFill><a:srgbClr val="000000"/></a:solidFill></a:ln>'


def _wsp(
    prst: str = "rect",
    *,
    off: tuple[int, int] = (0, 0),
    ext: tuple[int, int] = (1905000, 952500),
    label: str | None = None,
    geometry: str | None = None,
    line: str = _LINE,
    properties: str = "",
) -> str:
    geometry = geometry or f'<a:prstGeom prst="{prst}"><a:avLst/></a:prstGeom>'
    text = ""
    if label is not None:
        text = (
            "<wps:txbx><w:txbxContent><w:p><w:r>"
            f"<w:t>{label}</w:t></w:r></w:p></w:txbxContent></wps:txbx>"
        )
    return (
        f'<wps:wsp><wps:cNvPr id="2" name="s">{properties}</wps:cNvPr><wps:cNvSpPr/><wps:spPr>'
        f'<a:xfrm><a:off x="{off[0]}" y="{off[1]}"/><a:ext cx="{ext[0]}" cy="{ext[1]}"/></a:xfrm>'
        f'{geometry}<a:solidFill><a:srgbClr val="FFCC00"/></a:solidFill>{line}</wps:spPr>'
        f"{text}<wps:bodyPr/></wps:wsp>"
    )


def _drawing(kind: str, shapes: str, *, ext: tuple[int, int] = (1905000, 952500)) -> str:
    graphic = shapes
    if kind == "wpg":
        graphic = (
            '<wpg:wgp><wpg:cNvPr id="3" name="g"/><wpg:cNvGrpSpPr/><wpg:grpSpPr>'
            f'<a:xfrm><a:off x="0" y="0"/><a:ext cx="{ext[0]}" cy="{ext[1]}"/>'
            f'<a:chOff x="0" y="0"/><a:chExt cx="{ext[0]}" cy="{ext[1]}"/></a:xfrm>'
            f"</wpg:grpSpPr>{shapes}</wpg:wgp>"
        )
    elif kind == "wpc":
        graphic = f"<wpc:wpc><wpc:bg><a:noFill/></wpc:bg>{shapes}</wpc:wpc>"
    return (
        "<w:drawing><wp:inline>"
        f'<wp:extent cx="{ext[0]}" cy="{ext[1]}"/><wp:docPr id="1" name="d"/>'
        f'<a:graphic><a:graphicData uri="{_URI[kind]}">{graphic}</a:graphicData></a:graphic>'
        "</wp:inline></w:drawing>"
    )


def _group(*labels: str) -> str:
    return _drawing(
        "wpg",
        "".join(
            _wsp("ellipse", off=(index * 950000, 0), ext=(900000, 900000), label=label)
            for index, label in enumerate(labels)
        ),
    )


def _paragraph_with(*run_contents: str, text: str = "") -> bytes:
    document = Document()
    document.add_paragraph("Fråga 1")
    paragraph = document.add_paragraph(text)
    for content in run_contents:
        paragraph._p.append(parse_xml(f"<w:r {_NS}>{content}</w:r>"))
    document.add_paragraph("Fråga 2")
    return _bytes(document)


def _drawn_chunk(content: bytes) -> DocxChunk:
    (chunk,) = [
        c for c in PythonDocxExamChunker().chunk(content=content) if c.anchor == "body:0002"
    ]
    return chunk


def test_a_group_of_shapes_with_labels_becomes_one_picture() -> None:
    chunk = _drawn_chunk(_paragraph_with(_group("Population", "Urval")))

    assert (chunk.kind, chunk.text, chunk.missing_images) == ("drawing", "", ())
    (image,) = chunk.images
    assert image.payload.startswith(_PNG_SIGNATURE)
    assert image.media_type == "image/png"
    assert (image.offset, image.marked_offset) == (0, 0)
    assert image.width and image.height
    assert image.docx_shapes is not None
    assert "wpg:wgp" in image.docx_shapes
    assert "Population" in image.docx_shapes


@pytest.mark.parametrize("kind", ["wps", "wpc"])
def test_a_single_shape_and_a_canvas_become_one_picture_each(kind: str) -> None:
    shapes = _wsp("rect", label="Ruta")
    chunk = _drawn_chunk(_paragraph_with(_drawing(kind, shapes)))

    (image,) = chunk.images
    assert (chunk.kind, chunk.text, chunk.missing_images) == ("drawing", "", ())
    assert image.payload.startswith(_PNG_SIGNATURE)
    assert image.docx_shapes is not None
    assert "wps:wsp" in image.docx_shapes


def test_the_choice_drawing_is_drawn_and_its_vml_fallback_ignored() -> None:
    content = (
        '<mc:AlternateContent><mc:Choice Requires="wpg">'
        f"{_group('Population')}"
        "</mc:Choice><mc:Fallback><w:pict><v:rect/></w:pict></mc:Fallback></mc:AlternateContent>"
    )

    chunk = _drawn_chunk(_paragraph_with(content))

    (image,) = chunk.images
    assert image.payload.startswith(_PNG_SIGNATURE)
    assert chunk.missing_images == ()


def test_a_group_with_a_custom_geometry_shape_keeps_the_drawing_chunk_without_a_picture() -> None:
    custom = _wsp(
        geometry=(
            "<a:custGeom><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/>"
            '<a:rect l="0" t="0" r="r" b="b"/><a:pathLst/></a:custGeom>'
        )
    )
    chunk = _drawn_chunk(_paragraph_with(_drawing("wpg", _wsp("rect") + custom)))

    assert (chunk.kind, chunk.text, chunk.images, chunk.missing_images) == ("drawing", "", (), ())


def test_a_drawing_of_plain_lines_gets_no_picture() -> None:
    lines = _wsp("line") + _wsp("straightConnector1", off=(0, 100000))
    chunk = _drawn_chunk(_paragraph_with(_drawing("wpg", lines)))

    assert (chunk.kind, chunk.images, chunk.missing_images) == ("drawing", (), ())


def test_two_figures_in_one_paragraph_give_two_pictures_in_order() -> None:
    first = _drawing("wps", _wsp("rect", ext=(1905000, 952500)))
    second = _drawing("wps", _wsp("ellipse", ext=(952500, 1905000)), ext=(952500, 1905000))

    chunk = _drawn_chunk(_paragraph_with(first, second))

    assert len(chunk.images) == 2
    wide, tall = chunk.images
    assert wide.width and wide.height and tall.width and tall.height
    assert wide.width > wide.height
    assert tall.height > tall.width
    assert wide.docx_shapes is not None and "rect" in wide.docx_shapes
    assert tall.docx_shapes is not None and "ellipse" in tall.docx_shapes


def test_a_lines_only_figure_beside_a_shape_figure_leaves_one_picture() -> None:
    lines = _drawing("wps", _wsp("line"))
    shape = _drawing("wps", _wsp("rect"))

    chunk = _drawn_chunk(_paragraph_with(lines, shape))

    (image,) = chunk.images
    assert image.docx_shapes is not None and "rect" in image.docx_shapes


def test_an_unsupported_figure_drops_the_whole_paragraphs_pictures() -> None:
    custom = _wsp(geometry="<a:custGeom><a:pathLst/></a:custGeom>")
    chunk = _drawn_chunk(_paragraph_with(_drawing("wps", _wsp("rect")), _drawing("wps", custom)))

    assert (chunk.kind, chunk.images, chunk.missing_images) == ("drawing", (), ())


def test_vml_content_beside_a_drawing_keeps_todays_chunk() -> None:
    chunk = _drawn_chunk(
        _paragraph_with(_drawing("wps", _wsp("rect")), "<w:pict><v:rect/></w:pict>")
    )

    assert (chunk.kind, chunk.images, chunk.missing_images) == ("drawing", (), ())


def test_a_paragraph_with_text_and_a_shape_is_unchanged() -> None:
    chunk = _drawn_chunk(_paragraph_with(_drawing("wps", _wsp("rect")), text="Rita figuren"))

    assert (chunk.kind, chunk.text, chunk.images, chunk.missing_images) == (
        "paragraph",
        "Rita figuren",
        (),
        (),
    )


# --- a figure tied to another part of the file (TASK-SKRIPT-39-04-22 D12) ----------------

_NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_XPATH_NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "wps": _URI["wps"],
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": _NS_R,
}


def _shape_with_hyperlink(*, linked: bool) -> bytes:
    """A paragraph holding one rectangle; when ``linked`` its ``wps:cNvPr`` carries an
    ``a:hlinkClick`` that cites a real hyperlink relationship of the document part."""
    document = Document()
    document.add_paragraph("Fråga 1")
    paragraph = document.add_paragraph()
    properties = ""
    if linked:
        relationship_id = document.part.relate_to(
            "https://example.org/", RELATIONSHIP_TYPE.HYPERLINK, is_external=True
        )
        properties = f'<a:hlinkClick xmlns:r="{_NS_R}" r:id="{relationship_id}"/>'
    run = f"<w:r {_NS}>{_drawing('wps', _wsp('rect', properties=properties))}</w:r>"
    paragraph._p.append(parse_xml(run))
    document.add_paragraph("Fråga 2")
    return _bytes(document)


def test_a_shape_that_cites_a_relationship_gets_a_picture_and_no_shapes() -> None:
    chunk = _drawn_chunk(_shape_with_hyperlink(linked=True))

    (image,) = chunk.images
    assert (chunk.kind, chunk.text, chunk.missing_images) == ("drawing", "", ())
    assert image.payload.startswith(_PNG_SIGNATURE)
    assert image.docx_shapes is None


def _figure_paragraph(chunk: DocxChunk) -> NativeParagraph:
    (image,) = chunk.images
    return NativeParagraph(
        heading_level=None,
        segments=(
            NativeTextSegment(text="Se figuren", marks=()),
            NativeAssetSegment(asset_id=image.asset_id),
        ),
    )


def _native_document(chunk: DocxChunk) -> tuple[NativeExamDocument, dict[str, bytes]]:
    """The asset manifest of the chunk's figure and a one-question document that cites it."""
    item = make_free_text_item().model_copy(update={"body": (_figure_paragraph(chunk),)})
    assets, payloads = asset_manifest((item,), (chunk,))
    document = NativeExamDocument(
        document_id=uuid4(),
        revision=1,
        title="Bildprov",
        items=(item,),
        assets=assets,
        origin=NativeExamDocumentOrigin(kind="created"),
    )
    return document, dict(payloads)


def test_the_asset_of_a_shape_that_cites_a_relationship_carries_no_shapes() -> None:
    document, payloads = _native_document(_drawn_chunk(_shape_with_hyperlink(linked=True)))

    (asset,) = document.assets
    assert asset.docx_shapes is None
    assert payloads[asset.asset_id].startswith(_PNG_SIGNATURE)


def _word_export(*, linked: bool) -> tuple[DocxChunk, bytes]:
    """Word import to Word export: the real chunker, asset manifest and Word writer."""
    chunk = _drawn_chunk(_shape_with_hyperlink(linked=linked))
    document, payloads = _native_document(chunk)
    writer = ExamWorkspaceDocxWriter()
    content = writer.build_docx_bytes(document, assets_by_id=payloads)
    writer.validate_docx_bytes(content, expected_item_count=1)
    return chunk, content


def _export_root(content: bytes) -> etree._Element:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        return etree.fromstring(archive.read("word/document.xml"))


def test_the_word_export_writes_a_shape_that_cites_a_relationship_as_its_picture() -> None:
    chunk, content = _word_export(linked=True)
    root = _export_root(content)

    assert root.xpath("//wps:wsp", namespaces=_XPATH_NS) == []
    (embed,) = root.xpath("//a:blip/@r:embed", namespaces=_XPATH_NS)
    assert Document(io.BytesIO(content)).part.related_parts[str(embed)].blob == (
        chunk.images[0].payload
    )


def test_the_word_export_writes_a_shape_without_a_relationship_as_shapes() -> None:
    """The control: the same figure without the hyperlink is written back as its shape."""
    _, content = _word_export(linked=False)
    root = _export_root(content)

    assert len(root.xpath("//wps:wsp", namespaces=_XPATH_NS)) == 1
    assert root.xpath("//a:blip", namespaces=_XPATH_NS) == []
