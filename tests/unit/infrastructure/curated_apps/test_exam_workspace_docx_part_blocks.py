"""DOCX export: a part with no questions prints its instruction at its own place (G5, R2).

The childless part's instruction paragraph sits between the two question headings, an empty
one adds no paragraph, and a picture inside it is embedded like any shared instruction.
"""

import hashlib
import io
from uuid import uuid4

from docx import Document
from PIL import Image

from skriptoteket.domain.curated_apps.exam_workspace.native_exam_body import (
    NativeAssetSegment,
    NativeParagraph,
    NativeTextSegment,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_document import (
    NativeExamAsset,
    NativeExamDocument,
    NativeExamDocumentOrigin,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_item import NativeExamItem
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_parent import (
    NativeExamParent,
)
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub import (
    exam_workspace_docx_writer,
)
from tests.unit.domain.curated_apps.exam_workspace.native_exam_test_builders import (
    make_free_text_item,
    make_paragraph,
)

ExamWorkspaceDocxWriter = exam_workspace_docx_writer.ExamWorkspaceDocxWriter

_PART_TEXT = "Delinstruktion: arbeta i par."


def _question(sequence: int) -> NativeExamItem:
    return make_free_text_item(sequence=sequence, item_id=f"item_{sequence:03d}")


def _part(
    *,
    place: int,
    instruction: tuple[NativeParagraph, ...] = (),
    children: tuple[str, ...] = (),
) -> NativeExamParent:
    return NativeExamParent(
        parent_id="parent_001",
        place=place,
        instruction=instruction,
        item_ids=children,
    )


def _document(
    items: list[NativeExamItem],
    parents: tuple[NativeExamParent, ...] = (),
    *,
    assets: tuple[NativeExamAsset, ...] = (),
) -> NativeExamDocument:
    return NativeExamDocument(
        document_id=uuid4(),
        revision=1,
        title="Delprov",
        items=tuple(items),
        parents=parents,
        assets=assets,
        origin=NativeExamDocumentOrigin(kind="created"),
    )


def _paragraph_texts(content: bytes) -> list[str]:
    return [paragraph.text for paragraph in Document(io.BytesIO(content)).paragraphs]


def _png() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGBA", (12, 6), (200, 20, 20, 255)).save(buffer, format="PNG")
    return buffer.getvalue()


def test_a_childless_part_instruction_sits_between_the_question_headings() -> None:
    document = _document(
        [_question(1), _question(2)],
        (_part(place=1, instruction=(make_paragraph(_PART_TEXT),)),),
    )

    content = ExamWorkspaceDocxWriter().build_docx_bytes(document)

    ExamWorkspaceDocxWriter().validate_docx_bytes(content, expected_item_count=2)
    texts = _paragraph_texts(content)
    assert texts.index("Fråga 1") < texts.index(_PART_TEXT) < texts.index("Fråga 2")


def test_an_empty_childless_part_adds_no_paragraph() -> None:
    items = [_question(1), _question(2)]
    without_part = ExamWorkspaceDocxWriter().build_docx_bytes(_document(items))
    with_empty_part = ExamWorkspaceDocxWriter().build_docx_bytes(
        _document(items, (_part(place=1),))
    )

    ExamWorkspaceDocxWriter().validate_docx_bytes(with_empty_part, expected_item_count=2)
    assert _paragraph_texts(with_empty_part) == _paragraph_texts(without_part)


def test_a_picture_in_a_childless_part_instruction_is_embedded() -> None:
    payload = _png()
    digest = hashlib.sha256(payload).hexdigest()
    asset_id = f"asset_{digest[:16]}"
    asset = NativeExamAsset(
        asset_id=asset_id,
        media_type="image/png",
        byte_length=len(payload),
        sha256=digest,
    )
    instruction = (
        NativeParagraph(
            heading_level=None,
            segments=(
                NativeTextSegment(text="Se figuren.", marks=()),
                NativeAssetSegment(asset_id=asset_id),
            ),
        ),
    )
    document = _document(
        [_question(1), _question(2)],
        (_part(place=1, instruction=instruction),),
        assets=(asset,),
    )

    content = ExamWorkspaceDocxWriter().build_docx_bytes(document, assets_by_id={asset_id: payload})

    ExamWorkspaceDocxWriter().validate_docx_bytes(content, expected_item_count=2)
    parsed = Document(io.BytesIO(content))
    assert len(parsed.inline_shapes) == 1
