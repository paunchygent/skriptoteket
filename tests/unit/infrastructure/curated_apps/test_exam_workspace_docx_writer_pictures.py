"""Shared-picture round-trip tests for the exam workspace DOCX writer (W1-W2).

Instruction and item-body pictures are read back from the produced bytes, checking each
shared instruction keeps its text runs and payload in place, the body picture is
embedded, an oversized picture is clamped and a missing payload fails closed.
"""

import hashlib
import io
import zipfile
from uuid import uuid4

import pytest
from docx import Document
from docx.shared import Mm

from skriptoteket.domain.curated_apps.exam_conversion.exam_print_house_template import (
    IMAGE_SPACING,
    TEXT_WIDTH_MM,
)
from skriptoteket.domain.curated_apps.exam_workspace.exam_language_labels import (
    SWEDISH_EXAM_LABELS,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_document import (
    NativeExamDocument,
    NativeExamDocumentOrigin,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_item import (
    NativeExamItem,
)
from skriptoteket.domain.errors import DomainError, ErrorCode
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub import (
    exam_workspace_docx_writer,
)
from tests.unit.domain.curated_apps.exam_workspace.native_exam_test_builders import (
    make_free_text_item,
)
from tests.unit.infrastructure.curated_apps.exam_workspace_docx_picture_support import (
    SharedPictureDocument,
    build_shared_picture_document,
    shared_asset,
)
from tests.unit.infrastructure.curated_apps.exam_workspace_docx_writer_test_support import (
    document_text,
    drawing_payload,
    heading_texts,
    only_paragraph,
    png,
    run_shape,
)

ExamWorkspaceDocxWriter = exam_workspace_docx_writer.ExamWorkspaceDocxWriter

_SHARED_GLOBAL_BEFORE = "Betrakta: "
_SHARED_GLOBAL_AFTER = " diagrammet."
_SHARED_PARENT_BEFORE = "Se figuren: "
_SHARED_PARENT_AFTER = " och svara."


@pytest.fixture
def shared_picture_document() -> SharedPictureDocument:
    return build_shared_picture_document()


class TestPictures:
    def test_two_shared_instructions_round_trip_in_place_with_payloads(
        self, shared_picture_document: SharedPictureDocument
    ) -> None:
        fixture = shared_picture_document
        writer = ExamWorkspaceDocxWriter()
        content = writer.build_docx_bytes(fixture.document, assets_by_id=fixture.assets_by_id)
        writer.validate_docx_bytes(content, expected_item_count=len(fixture.document.items))

        parsed = Document(io.BytesIO(content))
        paragraphs = list(parsed.paragraphs)
        texts = [paragraph.text for paragraph in paragraphs]

        global_paragraph = only_paragraph(
            paragraphs, f"{_SHARED_GLOBAL_BEFORE}{_SHARED_GLOBAL_AFTER}"
        )
        assert run_shape(global_paragraph) == [
            _SHARED_GLOBAL_BEFORE,
            "image",
            _SHARED_GLOBAL_AFTER,
        ]
        assert (
            drawing_payload(parsed, global_paragraph)
            == fixture.assets_by_id[fixture.global_asset_id]
        )
        assert texts.count(global_paragraph.text) == 1

        parent_paragraph = only_paragraph(
            paragraphs, f"{_SHARED_PARENT_BEFORE}{_SHARED_PARENT_AFTER}"
        )
        assert run_shape(parent_paragraph) == [
            _SHARED_PARENT_BEFORE,
            "image",
            _SHARED_PARENT_AFTER,
        ]
        assert (
            drawing_payload(parsed, parent_paragraph)
            == fixture.assets_by_id[fixture.parent_asset_id]
        )
        assert texts.count(parent_paragraph.text) == 1
        first_heading_at = texts.index(heading_texts(parsed, "Heading 3")[0])
        assert texts.index(parent_paragraph.text) == first_heading_at - 1

        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            media_payloads = [
                archive.read(name) for name in archive.namelist() if name.startswith("word/media/")
            ]
        expected_digests = sorted(
            hashlib.sha256(payload).hexdigest() for payload in fixture.assets_by_id.values()
        )
        assert sorted(hashlib.sha256(payload).hexdigest() for payload in media_payloads) == (
            expected_digests
        )

        assert SWEDISH_EXAM_LABELS.answer_key not in heading_texts(parsed, "Heading 2")
        assert "Godtagbar förklaring." not in document_text(parsed)

    def test_item_body_picture_is_embedded(
        self, shared_picture_document: SharedPictureDocument
    ) -> None:
        fixture = shared_picture_document
        content = ExamWorkspaceDocxWriter().build_docx_bytes(
            fixture.document, assets_by_id=fixture.assets_by_id
        )

        parsed = Document(io.BytesIO(content))
        # The global instruction, item_001's body and the parent instruction each embed one.
        assert len(parsed.inline_shapes) == 3
        body_paragraph = only_paragraph(list(parsed.paragraphs), "Fråga 1. (1 p)")
        assert run_shape(body_paragraph) == ["Fråga 1.", "image", " (1 p)"]
        assert (
            drawing_payload(parsed, body_paragraph) == fixture.assets_by_id[fixture.global_asset_id]
        )

    def test_picture_wider_than_the_text_column_is_clamped(self) -> None:
        payload = png((10, 90, 10, 255), (4000, 40))
        asset_id, asset = shared_asset(payload)
        item = make_free_text_item()
        item = NativeExamItem.model_validate(
            {
                **item.model_dump(),
                "body": [
                    {
                        "heading_level": None,
                        "segments": [
                            {"kind": "text", "text": "Se bilden", "marks": []},
                            {"kind": "asset", "asset_id": asset_id},
                        ],
                    }
                ],
            }
        )
        document = NativeExamDocument(
            document_id=uuid4(),
            revision=1,
            title="Bildprov",
            items=(item,),
            assets=(asset,),
            origin=NativeExamDocumentOrigin(kind="created"),
        )

        content = ExamWorkspaceDocxWriter().build_docx_bytes(
            document, assets_by_id={asset_id: payload}
        )
        parsed = Document(io.BytesIO(content))
        assert len(parsed.inline_shapes) == 1
        assert parsed.inline_shapes[0].width == Mm(TEXT_WIDTH_MM)

    def test_a_body_picture_paragraph_takes_the_template_image_spacing(
        self, shared_picture_document: SharedPictureDocument
    ) -> None:
        fixture = shared_picture_document
        content = ExamWorkspaceDocxWriter().build_docx_bytes(
            fixture.document, assets_by_id=fixture.assets_by_id
        )
        parsed = Document(io.BytesIO(content))
        body_paragraph = only_paragraph(list(parsed.paragraphs), "Fråga 1. (1 p)")
        spacing = body_paragraph.paragraph_format

        assert spacing.space_before.mm == pytest.approx(IMAGE_SPACING.space_before_mm, abs=0.02)
        assert spacing.space_after.mm == pytest.approx(IMAGE_SPACING.space_after_mm, abs=0.02)

    def test_picture_without_payload_fails_closed(
        self, shared_picture_document: SharedPictureDocument
    ) -> None:
        fixture = shared_picture_document
        with pytest.raises(DomainError) as exc_info:
            ExamWorkspaceDocxWriter().build_docx_bytes(fixture.document)

        assert exc_info.value.code is ErrorCode.VALIDATION_ERROR
        assert exc_info.value.details["asset_ids"] == [fixture.global_asset_id]
