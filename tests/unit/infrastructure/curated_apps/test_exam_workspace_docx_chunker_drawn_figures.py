"""Drawn figures: Word shapes the DOCX chunking adapter turns into pictures."""

import io
import zipfile
from uuid import uuid4

import pytest
from docx import Document
from docx.document import Document as DocxDocument
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import parse_xml
from lxml import etree

from skriptoteket.domain.curated_apps.exam_workspace.docx_build_assets import asset_manifest
from skriptoteket.domain.curated_apps.exam_workspace.docx_chunks import DocxChunk
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_body import (
    NativeAssetSegment,
    NativeParagraph,
    NativeTextSegment,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_document import (
    NativeExamDocument,
    NativeExamDocumentOrigin,
)
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_docx_chunker import (  # noqa: E501
    PythonDocxExamChunker,
)
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_docx_writer import (  # noqa: E501
    ExamWorkspaceDocxWriter,
)
from tests.unit.domain.curated_apps.exam_workspace.native_exam_test_builders import (
    make_free_text_item,
)


def _bytes(document: DocxDocument) -> bytes:
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


# --- drawn figures: Word shapes become pictures (TASK-SKRIPT-39-04-22 D1-D4) ---------------

_PNG_SIGNATURE = b"\x89PNG\r\n\x1a\n"
_NS = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:v="urn:schemas-microsoft-com:vml" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:wpg="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup" '
    'xmlns:wpc="http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas"'
)
_URI = {
    "wps": "http://schemas.microsoft.com/office/word/2010/wordprocessingShape",
    "wpg": "http://schemas.microsoft.com/office/word/2010/wordprocessingGroup",
    "wpc": "http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas",
}
_LINE = '<a:ln w="19050"><a:solidFill><a:srgbClr val="000000"/></a:solidFill></a:ln>'


def _wsp(
    prst: str = "rect",
    *,
    off: tuple[int, int] = (0, 0),
    ext: tuple[int, int] = (1905000, 952500),
    label: str | None = None,
    geometry: str | None = None,
    line: str = _LINE,
    properties: str = "",
) -> str:
    geometry = geometry or f'<a:prstGeom prst="{prst}"><a:avLst/></a:prstGeom>'
    text = ""
    if label is not None:
        text = (
            "<wps:txbx><w:txbxContent><w:p><w:r>"
            f"<w:t>{label}</w:t></w:r></w:p></w:txbxContent></wps:txbx>"
        )
    return (
        f'<wps:wsp><wps:cNvPr id="2" name="s">{properties}</wps:cNvPr><wps:cNvSpPr/><wps:spPr>'
        f'<a:xfrm><a:off x="{off[0]}" y="{off[1]}"/><a:ext cx="{ext[0]}" cy="{ext[1]}"/></a:xfrm>'
        f'{geometry}<a:solidFill><a:srgbClr val="FFCC00"/></a:solidFill>{line}</wps:spPr>'
        f"{text}<wps:bodyPr/></wps:wsp>"
    )


def _drawing(kind: str, shapes: str, *, ext: tuple[int, int] = (1905000, 952500)) -> str:
    graphic = shapes
    if kind == "wpg":
        graphic = (
            '<wpg:wgp><wpg:cNvPr id="3" name="g"/><wpg:cNvGrpSpPr/><wpg:grpSpPr>'
            f'<a:xfrm><a:off x="0" y="0"/><a:ext cx="{ext[0]}" cy="{ext[1]}"/>'
            f'<a:chOff x="0" y="0"/><a:chExt cx="{ext[0]}" cy="{ext[1]}"/></a:xfrm>'
            f"</wpg:grpSpPr>{shapes}</wpg:wgp>"
        )
    elif kind == "wpc":
        graphic = f"<wpc:wpc><wpc:bg><a:noFill/></wpc:bg>{shapes}</wpc:wpc>"
    return (
        "<w:drawing><wp:inline>"
        f'<wp:extent cx="{ext[0]}" cy="{ext[1]}"/><wp:docPr id="1" name="d"/>'
        f'<a:graphic><a:graphicData uri="{_URI[kind]}">{graphic}</a:graphicData></a:graphic>'
        "</wp:inline></w:drawing>"
    )


def _group(*labels: str) -> str:
    return _drawing(
        "wpg",
        "".join(
            _wsp("ellipse", off=(index * 950000, 0), ext=(900000, 900000), label=label)
            for index, label in enumerate(labels)
        ),
    )


def _paragraph_with(*run_contents: str, text: str = "") -> bytes:
    document = Document()
    document.add_paragraph("Fråga 1")
    paragraph = document.add_paragraph(text)
    for content in run_contents:
        paragraph._p.append(parse_xml(f"<w:r {_NS}>{content}</w:r>"))
    document.add_paragraph("Fråga 2")
    return _bytes(document)


def _drawn_chunk(content: bytes) -> DocxChunk:
    (chunk,) = [
        c for c in PythonDocxExamChunker().chunk(content=content).chunks if c.anchor == "body:0002"
    ]
    return chunk


def test_a_group_of_shapes_with_labels_becomes_one_picture() -> None:
    chunk = _drawn_chunk(_paragraph_with(_group("Population", "Urval")))

    assert (chunk.kind, chunk.text, chunk.missing_images) == ("drawing", "", ())
    (image,) = chunk.images
    assert image.payload.startswith(_PNG_SIGNATURE)
    assert image.media_type == "image/png"
    assert (image.offset, image.marked_offset) == (0, 0)
    assert image.width and image.height
    assert image.docx_shapes is not None
    assert "wpg:wgp" in image.docx_shapes
    assert "Population" in image.docx_shapes


@pytest.mark.parametrize("kind", ["wps", "wpc"])
def test_a_single_shape_and_a_canvas_become_one_picture_each(kind: str) -> None:
    shapes = _wsp("rect", label="Ruta")
    chunk = _drawn_chunk(_paragraph_with(_drawing(kind, shapes)))

    (image,) = chunk.images
    assert (chunk.kind, chunk.text, chunk.missing_images) == ("drawing", "", ())
    assert image.payload.startswith(_PNG_SIGNATURE)
    assert image.docx_shapes is not None
    assert "wps:wsp" in image.docx_shapes


def test_the_choice_drawing_is_drawn_and_its_vml_fallback_ignored() -> None:
    content = (
        '<mc:AlternateContent><mc:Choice Requires="wpg">'
        f"{_group('Population')}"
        "</mc:Choice><mc:Fallback><w:pict><v:rect/></w:pict></mc:Fallback></mc:AlternateContent>"
    )

    chunk = _drawn_chunk(_paragraph_with(content))

    (image,) = chunk.images
    assert image.payload.startswith(_PNG_SIGNATURE)
    assert chunk.missing_images == ()


def test_a_group_with_a_custom_geometry_shape_keeps_the_drawing_chunk_without_a_picture() -> None:
    custom = _wsp(
        geometry=(
            "<a:custGeom><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/>"
            '<a:rect l="0" t="0" r="r" b="b"/><a:pathLst/></a:custGeom>'
        )
    )
    chunk = _drawn_chunk(_paragraph_with(_drawing("wpg", _wsp("rect") + custom)))

    assert (chunk.kind, chunk.text, chunk.images, chunk.missing_images) == ("drawing", "", (), ())


def test_a_drawing_of_plain_lines_gets_no_picture() -> None:
    lines = _wsp("line") + _wsp("straightConnector1", off=(0, 100000))
    chunk = _drawn_chunk(_paragraph_with(_drawing("wpg", lines)))

    assert (chunk.kind, chunk.images, chunk.missing_images) == ("drawing", (), ())


def test_two_figures_in_one_paragraph_give_two_pictures_in_order() -> None:
    first = _drawing("wps", _wsp("rect", ext=(1905000, 952500)))
    second = _drawing("wps", _wsp("ellipse", ext=(952500, 1905000)), ext=(952500, 1905000))

    chunk = _drawn_chunk(_paragraph_with(first, second))

    assert len(chunk.images) == 2
    wide, tall = chunk.images
    assert wide.width and wide.height and tall.width and tall.height
    assert wide.width > wide.height
    assert tall.height > tall.width
    assert wide.docx_shapes is not None and "rect" in wide.docx_shapes
    assert tall.docx_shapes is not None and "ellipse" in tall.docx_shapes


def test_a_lines_only_figure_beside_a_shape_figure_leaves_one_picture() -> None:
    lines = _drawing("wps", _wsp("line"))
    shape = _drawing("wps", _wsp("rect"))

    chunk = _drawn_chunk(_paragraph_with(lines, shape))

    (image,) = chunk.images
    assert image.docx_shapes is not None and "rect" in image.docx_shapes


def test_an_unsupported_figure_drops_the_whole_paragraphs_pictures() -> None:
    custom = _wsp(geometry="<a:custGeom><a:pathLst/></a:custGeom>")
    chunk = _drawn_chunk(_paragraph_with(_drawing("wps", _wsp("rect")), _drawing("wps", custom)))

    assert (chunk.kind, chunk.images, chunk.missing_images) == ("drawing", (), ())


def test_vml_content_beside_a_drawing_keeps_todays_chunk() -> None:
    chunk = _drawn_chunk(
        _paragraph_with(_drawing("wps", _wsp("rect")), "<w:pict><v:rect/></w:pict>")
    )

    assert (chunk.kind, chunk.images, chunk.missing_images) == ("drawing", (), ())


def test_a_paragraph_with_text_and_a_shape_is_unchanged() -> None:
    chunk = _drawn_chunk(_paragraph_with(_drawing("wps", _wsp("rect")), text="Rita figuren"))

    assert (chunk.kind, chunk.text, chunk.images, chunk.missing_images) == (
        "paragraph",
        "Rita figuren",
        (),
        (),
    )


# --- a figure tied to another part of the file (TASK-SKRIPT-39-04-22 D12) ----------------

_NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_XPATH_NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "wps": _URI["wps"],
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": _NS_R,
}


def _shape_with_hyperlink(*, linked: bool) -> bytes:
    """A paragraph holding one rectangle; when ``linked`` its ``wps:cNvPr`` carries an
    ``a:hlinkClick`` that cites a real hyperlink relationship of the document part."""
    document = Document()
    document.add_paragraph("Fråga 1")
    paragraph = document.add_paragraph()
    properties = ""
    if linked:
        relationship_id = document.part.relate_to(
            "https://example.org/", RELATIONSHIP_TYPE.HYPERLINK, is_external=True
        )
        properties = f'<a:hlinkClick xmlns:r="{_NS_R}" r:id="{relationship_id}"/>'
    run = f"<w:r {_NS}>{_drawing('wps', _wsp('rect', properties=properties))}</w:r>"
    paragraph._p.append(parse_xml(run))
    document.add_paragraph("Fråga 2")
    return _bytes(document)


def test_a_shape_that_cites_a_relationship_gets_a_picture_and_no_shapes() -> None:
    chunk = _drawn_chunk(_shape_with_hyperlink(linked=True))

    (image,) = chunk.images
    assert (chunk.kind, chunk.text, chunk.missing_images) == ("drawing", "", ())
    assert image.payload.startswith(_PNG_SIGNATURE)
    assert image.docx_shapes is None


def _figure_paragraph(chunk: DocxChunk) -> NativeParagraph:
    (image,) = chunk.images
    return NativeParagraph(
        heading_level=None,
        segments=(
            NativeTextSegment(text="Se figuren", marks=()),
            NativeAssetSegment(asset_id=image.asset_id),
        ),
    )


def _native_document(chunk: DocxChunk) -> tuple[NativeExamDocument, dict[str, bytes]]:
    """The asset manifest of the chunk's figure and a one-question document that cites it."""
    item = make_free_text_item().model_copy(update={"body": (_figure_paragraph(chunk),)})
    assets, payloads = asset_manifest((item,), (chunk,))
    document = NativeExamDocument(
        document_id=uuid4(),
        revision=1,
        title="Bildprov",
        items=(item,),
        assets=assets,
        origin=NativeExamDocumentOrigin(kind="created"),
    )
    return document, dict(payloads)


def test_the_asset_of_a_shape_that_cites_a_relationship_carries_no_shapes() -> None:
    document, payloads = _native_document(_drawn_chunk(_shape_with_hyperlink(linked=True)))

    (asset,) = document.assets
    assert asset.docx_shapes is None
    assert payloads[asset.asset_id].startswith(_PNG_SIGNATURE)


def _word_export(*, linked: bool) -> tuple[DocxChunk, bytes]:
    """Word import to Word export: the real chunker, asset manifest and Word writer."""
    chunk = _drawn_chunk(_shape_with_hyperlink(linked=linked))
    document, payloads = _native_document(chunk)
    writer = ExamWorkspaceDocxWriter()
    content = writer.build_docx_bytes(document, assets_by_id=payloads)
    writer.validate_docx_bytes(content, expected_item_count=1)
    return chunk, content


def _export_root(content: bytes) -> etree._Element:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        return etree.fromstring(archive.read("word/document.xml"))


def test_the_word_export_writes_a_shape_that_cites_a_relationship_as_its_picture() -> None:
    chunk, content = _word_export(linked=True)
    root = _export_root(content)

    assert root.xpath("//wps:wsp", namespaces=_XPATH_NS) == []
    (embed,) = root.xpath("//a:blip/@r:embed", namespaces=_XPATH_NS)
    assert Document(io.BytesIO(content)).part.related_parts[str(embed)].blob == (
        chunk.images[0].payload
    )


def test_the_word_export_writes_a_shape_without_a_relationship_as_shapes() -> None:
    """The control: the same figure without the hyperlink is written back as its shape."""
    _, content = _word_export(linked=False)
    root = _export_root(content)

    assert len(root.xpath("//wps:wsp", namespaces=_XPATH_NS)) == 1
    assert root.xpath("//a:blip", namespaces=_XPATH_NS) == []
