"""Fritext criteria survive storage and print with points in DOCX and PDF."""

import io
from uuid import uuid4

import pytest
from docx import Document

from skriptoteket.domain.curated_apps.exam_conversion.digiexam_examnet_pdf_contracts import (
    DigiExamExamNetPdfStatus,
)
from skriptoteket.domain.curated_apps.exam_workspace.container import (
    ExamWorkspaceContainerContent,
)
from skriptoteket.domain.curated_apps.exam_workspace.docx_chunks import DocxChunk
from skriptoteket.domain.curated_apps.exam_workspace.docx_own_instructions import wording_left_out
from skriptoteket.domain.curated_apps.exam_workspace.examnet_pdf_export import (
    native_exam_to_pdf_document,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_body import (
    NativeParagraph,
    NativeTextSegment,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_criteria import (
    NativeExamCriterion,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_document import (
    NativeExamDocument,
    NativeExamDocumentOrigin,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_item import (
    NativeExamItem,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_values import (
    NativeAnswerKey,
    NativeItemReview,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_vocabulary import (
    NativeAnswerKeyOrigin,
    NativeExamItemKind,
    NativeItemReviewState,
    NativeParseOrigin,
)
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub import (
    exam_workspace_docx_writer,
)
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_container import (
    ExamWorkspaceContainerCodec,
)

_SWEDISH_BODY = "Förklara vad som är en bisats och hur den används för att binda ihop en text."
_ENGLISH_BODY = "Explain what the clause is and how it is used to connect the text."
_CRITERIA = (
    NativeExamCriterion(text="Definitionen", points=2),
    NativeExamCriterion(text="<exemplet> & motiveringen\nmed detaljer", points=1),
)


def _document(body: str, *, criteria: tuple[NativeExamCriterion, ...]) -> NativeExamDocument:
    return NativeExamDocument(
        document_id=uuid4(),
        revision=1,
        title="Prov",
        items=(
            NativeExamItem(
                item_id="item_001",
                sequence=1,
                kind=NativeExamItemKind.FREE_TEXT,
                body=(NativeParagraph(segments=(NativeTextSegment(text=body),)),),
                points=3,
                answer_key=NativeAnswerKey(origin=NativeAnswerKeyOrigin.NOT_APPLICABLE),
                criteria=criteria,
                review=NativeItemReview(
                    state=NativeItemReviewState.REVIEW_COMPLETE,
                    parse_origin=NativeParseOrigin.TEACHER_CREATED,
                ),
            ),
        ),
        origin=NativeExamDocumentOrigin(kind="created"),
    )


def _docx_lines(document: NativeExamDocument) -> list[str]:
    content = exam_workspace_docx_writer.ExamWorkspaceDocxWriter().build_docx_bytes(document)
    return [paragraph.text for paragraph in Document(io.BytesIO(content)).paragraphs]


def test_own_instructions_count_criterion_text_as_preserved_wording() -> None:
    item = _document(_SWEDISH_BODY, criteria=_CRITERIA).items[0]
    preserved = DocxChunk(
        anchor="body:0002",
        kind="paragraph",
        text="Definitionen",
        marked_text="Definitionen",
    )
    omitted = DocxChunk(
        anchor="body:0003",
        kind="paragraph",
        text="Ett separat krav som saknas",
        marked_text="Ett separat krav som saknas",
    )
    assert wording_left_out(item, (preserved, omitted)) == (omitted,)


def test_criteria_survive_the_container_round_trip() -> None:
    document = _document(_SWEDISH_BODY, criteria=_CRITERIA)
    codec = ExamWorkspaceContainerCodec()
    built = codec.build(content=ExamWorkspaceContainerContent(document=document))
    assert codec.parse(content=built).document.items[0].criteria == _CRITERIA


@pytest.mark.parametrize(
    ("body", "label"),
    [(_SWEDISH_BODY, "Rättningsmall"), (_ENGLISH_BODY, "Marking guide")],
)
def test_docx_answer_key_prints_each_criterion_with_points(body: str, label: str) -> None:
    lines = _docx_lines(_document(body, criteria=_CRITERIA))
    assert f"{label}:" in lines
    assert "• Definitionen (2 p)" in lines
    assert "• <exemplet> & motiveringen\nmed detaljer (1 p)" in lines
    other = "Marking guide" if label == "Rättningsmall" else "Rättningsmall"
    assert not any(line.startswith(other) for line in lines)


@pytest.mark.parametrize("body", [_SWEDISH_BODY, _ENGLISH_BODY])
def test_docx_without_criteria_has_no_answer_key_section(body: str) -> None:
    lines = _docx_lines(_document(body, criteria=()))
    assert not any(line in ("Facit", "Answer key") for line in lines)
    assert not any(line.startswith(("Rättningsmall", "Marking guide")) for line in lines)


@pytest.mark.parametrize(
    ("body", "label"),
    [(_SWEDISH_BODY, "Rättningsmall"), (_ENGLISH_BODY, "Marking guide")],
)
def test_pdf_prints_each_escaped_criterion_with_points(body: str, label: str) -> None:
    result = native_exam_to_pdf_document(_document(body, criteria=_CRITERIA))
    assert result.status is DigiExamExamNetPdfStatus.SUCCESS
    assert (
        f'<div class="answer-key"><p>{label}:</p><ul>'
        "<li>Definitionen (2 p)</li>"
        "<li>&lt;exemplet&gt; &amp; motiveringen<br>med detaljer (1 p)</li>"
        "</ul></div>"
    ) in result.html


@pytest.mark.parametrize("body", [_SWEDISH_BODY, _ENGLISH_BODY])
def test_pdf_without_criteria_prints_no_answer_key(body: str) -> None:
    result = native_exam_to_pdf_document(_document(body, criteria=()))
    assert result.status is DigiExamExamNetPdfStatus.SUCCESS
    assert 'class="answer-key"' not in result.html
