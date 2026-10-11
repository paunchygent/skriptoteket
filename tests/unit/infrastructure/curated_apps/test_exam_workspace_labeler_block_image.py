"""Challenge the D14 block-image renderer with real python-docx documents built in memory."""

from __future__ import annotations

import io

import pytest
from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.oxml.text.paragraph import CT_P
from PIL import Image

from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_labeler_block_image import (  # noqa: E501
    DocxBlockImageRenderer,
    LabelerSourceImageUnavailable,
)

pytestmark = pytest.mark.unit


def _png(content: bytes) -> Image.Image:
    image = Image.open(io.BytesIO(content))
    image.load()
    assert image.format == "PNG"
    return image


def _ink(image: Image.Image) -> int:
    histogram = image.convert("L").histogram()
    return sum(histogram[:200])


def _numbered(document: DocxDocument, num_id: int, formats: list[tuple[str, str]]) -> None:
    levels = "".join(
        f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/>'
        f'<w:lvlText w:val="{text}"/></w:lvl>'
        for i, (fmt, text) in enumerate(formats)
    )
    numbering = document.part.numbering_part.element
    abstract = parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{900 + num_id}">{levels}</w:abstractNum>'
    )
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        first_num.addprevious(abstract)
    numbering.append(
        parse_xml(
            f'<w:num {nsdecls("w")} w:numId="{num_id}">'
            f'<w:abstractNumId w:val="{900 + num_id}"/></w:num>'
        )
    )


def _item(document: DocxDocument, text: str, num_id: int, level: int = 0) -> CT_P:
    paragraph = document.add_paragraph(text)
    paragraph._p.get_or_add_pPr().append(
        parse_xml(
            f'<w:numPr {nsdecls("w")}><w:ilvl w:val="{level}"/>'
            f'<w:numId w:val="{num_id}"/></w:numPr>'
        )
    )
    return paragraph._p


def _picture_document(size: tuple[int, int]) -> tuple[DocxDocument, CT_P]:
    stream = io.BytesIO()
    Image.new("RGB", size, (200, 20, 20)).save(stream, format="PNG")
    stream.seek(0)
    document = Document()
    document.add_picture(stream)
    return document, document.paragraphs[-1]._p


def test_horizontally_merged_cell_renders_wide_png() -> None:
    document = Document()
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 2)).text = "Rubrik över tre kolumner"
    for column, text in enumerate(("Fråga 1", "Svar", "Poäng")):
        table.cell(1, column).text = text
    image = _png(DocxBlockImageRenderer(document).render(table._tbl, anchor="body:0001"))
    assert _ink(image) > 0
    assert image.width > image.height


def test_vertical_merge_with_nested_table_renders() -> None:
    document = Document()
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).merge(table.cell(2, 0)).text = "Sammanslagen"
    nested = table.cell(1, 1).add_table(rows=2, cols=2)
    nested.cell(0, 0).text = "Inre"
    renderer = DocxBlockImageRenderer(document)
    image = _png(renderer.render(table._tbl, anchor="body:0002"))
    assert _ink(image) > 0


def test_decimal_list_counts_and_sublevel_restarts() -> None:
    document = Document()
    _numbered(document, 41, [("decimal", "%1."), ("lowerLetter", "%2)")])
    items = [
        _item(document, "Ett", 41),
        _item(document, "Under ett", 41, 1),
        _item(document, "Under ett igen", 41, 1),
        _item(document, "Två", 41),
        _item(document, "Under två", 41, 1),
    ]
    plain = document.add_paragraph("Ingen numrering")._p
    renderer = DocxBlockImageRenderer(document)
    assert [renderer.list_label(p) for p in items] == ["1.", "a)", "b)", "2.", "a)"]
    assert renderer.list_label(plain) is None
    assert _ink(_png(renderer.render(items[3], anchor="body:0004"))) > 0


def test_lower_letter_label_and_numbers_inside_tables() -> None:
    document = Document()
    _numbered(document, 42, [("lowerLetter", "%1)")])
    first = _item(document, "a-fråga", 42)
    table = document.add_table(rows=1, cols=1)
    inner = table.cell(0, 0).paragraphs[0]
    inner._p.get_or_add_pPr().append(
        parse_xml(f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="42"/></w:numPr>')
    )
    after = _item(document, "c-fråga", 42)
    renderer = DocxBlockImageRenderer(document)
    assert [renderer.list_label(p) for p in (first, inner._p, after)] == ["a)", "b)", "c)"]


def test_picture_size_carries_into_render() -> None:
    document, paragraph = _picture_document((200, 100))
    image = _png(DocxBlockImageRenderer(document).render(paragraph, anchor="body:0001"))
    # 200 px at 72 dpi is 200 pt wide; scale 2 gives 400 px plus 16 px margin per side.
    assert abs(image.width - 432) <= 6
    assert abs(image.height - 232) <= 6


def test_emf_picture_is_unavailable() -> None:
    document, paragraph = _picture_document((40, 40))
    for part in document.part.related_parts.values():
        if part.content_type == "image/png":
            part._content_type = "image/x-emf"
    with pytest.raises(LabelerSourceImageUnavailable) as raised:
        DocxBlockImageRenderer(document).render(paragraph, anchor="body:0007")
    assert raised.value.anchor == "body:0007"
    assert "image/x-emf" in raised.value.reason


def test_drawing_without_picture_is_unavailable() -> None:
    document = Document()
    paragraph = document.add_paragraph("Diagram: ")
    chart = "http://schemas.openxmlformats.org/drawingml/2006/chart"
    paragraph.add_run()._r.append(
        parse_xml(
            f"<w:drawing {nsdecls('w', 'wp', 'a')}><wp:inline>"
            '<wp:extent cx="952500" cy="952500"/><wp:docPr id="1" name="Diagram"/>'
            f'<a:graphic><a:graphicData uri="{chart}"/></a:graphic>'
            "</wp:inline></w:drawing>"
        )
    )
    with pytest.raises(LabelerSourceImageUnavailable) as raised:
        DocxBlockImageRenderer(document).render(paragraph._p, anchor="body:0003")
    assert raised.value.anchor == "body:0003"
    assert chart in raised.value.reason


def test_unknown_number_format_is_unavailable() -> None:
    document = Document()
    _numbered(document, 43, [("ordinalText", "%1")])
    item = _item(document, "Första", 43)
    with pytest.raises(LabelerSourceImageUnavailable) as raised:
        DocxBlockImageRenderer(document).render(item, anchor="body:0009")
    assert raised.value.anchor == "body:0009"
    assert raised.value.reason == "numbering format ordinalText"


def test_text_box_paragraph_counts_once() -> None:
    document = Document()
    _numbered(document, 44, [("decimal", "%1.")])
    host = document.add_paragraph()
    box = (
        f"<mc:AlternateContent {nsdecls('w')} "
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
        '<mc:Choice Requires="wps"><w:drawing><w:txbxContent><w:p><w:pPr><w:numPr>'
        '<w:ilvl w:val="0"/><w:numId w:val="44"/></w:numPr></w:pPr>'
        "<w:r><w:t>I rutan</w:t></w:r></w:p></w:txbxContent></w:drawing></mc:Choice>"
        "<mc:Fallback><w:pict><w:txbxContent><w:p><w:pPr><w:numPr>"
        '<w:ilvl w:val="0"/><w:numId w:val="44"/></w:numPr></w:pPr>'
        "<w:r><w:t>I rutan</w:t></w:r></w:p></w:txbxContent></w:pict></mc:Fallback>"
        "</mc:AlternateContent>"
    )
    host.add_run()._r.append(parse_xml(box))
    after = _item(document, "Efter rutan", 44)
    renderer = DocxBlockImageRenderer(document)
    assert renderer.list_label(after) == "2."
    assert _ink(_png(renderer.render(host._p, anchor="body:0001"))) > 0
