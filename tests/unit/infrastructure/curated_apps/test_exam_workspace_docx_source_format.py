"""Source-format reader proof: run marks, paragraph headings and their chunk offsets."""

import io

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_UNDERLINE
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt
from PIL import Image

from skriptoteket.domain.curated_apps.exam_workspace.docx_chunks import (
    DocxChunk,
    DocxSourceParagraph,
    DocxSourceRun,
)
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_docx_chunker import (  # noqa: E501
    PythonDocxExamChunker,
)
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_docx_source_format import (  # noqa: E501
    document_body_points,
)


def _bytes(document: DocxDocument) -> bytes:
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _chunks(document: DocxDocument) -> tuple[DocxChunk, ...]:
    return PythonDocxExamChunker().chunk(content=_bytes(document)).chunks


def _defaults_rpr_default(document: DocxDocument):
    defaults = document.styles.element.find(qn("w:docDefaults"))
    return defaults.find(qn("w:rPrDefault"))


def _png() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (10, 10), (10, 20, 30)).save(buffer, format="PNG")
    return buffer.getvalue()


def test_direct_run_marks_are_kept_in_canonical_order() -> None:
    document = Document()
    paragraph = document.add_paragraph()
    all_marks = paragraph.add_run("alla")
    all_marks.bold = True
    all_marks.italic = True
    all_marks.underline = True
    all_marks.font.superscript = True
    paragraph.add_run("upp").font.superscript = True
    paragraph.add_run("ner").font.subscript = True

    (chunk,) = _chunks(document)

    assert chunk.paragraphs == (
        DocxSourceParagraph(offset=0, text="allauppner", heading_level=None),
    )
    assert chunk.runs == (
        DocxSourceRun(0, "alla", ("bold", "italic", "underline", "superscript")),
        DocxSourceRun(4, "upp", ("superscript",)),
        DocxSourceRun(7, "ner", ("subscript",)),
    )


def test_underline_styles_count_and_explicit_false_cancels() -> None:
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("dubbel").font.underline = WD_UNDERLINE.DOUBLE
    plain = paragraph.add_run(" utan")
    plain.font.underline = False

    (chunk,) = _chunks(document)

    assert chunk.runs == (
        DocxSourceRun(0, "dubbel", ("underline",)),
        DocxSourceRun(6, " utan", ()),
    )


def test_character_style_chain_inherits_marks_and_direct_false_cancels_one() -> None:
    document = Document()
    base = document.styles.add_style("Bas", WD_STYLE_TYPE.CHARACTER)
    base.font.underline = True
    derived = document.styles.add_style("Härledd", WD_STYLE_TYPE.CHARACTER)
    derived.base_style = base
    derived.font.italic = True
    paragraph = document.add_paragraph()
    paragraph.add_run("ärvd").style = derived
    cancelled = paragraph.add_run(" rak")
    cancelled.style = derived
    cancelled.font.underline = False

    (chunk,) = _chunks(document)

    assert chunk.runs == (
        DocxSourceRun(0, "ärvd", ("italic", "underline")),
        DocxSourceRun(4, " rak", ("italic",)),
    )


def test_subscript_in_a_style_never_adds_to_a_direct_superscript() -> None:
    document = Document()
    lowered = document.styles.add_style("Sänkt", WD_STYLE_TYPE.CHARACTER)
    lowered.font.subscript = True
    paragraph = document.add_paragraph()
    raised = paragraph.add_run("upp")
    raised.style = lowered
    raised.font.superscript = True

    (chunk,) = _chunks(document)

    assert chunk.runs == (DocxSourceRun(0, "upp", ("superscript",)),)


def test_paragraph_style_inherits_and_a_character_style_false_cancels_it() -> None:
    document = Document()
    bold_style = document.styles.add_style("Fet text", WD_STYLE_TYPE.PARAGRAPH)
    bold_style.font.bold = True
    plain_style = document.styles.add_style("Rak", WD_STYLE_TYPE.CHARACTER)
    plain_style.font.bold = False
    paragraph = document.add_paragraph("fet", style=bold_style)
    plain = paragraph.add_run(" rak")
    plain.style = plain_style

    (chunk,) = _chunks(document)

    assert chunk.runs == (
        DocxSourceRun(0, "fet", ("bold",)),
        DocxSourceRun(3, " rak", ()),
    )


def test_document_defaults_apply_last_and_an_explicit_false_cancels_them() -> None:
    document = Document()
    rpr_default = _defaults_rpr_default(document)
    existing = rpr_default.find(qn("w:rPr"))
    rpr_default.replace(
        existing,
        parse_xml(
            f'<w:rPr {nsdecls("w")}><w:b/><w:i/><w:u w:val="single"/><w:sz w:val="18"/></w:rPr>'
        ),
    )
    paragraph = document.add_paragraph()
    paragraph.add_run("standard")
    cancelled = paragraph.add_run(" rak")
    cancelled.font.bold = False

    (chunk,) = _chunks(document)

    assert document_body_points(document) == 9.0
    assert chunk.runs == (
        DocxSourceRun(0, "standard", ("bold", "italic", "underline")),
        DocxSourceRun(8, " rak", ("italic", "underline")),
    )


def test_model_marked_text_stays_direct_italics_while_runs_resolve_styles() -> None:
    document = Document()
    italic_style = document.styles.add_style("Kursiv text", WD_STYLE_TYPE.PARAGRAPH)
    italic_style.font.italic = True
    document.add_paragraph("kursiv", style=italic_style)

    (chunk,) = _chunks(document)

    assert chunk.marked_text == "kursiv"
    assert chunk.runs == (DocxSourceRun(0, "kursiv", ("italic",)),)


def test_title_and_heading_styles_map_to_levels_one_to_three() -> None:
    document = Document()
    document.add_paragraph("Titel", style="Title")
    document.add_heading("Ett", level=1)
    document.add_heading("Två", level=2)
    document.add_heading("Tre", level=3)
    document.add_paragraph("Fyra", style="Heading 4")

    chunks = _chunks(document)

    assert [chunk.paragraphs[0].heading_level for chunk in chunks] == [1, 1, 2, 3, None]


def test_size_headings_use_fixed_ratios_and_need_every_run_larger() -> None:
    document = Document()
    document.styles["Normal"].font.size = Pt(12)
    document.add_paragraph("Brödtext " * 12)
    document.add_paragraph().add_run("Störst").font.size = Pt(24)
    document.add_paragraph().add_run("Större").font.size = Pt(18)
    document.add_paragraph().add_run("Stor").font.size = Pt(15)
    document.add_paragraph().add_run("Nästan").font.size = Pt(13)
    mixed = document.add_paragraph()
    mixed.add_run("Stor").font.size = Pt(30)
    mixed.add_run(" blandad").font.size = Pt(12)
    inline = document.add_paragraph("Vanlig text med ")
    inline.add_run("stort").font.size = Pt(30)

    chunks = _chunks(document)

    assert [chunk.paragraphs[0].heading_level for chunk in chunks] == [
        None,
        1,
        2,
        3,
        None,
        None,
        None,
    ]


def test_size_heading_resolves_the_size_through_the_paragraph_style() -> None:
    document = Document()
    document.add_paragraph("Brödtext " * 12)
    large_style = document.styles.add_style("Stor text", WD_STYLE_TYPE.PARAGRAPH)
    large_style.font.size = Pt(22)
    document.add_paragraph("Stor stil", style=large_style)

    chunks = _chunks(document)

    assert [chunk.paragraphs[0].heading_level for chunk in chunks] == [None, 1]


def test_body_size_is_the_most_used_size_and_ties_go_to_the_smaller_size() -> None:
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("elva").font.size = Pt(11)
    paragraph.add_run("tolv").font.size = Pt(12)

    assert document_body_points(document) == 11.0

    paragraph.add_run("tolv igen").font.size = Pt(12)

    assert document_body_points(document) == 12.0


def test_body_size_falls_back_to_the_normal_style_then_eleven_points() -> None:
    document = Document()
    assert document_body_points(document) == 11.0

    document.styles["Normal"].font.size = Pt(10)

    assert document_body_points(document) == 10.0


def test_table_offsets_tile_joined_text_and_keep_picture_places() -> None:
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Vänster"
    cell = table.cell(0, 1)
    paragraph = cell.paragraphs[0]
    bold = paragraph.add_run("fet")
    bold.bold = True
    paragraph.add_run(" och ")
    italic = paragraph.add_run("snett")
    italic.italic = True
    paragraph.add_run().add_picture(io.BytesIO(_png()))
    cell.add_paragraph("Andra")

    (chunk,) = _chunks(document)

    assert chunk.kind == "table"
    assert chunk.text == "Vänster | fet och snett\nAndra"
    assert chunk.marked_text == "Vänster | fet och <i>snett</i>\nAndra"
    assert chunk.paragraphs == (
        DocxSourceParagraph(offset=0, text="Vänster", heading_level=None),
        DocxSourceParagraph(offset=10, text="fet och snett", heading_level=None),
        DocxSourceParagraph(offset=24, text="Andra", heading_level=None),
    )
    assert chunk.runs == (
        DocxSourceRun(offset=0, text="Vänster", marks=()),
        DocxSourceRun(offset=10, text="fet", marks=("bold",)),
        DocxSourceRun(offset=13, text=" och ", marks=()),
        DocxSourceRun(offset=18, text="snett", marks=("italic",)),
        DocxSourceRun(offset=24, text="Andra", marks=()),
    )
    for run in chunk.runs:
        assert chunk.text[run.offset : run.offset + len(run.text)] == run.text
    for source in chunk.paragraphs:
        assert chunk.text[source.offset : source.offset + len(source.text)] == source.text
    (image,) = chunk.images
    assert image.offset == 23
    assert image.marked_offset == chunk.marked_text.index("</i>") + len("</i>")
