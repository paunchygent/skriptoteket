"""Content, house-template and fail-closed tests for the exam workspace DOCX writer (D3).

The Word export is the student copy (W2) rendered from the print document; these tests read
the produced bytes back with python-docx and inspect the paragraph XML. The printed
answer-form geometry lives in the answer-forms suite and the shared-picture round trips in
the pictures suite.
"""

import copy
import io
import zipfile
from collections.abc import Callable
from pathlib import Path
from uuid import uuid4

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm
from lxml import etree  # type: ignore[import-untyped]  # transitive of python-docx, no stubs

from skriptoteket.domain.curated_apps.exam_conversion.exam_print_house_template import (
    BODY_PARAGRAPH_SPACING,
    PARAGRAPH_SPACING,
    TEXT_WIDTH_MM,
)
from skriptoteket.domain.curated_apps.exam_workspace.exam_language_labels import (
    SWEDISH_EXAM_LABELS,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_body import (
    NativeAssetSegment,
    NativeParagraph,
    NativeTextSegment,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_document import (
    NativeExamAsset,
    NativeExamDocument,
    NativeExamDocumentOrigin,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_item import (
    NativeExamItem,
)
from skriptoteket.domain.errors import DomainError, ErrorCode
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub import (
    exam_workspace_docx_writer,
)
from tests.unit.domain.curated_apps.exam_workspace.native_exam_test_builders import (
    make_document,
    make_free_text_item,
)
from tests.unit.infrastructure.curated_apps.exam_workspace_docx_writer_test_support import (
    STEP_BOX,
    all_kinds_document,
    document_text,
    heading_texts,
    only_paragraph,
    png,
)

ExamWorkspaceDocxWriter = exam_workspace_docx_writer.ExamWorkspaceDocxWriter

_FIXTURE = Path("tests/fixtures/exam_conversion/native_workspace_grammatik.json")


@pytest.fixture(scope="module")
def document() -> NativeExamDocument:
    return all_kinds_document()


@pytest.fixture(scope="module")
def docx_bytes(document: NativeExamDocument) -> bytes:
    return ExamWorkspaceDocxWriter().build_docx_bytes(document)


class TestDocxContent:
    def test_round_trip_validation_passes(
        self, docx_bytes: bytes, document: NativeExamDocument
    ) -> None:
        ExamWorkspaceDocxWriter().validate_docx_bytes(
            docx_bytes, expected_item_count=len(document.items)
        )

    def test_title_instructions_and_item_headings(self, docx_bytes: bytes) -> None:
        parsed = Document(io.BytesIO(docx_bytes))
        texts = [paragraph.text for paragraph in parsed.paragraphs]

        assert "Grammatikprov" in texts
        assert "Svenska 2" in texts
        # Items carry no title (39-05-08); headings are the question labels.
        assert "Fråga 1" in texts
        assert "Fråga 2" in texts

    def test_points_and_tick_options_without_gap_markers(self, docx_bytes: bytes) -> None:
        parsed = Document(io.BytesIO(docx_bytes))
        texts = [paragraph.text for paragraph in parsed.paragraphs]

        assert "Resonera om bisatser. (4 p)" in texts
        assert "Välj rätt alternativ. (1 p)" in texts
        assert "Markera alla fraser. (2 p)" in texts
        assert any(text.startswith("a) Frastyp: ") and text.endswith(" (2,5 p)") for text in texts)
        assert f"{STEP_BOX} A) Nominalfras" in texts
        assert f"{STEP_BOX} B) Verbfras" in texts
        assert "[___]" not in document_text(parsed)

    def test_student_copy_carries_no_answer_key_or_criterion(self, docx_bytes: bytes) -> None:
        parsed = Document(io.BytesIO(docx_bytes))
        text = document_text(parsed)

        assert SWEDISH_EXAM_LABELS.answer_key not in heading_texts(parsed, "Heading 2")
        assert "Facit" not in text
        assert "Kriteriet nämner bisatser" not in text
        assert "nominalfras" not in text
        # The correct choice is still a printed option, so the student sees the wording.
        assert f"{STEP_BOX} B) Verbfras" in text

    def test_real_fixture_document_round_trips(self) -> None:
        document = NativeExamDocument.model_validate_json(_FIXTURE.read_bytes())
        content = ExamWorkspaceDocxWriter().build_docx_bytes(document)

        ExamWorkspaceDocxWriter().validate_docx_bytes(
            content, expected_item_count=len(document.items)
        )


class TestHouseTemplate:
    def test_normal_style_is_the_house_font_and_margins(self, docx_bytes: bytes) -> None:
        parsed = Document(io.BytesIO(docx_bytes))

        normal = parsed.styles["Normal"]
        assert normal.font.name == "Liberation Sans"
        assert normal.font.size.pt == 12
        assert normal.paragraph_format.line_spacing == 1.5

        section = parsed.sections[0]
        for margin in (
            section.left_margin,
            section.right_margin,
            section.top_margin,
            section.bottom_margin,
        ):
            assert margin is not None
            assert margin.mm == pytest.approx(18, abs=0.05)


class TestValidationFailsClosed:
    def test_corrupt_bytes_rejected(self) -> None:
        with pytest.raises(DomainError) as exc_info:
            ExamWorkspaceDocxWriter().validate_docx_bytes(b"not a docx", expected_item_count=1)

        assert exc_info.value.code is ErrorCode.VALIDATION_ERROR

    def test_wrong_item_heading_count_rejected(self, docx_bytes: bytes) -> None:
        with pytest.raises(DomainError) as exc_info:
            ExamWorkspaceDocxWriter().validate_docx_bytes(docx_bytes, expected_item_count=5)

        assert exc_info.value.code is ErrorCode.VALIDATION_ERROR
        assert exc_info.value.details["found_item_headings"] == 4


def _marked_body_item() -> NativeExamItem:
    body = (
        NativeParagraph(
            heading_level=None,
            segments=(
                NativeTextSegment(text="Diskutera", marks=("bold",)),
                NativeTextSegment(text=" hur ", marks=()),
            ),
        ),
        NativeParagraph(
            heading_level=None,
            segments=(NativeTextSegment(text="Andra stycket.", marks=()),),
        ),
    )
    return make_free_text_item().model_copy(update={"body": body})


class TestPrintedMarksAndBodySpacing:
    def test_bold_run_keeps_bold_and_plain_run_keeps_none(self) -> None:
        document = make_document(_marked_body_item())
        parsed = Document(io.BytesIO(ExamWorkspaceDocxWriter().build_docx_bytes(document)))
        paragraph = only_paragraph(list(parsed.paragraphs), "Diskutera hur ")
        by_text = {run.text: run for run in paragraph.runs}

        assert by_text["Diskutera"].bold is True
        assert by_text[" hur "].bold is None

    def test_two_body_paragraphs_carry_the_body_spacing(self) -> None:
        document = make_document(_marked_body_item())
        parsed = Document(io.BytesIO(ExamWorkspaceDocxWriter().build_docx_bytes(document)))

        for text in ("Diskutera hur ", "Andra stycket. (4 p)"):
            paragraph = only_paragraph(list(parsed.paragraphs), text)
            assert paragraph.paragraph_format.space_after is not None
            assert paragraph.paragraph_format.space_after.mm == pytest.approx(
                BODY_PARAGRAPH_SPACING.space_after_mm, abs=0.02
            )

    def test_each_mark_sets_only_its_word_property(self) -> None:
        body = (
            NativeParagraph(
                heading_level=None,
                segments=(
                    NativeTextSegment(text="kursiv", marks=("italic",)),
                    NativeTextSegment(text="under", marks=("underline",)),
                    NativeTextSegment(text="upp", marks=("superscript",)),
                    NativeTextSegment(text="ner", marks=("subscript",)),
                    NativeTextSegment(text="vanlig", marks=()),
                ),
            ),
        )
        item = make_free_text_item().model_copy(update={"body": body, "points": None})
        parsed = Document(
            io.BytesIO(ExamWorkspaceDocxWriter().build_docx_bytes(make_document(item)))
        )
        paragraph = only_paragraph(list(parsed.paragraphs), "kursivunderuppnervanlig")
        by_text = {run.text: run for run in paragraph.runs}

        assert by_text["kursiv"].italic is True
        assert by_text["kursiv"].bold is None
        assert by_text["under"].underline is True
        assert by_text["under"].bold is None
        assert by_text["upp"].font.superscript is True
        assert by_text["upp"].font.subscript is False
        assert by_text["ner"].font.subscript is True
        assert by_text["ner"].font.superscript is False
        plain = by_text["vanlig"]
        assert plain.bold is None
        assert plain.italic is None
        assert plain.underline is None
        assert plain.font.superscript is None
        assert plain.font.subscript is None

    def test_hard_break_inside_a_run_writes_a_word_line_break(self) -> None:
        body = (
            NativeParagraph(
                heading_level=None,
                segments=(NativeTextSegment(text="a\nb", marks=()),),
            ),
        )
        item = make_free_text_item().model_copy(update={"body": body, "points": None})
        parsed = Document(
            io.BytesIO(ExamWorkspaceDocxWriter().build_docx_bytes(make_document(item)))
        )
        paragraph = only_paragraph(list(parsed.paragraphs), "a\nb")

        breaks = [element for run in paragraph.runs for element in run._element.iter(qn("w:br"))]
        assert len(breaks) == 1

    def test_header_detail_line_keeps_the_plain_paragraph_spacing(self, docx_bytes: bytes) -> None:
        parsed = Document(io.BytesIO(docx_bytes))
        detail = only_paragraph(list(parsed.paragraphs), "Totalt: 9,5 poäng")

        assert detail.paragraph_format.space_after is not None
        assert detail.paragraph_format.space_after.mm == pytest.approx(
            PARAGRAPH_SPACING.space_after_mm, abs=0.02
        )


_SHAPES_FIXTURE = Path("tests/fixtures/exam_conversion/pdf_corpus/metodprov_med_bilder.docx")
_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
_WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_NAMESPACES = {
    "w": _W,
    "mc": _MC,
    "wp": _WP,
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "wpg": "http://schemas.microsoft.com/office/word/2010/wordprocessingGroup",
    "wps": "http://schemas.microsoft.com/office/word/2010/wordprocessingShape",
    "w14": "http://schemas.microsoft.com/office/word/2010/wordml",
    "wp14": "http://schemas.microsoft.com/office/word/2010/wordprocessingDrawing",
}
_ASSET_A = "asset_" + "a" * 16
_ASSET_B = "asset_" + "b" * 16
_ASSET_C = "asset_" + "c" * 16
_LABELS = ("Svenska gymnasieelever (populationen)", "Elever på ett program (urvalet)")


def _stored_group_shapes() -> str:
    """The fixture's drawn group (two labelled ellipses) as the import stores it."""
    with zipfile.ZipFile(_SHAPES_FIXTURE) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    drawings = root.xpath("//mc:Choice/w:drawing", namespaces=_NAMESPACES)
    assert drawings, "the fixture holds a drawn group in mc:Choice"
    return str(etree.tostring(copy.deepcopy(drawings[0]), encoding="unicode"))


def _edited_shapes(edit: Callable[[etree._Element], None]) -> str:
    drawing = etree.fromstring(_stored_group_shapes())
    edit(drawing)
    return str(etree.tostring(drawing, encoding="unicode"))


def _picture_document(
    *pictures: tuple[str, str | None],
) -> tuple[NativeExamDocument, dict[str, bytes]]:
    """One free-text question per (asset id, stored shapes), each with a picture in its body."""
    payloads = {asset_id: png((20, 90, 20, 255), (40, 20)) for asset_id, _ in pictures}
    items = []
    for sequence, (asset_id, _) in enumerate(pictures, start=1):
        item = make_free_text_item(sequence=sequence, item_id=f"item_{sequence:03d}")
        items.append(
            item.model_copy(
                update={
                    "body": (
                        NativeParagraph(
                            heading_level=None,
                            segments=(
                                NativeTextSegment(text=f"Se bild {sequence}", marks=()),
                                NativeAssetSegment(asset_id=asset_id),
                            ),
                        ),
                    )
                }
            )
        )
    assets = tuple(
        NativeExamAsset(
            asset_id=asset_id,
            sha256="0" * 64,
            media_type="image/png",
            byte_length=len(payloads[asset_id]),
            docx_shapes=shapes,
        )
        for asset_id, shapes in pictures
    )
    document = NativeExamDocument(
        document_id=uuid4(),
        revision=1,
        title="Bildprov",
        items=tuple(items),
        assets=assets,
        origin=NativeExamDocumentOrigin(kind="created"),
    )
    return document, payloads


def _export(*pictures: tuple[str, str | None]) -> tuple[bytes, etree._Element]:
    """Export, re-open and validate the file; return its bytes and its document.xml root."""
    document, payloads = _picture_document(*pictures)
    writer = ExamWorkspaceDocxWriter()
    content = writer.build_docx_bytes(document, assets_by_id=payloads)
    Document(io.BytesIO(content))
    writer.validate_docx_bytes(content, expected_item_count=len(document.items))
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        return content, etree.fromstring(archive.read("word/document.xml"))


_DRAWING_IDS = "//wp:docPr/@id | //wps:cNvPr/@id | //wpg:cNvPr/@id"


def _media_names(content: bytes) -> list[str]:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        return [name for name in archive.namelist() if name.startswith("word/media/")]


def _picture_paragraphs(root: etree._Element, label: str) -> list[etree._Element]:
    paragraphs: list[etree._Element] = root.xpath(
        f"//w:p[w:r/w:t[starts-with(., '{label}')]]",
        namespaces=_NAMESPACES,
    )
    return paragraphs


class TestStoredShapesReplaceThePicture:
    def test_group_shapes_are_written_as_an_inline_drawing_in_the_question_paragraph(
        self,
    ) -> None:
        content, root = _export((_ASSET_A, _stored_group_shapes()))

        (paragraph,) = _picture_paragraphs(root, "Se bild 1")
        (inline,) = paragraph.xpath(
            "./w:r/mc:AlternateContent/mc:Choice[@Requires='wpg']/w:drawing/wp:inline",
            namespaces=_NAMESPACES,
        )
        assert inline.xpath("./a:graphic/a:graphicData/wpg:wgp", namespaces=_NAMESPACES)
        written = "".join(inline.xpath(".//w:t/text()", namespaces=_NAMESPACES))
        for label in _LABELS:
            assert label in written
        assert (
            paragraph.xpath("./w:r/mc:AlternateContent/mc:Fallback", namespaces=_NAMESPACES) == []
        )
        assert root.xpath("//a:blip", namespaces=_NAMESPACES) == []
        assert _media_names(content) == []

    def test_the_picture_paragraph_keeps_the_picture_spacing(self) -> None:
        _, shape_root = _export((_ASSET_A, _stored_group_shapes()))
        _, plain_root = _export((_ASSET_A, None))

        (shape_paragraph,) = _picture_paragraphs(shape_root, "Se bild 1")
        (plain_paragraph,) = _picture_paragraphs(plain_root, "Se bild 1")
        spacing = "./w:pPr/w:spacing"
        assert dict(shape_paragraph.xpath(spacing, namespaces=_NAMESPACES)[0].attrib) == dict(
            plain_paragraph.xpath(spacing, namespaces=_NAMESPACES)[0].attrib
        )

    def test_two_pictures_get_unique_drawing_ids_and_no_word_revision_ids(self) -> None:
        shapes = _stored_group_shapes()
        _, root = _export((_ASSET_A, shapes), (_ASSET_B, shapes))

        doc_pr_ids = root.xpath("//wp:docPr/@id", namespaces=_NAMESPACES)
        assert len(doc_pr_ids) == 2
        every_shape_id = root.xpath(_DRAWING_IDS, namespaces=_NAMESPACES)
        assert len(every_shape_id) == 6  # two docPr and two shapes in each of two drawings
        assert len(every_shape_id) == len(set(every_shape_id))
        drawings = root.xpath("//w:drawing", namespaces=_NAMESPACES)
        assert len(drawings) == 2
        for drawing in drawings:
            assert drawing.xpath(".//@w14:paraId", namespaces=_NAMESPACES) == []
            assert drawing.xpath(".//@w14:textId", namespaces=_NAMESPACES) == []
            assert drawing.xpath(".//@wp14:anchorId", namespaces=_NAMESPACES) == []
            assert drawing.xpath(".//@wp14:editId", namespaces=_NAMESPACES) == []
            assert drawing.xpath(".//@*[starts-with(local-name(), 'rsid')]") == []
        assert set(every_shape_id).isdisjoint({"9001", "9101", "9102"})

    def test_shape_ids_stay_unique_beside_ordinary_pictures(self) -> None:
        shapes = _stored_group_shapes()
        content, root = _export((_ASSET_A, None), (_ASSET_B, shapes), (_ASSET_C, None))

        every_shape_id = root.xpath(_DRAWING_IDS, namespaces=_NAMESPACES)
        assert len(every_shape_id) == 5  # two plain docPr and the group's docPr and two shapes
        assert len(every_shape_id) == len(set(every_shape_id))
        assert len(root.xpath("//a:blip", namespaces=_NAMESPACES)) == 2
        assert len(_media_names(content)) == 1  # identical PNG bytes share one image part
        assert len(root.xpath("//mc:AlternateContent", namespaces=_NAMESPACES)) == 1

    def test_a_stored_anchor_is_written_as_an_inline(self) -> None:
        def to_anchor(drawing: etree._Element) -> None:
            (inline,) = drawing
            anchor = etree.Element(
                f"{{{_WP}}}anchor",
                {
                    "distT": "0",
                    "distB": "0",
                    "distL": "0",
                    "distR": "0",
                    "simplePos": "0",
                    "relativeHeight": "251658240",
                    "behindDoc": "0",
                    "locked": "0",
                    "layoutInCell": "1",
                    "allowOverlap": "1",
                    f"{{{_NAMESPACES['wp14']}}}anchorId": "2BABB811",
                },
                nsmap={"wp": _WP, "wp14": _NAMESPACES["wp14"]},
            )
            etree.SubElement(anchor, f"{{{_WP}}}simplePos", {"x": "0", "y": "0"})
            horizontal = etree.SubElement(anchor, f"{{{_WP}}}positionH", {"relativeFrom": "column"})
            etree.SubElement(horizontal, f"{{{_WP}}}posOffset").text = "100"
            vertical = etree.SubElement(
                anchor, f"{{{_WP}}}positionV", {"relativeFrom": "paragraph"}
            )
            etree.SubElement(vertical, f"{{{_WP}}}posOffset").text = "200"
            for child in list(inline):
                anchor.append(child)
                if child.tag == f"{{{_WP}}}effectExtent":
                    etree.SubElement(anchor, f"{{{_WP}}}wrapSquare", {"wrapText": "bothSides"})
            drawing.replace(inline, anchor)

        _, root = _export((_ASSET_A, _edited_shapes(to_anchor)))

        (drawing,) = root.xpath("//mc:Choice/w:drawing", namespaces=_NAMESPACES)
        (inline,) = drawing
        assert inline.tag == f"{{{_WP}}}inline"
        assert [child.tag.rpartition("}")[2] for child in inline] == [
            "extent",
            "effectExtent",
            "docPr",
            "cNvGraphicFramePr",
            "graphic",
        ]
        assert {name: inline.get(name) for name in ("distT", "distB", "distL", "distR")} == {
            "distT": "0",
            "distB": "0",
            "distL": "0",
            "distR": "0",
        }
        assert inline.get("simplePos") is None
        assert drawing.xpath(".//@wp14:anchorId", namespaces=_NAMESPACES) == []
        assert inline.xpath("./a:graphic//wpg:wgp", namespaces=_NAMESPACES)


def _add_embed(drawing: etree._Element) -> None:
    graphic_data = drawing.xpath(".//a:graphicData", namespaces=_NAMESPACES)[0]
    blip = etree.SubElement(
        graphic_data, f"{{{_NAMESPACES['a']}}}blip", {f"{{{_R}}}embed": "rId99"}, nsmap={"r": _R}
    )
    assert blip is not None


def _unknown_graphic(drawing: etree._Element) -> None:
    graphic_data = drawing.xpath(".//a:graphicData", namespaces=_NAMESPACES)[0]
    graphic_data.set("uri", "http://schemas.openxmlformats.org/drawingml/2006/chart")


class TestUnusableShapesWriteThePicture:
    @pytest.mark.parametrize(
        "stored",
        [
            pytest.param("<w:drawing xmlns:w=", id="malformed-xml"),
            pytest.param(
                f'<w:p xmlns:w="{_W}"><w:r><w:t>text</w:t></w:r></w:p>', id="root-is-not-a-drawing"
            ),
            pytest.param(_edited_shapes(_add_embed), id="relationship-reference"),
            pytest.param(_edited_shapes(_unknown_graphic), id="not-a-shape-graphic"),
        ],
    )
    def test_the_picture_is_written_instead(self, stored: str) -> None:
        content, root = _export((_ASSET_A, stored))

        assert len(root.xpath("//a:blip", namespaces=_NAMESPACES)) == 1
        assert root.xpath("//mc:AlternateContent", namespaces=_NAMESPACES) == []
        assert len(_media_names(content)) == 1
        assert root.xpath("//wpg:wgp", namespaces=_NAMESPACES) == []


_WPC = "http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas"
_NAMESPACES["wpc"] = _WPC
_SOURCE_NAMESPACES = (
    f'xmlns:w="{_W}" xmlns:wp="{_WP}" xmlns:a="{_NAMESPACES["a"]}" '
    f'xmlns:wpg="{_NAMESPACES["wpg"]}" xmlns:wps="{_NAMESPACES["wps"]}" xmlns:wpc="{_WPC}"'
)


def _rectangle(shape_id: int, name: str, x: int) -> str:
    return (
        f'<wps:wsp><wps:cNvPr id="{shape_id}" name="{name}"/><wps:cNvSpPr/><wps:spPr>'
        f'<a:xfrm><a:off x="{x}" y="0"/><a:ext cx="1000" cy="500"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></wps:spPr><wps:bodyPr/></wps:wsp>'
    )


def _connector(shape_id: int, start: int, end: int) -> str:
    return (
        f'<wps:wsp><wps:cNvPr id="{shape_id}" name="Pil"/><wps:cNvCnPr>'
        f'<a:stCxn id="{start}" idx="3"/><a:endCxn id="{end}" idx="1"/></wps:cNvCnPr><wps:spPr>'
        '<a:xfrm><a:off x="1000" y="250"/><a:ext cx="1000" cy="0"/></a:xfrm>'
        '<a:prstGeom prst="straightConnector1"><a:avLst/></a:prstGeom></wps:spPr>'
        "<wps:bodyPr/></wps:wsp>"
    )


def _drawing(graphic_uri: str, root: str, *, cx: int, cy: int) -> str:
    return (
        f"<w:drawing {_SOURCE_NAMESPACES}>"
        '<wp:inline distT="0" distB="0" distL="0" distR="0">'
        f'<wp:extent cx="{cx}" cy="{cy}"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
        '<wp:docPr id="7" name="Figur"/><wp:cNvGraphicFramePr/>'
        f'<a:graphic><a:graphicData uri="{graphic_uri}">{root}</a:graphicData></a:graphic>'
        "</wp:inline></w:drawing>"
    )


def _group_xml(children: str, *, cx: int = 3000, cy: int = 1000) -> str:
    return _drawing(
        _NAMESPACES["wpg"],
        "<wpg:wgp><wpg:cNvGrpSpPr/><wpg:grpSpPr><a:xfrm>"
        f'<a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/>'
        f'<a:chOff x="0" y="0"/><a:chExt cx="{cx}" cy="{cy}"/>'
        f"</a:xfrm></wpg:grpSpPr>{children}</wpg:wgp>",
        cx=cx,
        cy=cy,
    )


def _single_xml(*, cx: int, cy: int) -> str:
    return _drawing(
        _NAMESPACES["wps"],
        '<wps:wsp><wps:cNvPr id="3" name="Ensam"/><wps:cNvSpPr/><wps:spPr>'
        f'<a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></wps:spPr><wps:bodyPr/></wps:wsp>',
        cx=cx,
        cy=cy,
    )


def _canvas_xml(*, cx: int, cy: int) -> str:
    return _drawing(
        _WPC,
        f"<wpc:wpc><wpc:bg/><wpc:whole/>{_rectangle(3, 'Duk', 0)}</wpc:wpc>",
        cx=cx,
        cy=cy,
    )


def _connected_group_xml(*, start: int, end: int) -> str:
    return _group_xml(
        _rectangle(100, "Vänster", 0) + _rectangle(101, "Höger", 2000) + _connector(102, start, end)
    )


def _connection_ids(drawing: etree._Element) -> tuple[list[str], list[str]]:
    starts: list[str] = drawing.xpath(".//a:stCxn/@id", namespaces=_NAMESPACES)
    ends: list[str] = drawing.xpath(".//a:endCxn/@id", namespaces=_NAMESPACES)
    return starts, ends


def _shape_id(drawing: etree._Element, name: str) -> str:
    ids: list[str] = drawing.xpath(f".//wps:cNvPr[@name='{name}']/@id", namespaces=_NAMESPACES)
    assert len(ids) == 1
    return ids[0]


class TestConnectorsKeepTheirShapes:
    def test_each_copy_connects_its_own_renumbered_rectangles(self) -> None:
        group = _connected_group_xml(start=100, end=101)
        _, root = _export((_ASSET_A, None), (_ASSET_B, group), (_ASSET_C, group))

        drawings = root.xpath("//mc:Choice/w:drawing", namespaces=_NAMESPACES)
        assert len(drawings) == 2
        for drawing in drawings:
            starts, ends = _connection_ids(drawing)
            assert starts == [_shape_id(drawing, "Vänster")]
            assert ends == [_shape_id(drawing, "Höger")]
        every_doc_pr_id = root.xpath("//wp:docPr/@id", namespaces=_NAMESPACES)
        assert len(every_doc_pr_id) == 3  # the plain picture and the two group copies
        every_id = root.xpath(_DRAWING_IDS, namespaces=_NAMESPACES)
        assert len(every_id) == len(set(every_id))

    def test_a_connection_naming_a_shape_outside_the_drawing_is_removed(self) -> None:
        _, root = _export((_ASSET_A, _connected_group_xml(start=999, end=101)))

        (drawing,) = root.xpath("//w:drawing", namespaces=_NAMESPACES)
        starts, ends = _connection_ids(drawing)
        assert starts == []
        assert ends == [_shape_id(drawing, "Höger")]


class TestWideDrawingsAreFittedNotReplaced:
    @pytest.mark.parametrize(
        ("stored", "root_path"),
        [
            pytest.param(
                lambda cx, cy: _group_xml(_rectangle(100, "Vänster", 0), cx=cx, cy=cy),
                ".//wpg:grpSpPr/a:xfrm/a:ext",
                id="group",
            ),
            pytest.param(
                lambda cx, cy: _single_xml(cx=cx, cy=cy),
                ".//wps:spPr/a:xfrm/a:ext",
                id="single-shape",
            ),
        ],
    )
    def test_a_wide_drawing_is_scaled_to_the_text_width_keeping_its_aspect(
        self, stored: Callable[[int, int], str], root_path: str
    ) -> None:
        limit = int(Mm(TEXT_WIDTH_MM))
        cx, cy = limit + 1000, 1234567
        content, root = _export((_ASSET_A, stored(limit - 1000, cy)), (_ASSET_B, stored(cx, cy)))

        assert root.xpath("//a:blip", namespaces=_NAMESPACES) == []
        assert _media_names(content) == []
        control, wide = root.xpath("//wp:inline", namespaces=_NAMESPACES)
        assert control.xpath("./wp:extent/@cx", namespaces=_NAMESPACES) == [str(limit - 1000)]
        (extent,) = wide.xpath("./wp:extent", namespaces=_NAMESPACES)
        assert int(extent.get("cx", "")) == limit
        expected_cy = cy * limit / cx
        assert abs(int(extent.get("cy", "")) - expected_cy) <= 1
        (box,) = wide.xpath(root_path, namespaces=_NAMESPACES)
        assert int(box.get("cx", "")) == limit
        assert abs(int(box.get("cy", "")) - expected_cy) <= 1
        assert control.xpath(root_path + "/@cy", namespaces=_NAMESPACES) == [str(cy)]

    def test_a_wide_group_keeps_its_child_coordinate_space(self) -> None:
        limit = int(Mm(TEXT_WIDTH_MM))
        _, root = _export((_ASSET_A, _group_xml(_rectangle(100, "Vänster", 0), cx=limit + 1000)))

        (transform,) = root.xpath(".//wpg:grpSpPr/a:xfrm", namespaces=_NAMESPACES)
        (child_extent,) = transform.xpath("./a:chExt", namespaces=_NAMESPACES)
        assert child_extent.get("cx") == str(limit + 1000)
        assert child_extent.get("cy") == "1000"

    def test_a_wide_canvas_keeps_its_source_size_and_stays_editable(self) -> None:
        limit = int(Mm(TEXT_WIDTH_MM))
        content, root = _export((_ASSET_A, _canvas_xml(cx=limit + 1000, cy=777)))

        assert root.xpath("//a:blip", namespaces=_NAMESPACES) == []
        assert _media_names(content) == []
        (extent,) = root.xpath("//wp:inline/wp:extent", namespaces=_NAMESPACES)
        assert (extent.get("cx"), extent.get("cy")) == (str(limit + 1000), "777")
        assert len(root.xpath("//wpc:wpc", namespaces=_NAMESPACES)) == 1
