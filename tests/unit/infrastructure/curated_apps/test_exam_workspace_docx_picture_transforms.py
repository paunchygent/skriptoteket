"""Word picture transforms: crop, flip and quarter turns, or a named gate (TASK-SKRIPT-39-05-09)."""

import io

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from PIL import Image

from skriptoteket.domain.curated_apps.exam_workspace.docx_chunks import DocxMissingImage
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_docx_chunker import (  # noqa: E501
    PythonDocxExamChunker,
)
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_docx_picture import (  # noqa: E501
    picture_from_relationship,
)
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_labeler_block_image import (  # noqa: E501
    DocxBlockImageRenderer,
    LabelerSourceImageUnavailable,
)
from tests.unit.infrastructure.curated_apps.exam_workspace_docx_picture_support import (
    ALPHA_ROWS,
    chunk_document,
    document_with_picture,
    halves,
    jpeg,
    pixels,
    render_pixels,
    rgba_pixels,
    rgba_png,
    save_document,
    set_src_rect,
)

pytestmark = pytest.mark.unit

_A, _B, _C, _D = (255, 0, 0), (0, 200, 0), (0, 0, 255), (255, 255, 0)
_E, _F, _G, _H = (255, 0, 255), (0, 255, 255), (128, 64, 32), (9, 18, 27)
_QUARTER_TURNS = [
    pytest.param({"flipH": "1"}, [[_D, _C, _B, _A], [_H, _G, _F, _E]], id="flipH"),
    pytest.param({"flipV": "1"}, [[_E, _F, _G, _H], [_A, _B, _C, _D]], id="flipV"),
    pytest.param({"rot": str(90 * 60000)}, [[_E, _A], [_F, _B], [_G, _C], [_H, _D]], id="rot90"),
    pytest.param({"rot": str(180 * 60000)}, [[_H, _G, _F, _E], [_D, _C, _B, _A]], id="rot180"),
    pytest.param({"rot": str(270 * 60000)}, [[_D, _H], [_C, _G], [_B, _F], [_A, _E]], id="rot270"),
]
_ALPHA_TRANSFORMS = [
    pytest.param(
        "crop",
        {"r": "50000"},
        [[(255, 0, 0, 0), (0, 255, 0, 0)], [(0, 0, 0, 0), (0, 0, 0, 128)]],
        id="crop",
    ),
    pytest.param(
        "frame",
        {"flipH": "1"},
        [
            [(255, 255, 0, 255), (0, 0, 255, 128), (0, 255, 0, 0), (255, 0, 0, 0)],
            [(40, 50, 60, 255), (10, 20, 30, 0), (0, 0, 0, 128), (0, 0, 0, 0)],
        ],
        id="flipH",
    ),
    pytest.param(
        "frame",
        {"rot": str(90 * 60000)},
        [
            [(0, 0, 0, 0), (255, 0, 0, 0)],
            [(0, 0, 0, 128), (0, 255, 0, 0)],
            [(10, 20, 30, 0), (0, 0, 255, 128)],
            [(40, 50, 60, 255), (255, 255, 0, 255)],
        ],
        id="rot90",
    ),
]


def _png(rows: list[list[tuple[int, int, int]]]) -> bytes:
    image = Image.new("RGB", (len(rows[0]), len(rows)))
    image.putdata([pixel for row in rows for pixel in row])
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def _checkerboard() -> bytes:
    return _png([[_A, _B, _C, _D], [_E, _F, _G, _H]])


def _set_xfrm(blip, **attributes: str) -> None:
    picture = next(blip.iterancestors(qn("pic:pic")))
    properties = picture.find(qn("pic:spPr"))
    assert properties is not None
    frame = properties.find(qn("a:xfrm"))
    assert frame is not None
    for name, value in attributes.items():
        frame.set(name, value)


def _set_geometry(blip, prst: str) -> None:
    picture = next(blip.iterancestors(qn("pic:pic")))
    properties = picture.find(qn("pic:spPr"))
    assert properties is not None
    geometry = properties.find(qn("a:prstGeom"))
    assert geometry is not None
    geometry.set("prst", prst)


def _vml_xml(payload: bytes, *, style: str, attributes: str = "") -> bytes:
    document = Document()
    part = Part(PackURI("/word/media/vml.png"), "image/png", payload, None)
    rid = document.part.relate_to(part, RELATIONSHIP_TYPE.IMAGE)
    paragraph = document.add_paragraph()
    spacing = f" {attributes}" if attributes else ""
    paragraph.add_run()._r.append(
        parse_xml(
            '<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
            ' xmlns:v="urn:schemas-microsoft-com:vml"'
            ' xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            f'<v:shape style="{style}"><v:imagedata r:id="{rid}"{spacing}/></v:shape></w:pict>'
        )
    )
    return save_document(document)


def test_a_cropped_picture_keeps_only_the_visible_half() -> None:
    document, blip, paragraph = document_with_picture(halves())
    set_src_rect(blip, r="50000")

    (chunk,) = chunk_document(document)

    assert chunk.missing_images == ()
    (image,) = chunk.images
    assert image.media_type == "image/png"
    assert (image.width, image.height) == (100, 100)
    assert set(pixels(image.payload)) == {(255, 0, 0)}

    rendered = render_pixels(document, paragraph)
    assert all(blue <= red + 40 for red, _, blue in rendered)
    assert any(red > blue + 100 for red, _, blue in rendered)


def test_a_non_right_rotation_is_missing_and_the_renderer_refuses_the_block() -> None:
    document, blip, paragraph = document_with_picture(halves())
    _set_xfrm(blip, rot=str(45 * 60000))

    (chunk,) = chunk_document(document)

    assert chunk.images == ()
    (missing,) = chunk.missing_images
    assert isinstance(missing, DocxMissingImage)
    assert missing.media_type == "image/png"
    assert missing.reference.startswith("/word/media/image")
    assert "rotation som inte är ett kvarts varv" in missing.reference
    with pytest.raises(LabelerSourceImageUnavailable) as raised:
        DocxBlockImageRenderer(document).render(paragraph, anchor="body:0001")
    assert "rotation" in raised.value.reason


@pytest.mark.parametrize(("attributes", "expected"), _QUARTER_TURNS)
def test_flips_and_quarter_turns_turn_the_pixels(attributes: dict[str, str], expected) -> None:
    document, blip, _ = document_with_picture(_checkerboard())
    _set_xfrm(blip, **attributes)

    (chunk,) = chunk_document(document)

    assert chunk.missing_images == ()
    (image,) = chunk.images
    assert image.media_type == "image/png"
    assert (image.width, image.height) == (len(expected[0]), len(expected))
    assert pixels(image.payload) == [pixel for row in expected for pixel in row]


@pytest.mark.parametrize(("kind", "attributes", "expected"), _ALPHA_TRANSFORMS)
def test_an_rgba_picture_keeps_its_alpha_through_crop_flip_and_quarter_turn(
    kind: str, attributes: dict[str, str], expected
) -> None:
    document, blip, _ = document_with_picture(rgba_png(ALPHA_ROWS))
    if kind == "crop":
        set_src_rect(blip, **attributes)
    else:
        _set_xfrm(blip, **attributes)

    (chunk,) = chunk_document(document)

    assert chunk.missing_images == ()
    (image,) = chunk.images
    assert image.media_type == "image/png"
    assert (image.width, image.height) == (len(expected[0]), len(expected))
    assert rgba_pixels(image.payload) == [pixel for row in expected for pixel in row]


def test_a_noop_transform_keeps_png_and_jpeg_bytes_exact() -> None:
    png = halves()
    jpeg_bytes = jpeg()
    document = Document()
    document.add_picture(io.BytesIO(png))
    document.paragraphs[0].add_run().add_picture(io.BytesIO(jpeg_bytes))
    for blip in document.element.body.iter(qn("a:blip")):
        set_src_rect(blip, l="0", t="0", r="0", b="0")
        _set_xfrm(blip, rot="0", flipH="0", flipV="0")

    (chunk,) = chunk_document(document)

    assert chunk.missing_images == ()
    assert [image.payload for image in chunk.images] == [png, jpeg_bytes]
    assert [image.media_type for image in chunk.images] == ["image/png", "image/jpeg"]


@pytest.mark.parametrize(
    "attributes",
    [{"l": "abc"}, {"l": "-1"}, {"l": "60000", "r": "60000"}],
    ids=["unparsable", "negative", "empty"],
)
def test_an_invalid_crop_is_a_named_missing(attributes: dict[str, str]) -> None:
    document, blip, _ = document_with_picture(halves())
    set_src_rect(blip, **attributes)

    (chunk,) = chunk_document(document)

    assert chunk.images == ()
    (missing,) = chunk.missing_images
    assert missing.media_type == "image/png"
    assert "beskärning" in missing.reference


def test_an_invalid_flip_is_a_named_missing() -> None:
    document, blip, _ = document_with_picture(halves())
    _set_xfrm(blip, flipH="maybe")

    (chunk,) = chunk_document(document)

    assert chunk.images == ()
    (missing,) = chunk.missing_images
    assert "spegling" in missing.reference


def test_a_non_rectangular_shape_is_a_named_missing() -> None:
    document, blip, _ = document_with_picture(halves())
    _set_geometry(blip, "ellipse")

    (chunk,) = chunk_document(document)

    assert chunk.images == ()
    (missing,) = chunk.missing_images
    assert "rektangel" in missing.reference


def test_an_image_effect_is_a_named_missing() -> None:
    document, blip, _ = document_with_picture(halves())
    blip.append(parse_xml(f'<a:lum {nsdecls("a")} bright="20000" contrast="0"/>'))

    (chunk,) = chunk_document(document)

    assert chunk.images == ()
    (missing,) = chunk.missing_images
    assert "bildeffekt" in missing.reference


@pytest.mark.parametrize(
    ("style", "attributes"),
    [
        ("width:200pt;height:100pt", ""),
        ("width:200pt;height:100pt", 'cropleft="0"'),
        ("width:200pt;height:100pt;rotation:0", ""),
        ("width:200pt;height:100pt", 'grayscale="f"'),
    ],
    ids=["plain", "zero-crop", "zero-rotation", "false-effect"],
)
def test_a_vml_picture_without_transforms_stays_byte_exact(style: str, attributes: str) -> None:
    payload = halves()
    document = Document(io.BytesIO(_vml_xml(payload, style=style, attributes=attributes)))

    (chunk,) = chunk_document(document)

    assert chunk.missing_images == ()
    (image,) = chunk.images
    assert image.payload == payload
    assert image.media_type == "image/png"


@pytest.mark.parametrize(
    ("style", "attributes", "note"),
    [
        ("width:10pt;height:10pt", 'cropleft="0.1"', "beskärning"),
        ("width:10pt;height:10pt;rotation:45", "", "rotation"),
        ("width:10pt;height:10pt;flip:x", "", "rotation"),
    ],
    ids=["crop", "rotation", "flip"],
)
def test_a_vml_transform_is_a_named_missing(style: str, attributes: str, note: str) -> None:
    content = _vml_xml(halves(), style=style, attributes=attributes)

    (chunk,) = PythonDocxExamChunker().chunk(content=content).chunks

    assert chunk.images == ()
    (missing,) = chunk.missing_images
    assert missing.media_type == "image/png"
    assert note in missing.reference


def test_a_missing_relationship_keeps_a_descriptive_name() -> None:
    missing = picture_from_relationship(Document(), None)

    assert isinstance(missing, DocxMissingImage)
    assert missing.reference == "bild utan inbäddad relation"
    assert missing.media_type == "application/octet-stream"


def test_a_drawing_without_an_embedded_relationship_is_a_named_missing() -> None:
    document, blip, _ = document_with_picture(halves())
    blip.attrib.pop(qn("r:embed"))

    (chunk,) = chunk_document(document)

    assert chunk.images == ()
    (missing,) = chunk.missing_images
    assert missing.reference == "bild utan inbäddad relation"
    assert missing.media_type == "application/octet-stream"
