"""A gap-count-mismatch review item drops its inline answer key and answer line (M4).

The source holds ``Fyll i: ____ år. Svar: 2000`` plus a printed ``_____`` paragraph.
The reply carries two gaps for the single blank run, so admission records
``gap_count_mismatch`` and falls back to the printed source: the D7 cut removes the
inline answer key and the W10 gate (plus the review fallback) removes the printed
answer line. Both the direct ``assemble_build_result`` path and the application
``admit_batch`` + ``result`` path end with the question text in review. The control
without the answer-line paragraph ends with the same body and review state.
"""

import io
from uuid import uuid4

from docx import Document
from docx.document import Document as DocxDocument

from skriptoteket.application.curated_apps.handlers.exam_workspace_docx_build_admission import (
    admit_batch,
)
from skriptoteket.application.curated_apps.handlers.exam_workspace_docx_build_state import (
    DocxBuildState,
)
from skriptoteket.domain.curated_apps.exam_workspace.docx_ai_build import assemble_build_result
from skriptoteket.domain.curated_apps.exam_workspace.docx_ai_contracts import (
    DocxBuildResult,
    ProducedItem,
    ProducedParagraph,
    ProducedSegment,
)
from skriptoteket.domain.curated_apps.exam_workspace.docx_answer_line_gate import (
    ANSWER_LINE_ONLY_REASON,
)
from skriptoteket.domain.curated_apps.exam_workspace.docx_chunks import DocxChunk
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_vocabulary import (
    NativeExamItemKind,
    NativeItemReviewState,
)
from skriptoteket.domain.curated_apps.exam_workspace.shared_body import (
    body_plain_text,
    body_text_lines,
)
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_docx_chunker import (  # noqa: E501
    PythonDocxExamChunker,
)

_MISMATCH_REASON = "Frågan följer inte provformatet: 2 luckor men 1 svarsplats."
_QUESTION_TEXT = "Fyll i: ____ år."


def _chunked(document: DocxDocument) -> tuple[DocxChunk, ...]:
    stream = io.BytesIO()
    document.save(stream)
    return PythonDocxExamChunker().chunk(content=stream.getvalue())


def _source_document(with_answer_line: bool) -> DocxDocument:
    document = Document()
    document.add_paragraph("Fyll i: ____ år. Svar: 2000")
    if with_answer_line:
        document.add_paragraph("_____")
    return document


def _two_gap_reply(anchor: str) -> ProducedItem:
    return ProducedItem(
        source_anchors=(anchor,),
        outcome="converted",
        review_reason=None,
        kind=NativeExamItemKind.GAP_FILL,
        printed_label=None,
        key_source="copied",
        facit_text="svar",
        body=(
            ProducedParagraph(
                segments=(
                    ProducedSegment(kind="text", text="Fyll i: ", accepted_values=()),
                    ProducedSegment(kind="gap", text=None, accepted_values=("2000",)),
                    ProducedSegment(kind="gap", text=None, accepted_values=("2001",)),
                    ProducedSegment(kind="text", text=" år", accepted_values=()),
                )
            ),
        ),
        points=1,
        choices=(),
        match_left=(),
        match_right=(),
        match_pairs=(),
        criteria=(),
    )


def _admitted(chunks: tuple[DocxChunk, ...]) -> tuple[DocxBuildState, DocxBuildResult]:
    anchors = tuple(chunk.anchor for chunk in chunks)
    state = DocxBuildState(
        job_id=uuid4(),
        document_id=uuid4(),
        filename="synthetic.docx",
        source_sha256="0" * 64,
        chunks=chunks,
        questions=(anchors,),
    )
    admit_batch({"items": [_two_gap_reply(chunks[0].anchor).model_dump(mode="json")]}, state)
    return state, state.result()


def _assembled(chunks: tuple[DocxChunk, ...], state: DocxBuildState) -> DocxBuildResult:
    anchors = tuple(chunk.anchor for chunk in chunks)
    return assemble_build_result(
        document_id=uuid4(),
        filename="synthetic.docx",
        source_sha256="0" * 64,
        chunks=chunks,
        accepted=tuple(state.accepted),
        review_reasons={},
        questions=(anchors,),
    )


def test_a_mismatch_review_item_with_an_answer_line_keeps_only_its_question_text() -> None:
    chunks = _chunked(_source_document(with_answer_line=True))
    assert [chunk.text for chunk in chunks] == ["Fyll i: ____ år. Svar: 2000", "_____"]

    state, admitted = _admitted(chunks)
    assert [
        (event.code, event.detail)
        for event in state.events
        if event.stage != "formatting" and event.code == "gap_count_mismatch"
    ] == [("gap_count_mismatch", "2 gaps for 1 blank runs")]

    (item,) = admitted.document.items
    assert item.review.state is NativeItemReviewState.REVIEW_REQUIRED
    assert item.review.reasons == (_MISMATCH_REASON,)
    assert ANSWER_LINE_ONLY_REASON not in item.review.reasons
    assert body_text_lines(item.body) == (_QUESTION_TEXT,)
    assert body_plain_text(item.body) == _QUESTION_TEXT
    assert "Svar:" not in body_plain_text(item.body)
    assert "2000" not in body_plain_text(item.body)
    assert "[___]" not in body_plain_text(item.body)

    assembled = _assembled(chunks, state)
    (direct,) = assembled.document.items
    assert direct.review.state is NativeItemReviewState.REVIEW_REQUIRED
    assert direct.review.reasons == (_MISMATCH_REASON,)
    assert body_plain_text(direct.body) == _QUESTION_TEXT


def test_the_same_mismatch_without_an_answer_line_ends_with_the_same_body() -> None:
    plain_chunks = _chunked(_source_document(with_answer_line=False))
    lined_chunks = _chunked(_source_document(with_answer_line=True))

    plain_state, plain = _admitted(plain_chunks)
    _, lined = _admitted(lined_chunks)

    (plain_item,) = plain.document.items
    (lined_item,) = lined.document.items
    assert body_plain_text(plain_item.body) == body_plain_text(lined_item.body)
    assert plain_item.review.state is lined_item.review.state
    assert plain_item.review.reasons == lined_item.review.reasons

    assembled = _assembled(plain_chunks, plain_state)
    (direct,) = assembled.document.items
    assert body_plain_text(direct.body) == body_plain_text(plain_item.body)
    assert direct.review.state is plain_item.review.state
