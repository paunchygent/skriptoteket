"""The stored Word shapes state what the picture drew, not where to look it up (D8, D9, D12).

A drawing's colours (``a:schemeClr``, ``wps:style`` references) and its label's paragraph and
character styles resolve against the source package. The Word export writes the stored markup
into a new document with other themes and styles, so the stored copy must carry the resolved
values itself. Each test chunks a source file whose theme and styles differ from the export's,
writes the stored shapes through the real Word writer and reads the exported file back with the
shape reader: the exported figure must equal the picture's model.
"""

from __future__ import annotations

import io
import zipfile
from uuid import uuid4

import pytest
from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from lxml import etree

from skriptoteket.domain.curated_apps.exam_workspace.native_exam_body import (
    NativeAssetSegment,
    NativeParagraph,
    NativeTextSegment,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_document import (
    NativeExamAsset,
    NativeExamDocument,
    NativeExamDocumentOrigin,
)
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_docx_chunker import (  # noqa: E501
    PythonDocxExamChunker,
)
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_docx_shape_figure import (  # noqa: E501
    ShapeFigure,
)
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_docx_shape_reader import (  # noqa: E501
    read_shape_figure,
)
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_docx_writer import (  # noqa: E501
    ExamWorkspaceDocxWriter,
)
from tests.unit.domain.curated_apps.exam_workspace.native_exam_test_builders import (
    make_free_text_item,
)
from tests.unit.infrastructure.curated_apps.exam_workspace_docx_writer_test_support import png

pytestmark = pytest.mark.unit

_NS = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:wpg="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup"'
)
_URI = {
    "wps": "http://schemas.microsoft.com/office/word/2010/wordprocessingShape",
    "wpg": "http://schemas.microsoft.com/office/word/2010/wordprocessingGroup",
}
_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_WPS = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
_ASSET_ID = "asset_0123456789abcdef"
_SOURCE_ACCENT1 = "4472C4"
_EXPORT_ACCENT1 = "4F81BD"  # accent1 of the new document the Word export writes
_BLACK_LINE = '<a:ln w="19050"><a:solidFill><a:srgbClr val="000000"/></a:solidFill></a:ln>'
_STYLE = (
    "<wps:style>"
    '<a:lnRef idx="2"><a:schemeClr val="accent1"><a:shade val="50000"/></a:schemeClr></a:lnRef>'
    '<a:fillRef idx="1"><a:schemeClr val="accent1"/></a:fillRef>'
    '<a:effectRef idx="0"><a:scrgbClr r="0" g="0" b="0"/></a:effectRef>'
    '<a:fontRef idx="minor"><a:schemeClr val="lt1"/></a:fontRef>'
    "</wps:style>"
)


def _source_theme() -> bytes:
    """An Office 2013 theme whose accent1 is 4472C4 (the export's default theme has 4F81BD)."""
    return (
        '<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" name="T">'
        "<a:themeElements>"
        '<a:clrScheme name="T">'
        '<a:dk1><a:sysClr val="windowText" lastClr="000000"/></a:dk1>'
        '<a:lt1><a:sysClr val="window" lastClr="FFFFFF"/></a:lt1>'
        '<a:dk2><a:srgbClr val="44546A"/></a:dk2>'
        '<a:lt2><a:srgbClr val="E7E6E6"/></a:lt2>'
        f'<a:accent1><a:srgbClr val="{_SOURCE_ACCENT1}"/></a:accent1>'
        '<a:accent2><a:srgbClr val="ED7D31"/></a:accent2>'
        "</a:clrScheme>"
        '<a:fmtScheme name="T">'
        '<a:fillStyleLst><a:solidFill><a:schemeClr val="phClr"/></a:solidFill>'
        '<a:solidFill><a:schemeClr val="phClr"/></a:solidFill>'
        '<a:solidFill><a:srgbClr val="123456"/></a:solidFill></a:fillStyleLst>'
        "<a:lnStyleLst>"
        '<a:ln w="6350"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill>'
        '<a:prstDash val="solid"/></a:ln>'
        '<a:ln w="12700"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill>'
        '<a:prstDash val="solid"/></a:ln>'
        '<a:ln w="19050"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:ln>'
        "</a:lnStyleLst>"
        "<a:effectStyleLst>"
        "<a:effectStyle><a:effectLst/></a:effectStyle>"
        "<a:effectStyle><a:effectLst/></a:effectStyle>"
        "<a:effectStyle><a:effectLst/></a:effectStyle>"
        "</a:effectStyleLst>"
        "</a:fmtScheme>"
        "</a:themeElements></a:theme>"
    ).encode()


def _wsp(
    *,
    prst: str = "rect",
    off: tuple[int, int] = (0, 0),
    ext: tuple[int, int] = (1905000, 952500),
    fill: str = "",
    line: str = _BLACK_LINE,
    style: str = "",
    paragraphs: str = "",
) -> str:
    text = f"<wps:txbx><w:txbxContent>{paragraphs}</w:txbxContent></wps:txbx>" if paragraphs else ""
    return (
        '<wps:wsp><wps:cNvPr id="2" name="s"/><wps:cNvSpPr/><wps:spPr>'
        f'<a:xfrm><a:off x="{off[0]}" y="{off[1]}"/><a:ext cx="{ext[0]}" cy="{ext[1]}"/></a:xfrm>'
        f'<a:prstGeom prst="{prst}"><a:avLst/></a:prstGeom>{fill}{line}</wps:spPr>'
        f"{style}{text}<wps:bodyPr/></wps:wsp>"
    )


def _drawing(kind: str, shapes: str, *, ext: tuple[int, int] = (1905000, 952500)) -> str:
    graphic = shapes
    if kind == "wpg":
        graphic = (
            '<wpg:wgp><wpg:cNvPr id="3" name="g"/><wpg:cNvGrpSpPr/><wpg:grpSpPr>'
            f'<a:xfrm><a:off x="0" y="0"/><a:ext cx="{ext[0]}" cy="{ext[1]}"/>'
            f'<a:chOff x="0" y="0"/><a:chExt cx="{ext[0]}" cy="{ext[1]}"/></a:xfrm>'
            f"</wpg:grpSpPr>{shapes}</wpg:wgp>"
        )
    return (
        "<w:drawing><wp:inline>"
        f'<wp:extent cx="{ext[0]}" cy="{ext[1]}"/><wp:docPr id="1" name="d"/>'
        f'<a:graphic><a:graphicData uri="{_URI[kind]}">{graphic}</a:graphicData></a:graphic>'
        "</wp:inline></w:drawing>"
    )


def _solid(inner: str) -> str:
    return f"<a:solidFill>{inner}</a:solidFill>"


def _source_file(drawing: str) -> bytes:
    """A Word file whose second paragraph holds ``drawing``, with the 4472C4 theme and styles."""
    document = Document()
    document.part.part_related_by(RELATIONSHIP_TYPE.THEME)._blob = _source_theme()
    label = document.styles.add_style("Etikett", WD_STYLE_TYPE.PARAGRAPH)
    label.font.size = Pt(14)
    label.font.bold = True
    label.font.color.rgb = RGBColor(0x33, 0x66, 0x99)
    label.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mark = document.styles.add_style("Markering", WD_STYLE_TYPE.CHARACTER)
    mark.font.italic = True
    mark.font.underline = True
    document.add_paragraph("Fråga 1")
    paragraph = document.add_paragraph("")
    paragraph._p.append(parse_xml(f"<w:r {_NS}>{drawing}</w:r>"))
    document.add_paragraph("Fråga 2")
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


class _Round:
    """One drawing from import to export: the picture's model and the exported file's shapes."""

    def __init__(self, drawing: str) -> None:
        self.source_bytes = _source_file(drawing)
        source = Document(io.BytesIO(self.source_bytes))
        (chunk,) = [
            c
            for c in PythonDocxExamChunker().chunk(content=self.source_bytes).chunks
            if c.anchor == "body:0002"
        ]
        (image,) = chunk.images
        assert image.docx_shapes is not None, "the drawing keeps its editable shapes"
        self.stored: str = image.docx_shapes
        self.model: ShapeFigure = read_shape_figure(_only_drawing(source), source)
        self.exported = _export(image.payload, self.stored)
        self.exported_model = read_shape_figure(_only_drawing(self.exported), self.exported)


def _only_drawing(document: DocxDocument) -> etree._Element:
    (drawing,) = document.element.body.iter(qn("w:drawing"))
    return drawing


def _export(payload: bytes, stored: str) -> DocxDocument:
    item = make_free_text_item(sequence=1, item_id="item_001")
    item = item.model_copy(
        update={
            "body": (
                NativeParagraph(
                    heading_level=None,
                    segments=(
                        NativeTextSegment(text="Se bild", marks=()),
                        NativeAssetSegment(asset_id=_ASSET_ID),
                    ),
                ),
            )
        }
    )
    document = NativeExamDocument(
        document_id=uuid4(),
        revision=1,
        title="Bildprov",
        items=(item,),
        assets=(
            NativeExamAsset(
                asset_id=_ASSET_ID,
                sha256="0" * 64,
                media_type="image/png",
                byte_length=len(payload),
                docx_shapes=stored,
            ),
        ),
        origin=NativeExamDocumentOrigin(kind="created"),
    )
    content = ExamWorkspaceDocxWriter().build_docx_bytes(
        document, assets_by_id={_ASSET_ID: png((20, 90, 20, 255), (40, 20))}
    )
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        assert "word/media/" not in "".join(archive.namelist()), "the shapes, not the picture"
    return Document(io.BytesIO(content))


def _stored_root(trip: _Round) -> etree._Element:
    return etree.fromstring(trip.stored.encode())


def _local_names(element: etree._Element | None) -> list[str]:
    assert element is not None
    return [etree.QName(child).localname for child in element]


def _assert_self_contained(stored: str) -> None:
    for reference in ("schemeClr", "<wps:style", "pStyle", "rStyle"):
        assert reference not in stored, f"{reference} is left in the stored shapes"


@pytest.mark.parametrize(
    "colour",
    [
        '<a:schemeClr val="accent1"/>',
        '<a:schemeClr val="accent1"><a:lumMod val="75000"/></a:schemeClr>',
        '<a:schemeClr val="accent1"><a:lumMod val="40000"/><a:lumOff val="60000"/></a:schemeClr>',
        '<a:schemeClr val="accent1"><a:shade val="50000"/></a:schemeClr>',
        '<a:schemeClr val="accent1"><a:tint val="40000"/></a:schemeClr>',
        '<a:schemeClr val="tx2"/>',
    ],
)
def test_a_scheme_colour_fill_exports_as_the_colour_the_picture_drew(colour: str) -> None:
    trip = _Round(_drawing("wps", _wsp(fill=_solid(colour))))

    _assert_self_contained(trip.stored)
    (shape,) = trip.model.shapes
    assert shape.fill is not None
    (exported,) = trip.exported_model.shapes
    assert exported.fill == shape.fill


def test_an_accent1_rectangle_keeps_the_source_accent_under_the_export_theme() -> None:
    trip = _Round(_drawing("wps", _wsp(fill=_solid('<a:schemeClr val="accent1"/>'))))

    assert trip.model.shapes[0].fill == _SOURCE_ACCENT1
    assert trip.exported_model.shapes[0].fill == _SOURCE_ACCENT1 != _EXPORT_ACCENT1


def test_a_shape_styled_only_through_wps_style_exports_its_resolved_fill_and_outline() -> None:
    trip = _Round(_drawing("wps", _wsp(line="", style=_STYLE)))

    _assert_self_contained(trip.stored)
    (shape,) = trip.model.shapes
    assert shape.fill == _SOURCE_ACCENT1
    assert shape.outline is not None and shape.outline.color == "2F528F"
    (exported,) = trip.exported_model.shapes
    assert (exported.fill, exported.outline) == (shape.fill, shape.outline)


def test_a_style_outline_with_no_fill_of_its_own_keeps_no_fill_and_the_style_outline() -> None:
    trip = _Round(_drawing("wps", _wsp(fill="<a:noFill/>", line="", style=_STYLE)))

    _assert_self_contained(trip.stored)
    (exported,) = trip.exported_model.shapes
    assert exported.fill is None
    assert exported.outline == trip.model.shapes[0].outline


def test_a_label_citing_source_only_styles_exports_with_explicit_marks() -> None:
    paragraph = (
        '<w:p><w:pPr><w:pStyle w:val="Etikett"/></w:pPr>'
        '<w:r><w:rPr><w:rStyle w:val="Markering"/></w:rPr><w:t>Hej</w:t></w:r>'
        "<w:r><w:t> där</w:t></w:r></w:p>"
    )
    trip = _Round(_drawing("wps", _wsp(paragraphs=paragraph)))

    _assert_self_contained(trip.stored)
    label = trip.model.shapes[0].label
    assert label is not None
    first, second = label.paragraphs[0].runs
    assert (first.size_pt, first.color, first.bold) == (14.0, "336699", True)
    assert (first.italic, first.underline, second.italic) == (True, True, False)
    assert label.paragraphs[0].align == "center"
    assert trip.exported_model.shapes[0].label == label
    exported_paragraph = trip.exported.element.body.find(f".//{qn('w:txbxContent')}/{qn('w:p')}")
    assert exported_paragraph is not None
    assert exported_paragraph.find(f".//{qn('w:pStyle')}") is None
    assert exported_paragraph.find(f".//{qn('w:rStyle')}") is None


def test_a_label_run_set_to_not_bold_stores_an_explicit_not_bold_under_a_bold_style() -> None:
    paragraph = (
        '<w:p><w:pPr><w:pStyle w:val="Etikett"/></w:pPr>'
        '<w:r><w:rPr><w:b w:val="0"/></w:rPr><w:t>Hej</w:t></w:r></w:p>'
    )
    trip = _Round(_drawing("wps", _wsp(paragraphs=paragraph)))

    (bold,) = _stored_root(trip).iter(f"{{{_W}}}b")
    assert bold.get(f"{{{_W}}}val") == "0"
    label = trip.exported_model.shapes[0].label
    assert label is not None and label.paragraphs[0].runs[0].bold is False


def test_a_run_colour_with_a_theme_colour_exports_its_resolved_value_only() -> None:
    paragraph = (
        '<w:p><w:r><w:rPr><w:color w:val="336699" w:themeColor="accent1" w:themeShade="BF"/>'
        "</w:rPr><w:t>Hej</w:t></w:r></w:p>"
    )
    trip = _Round(_drawing("wps", _wsp(paragraphs=paragraph)))

    (color,) = _stored_root(trip).iter(f"{{{_W}}}color")
    assert dict(color.attrib) == {f"{{{_W}}}val": "336699"}


def test_the_stored_shapes_keep_the_schema_order_of_the_elements_they_rewrite() -> None:
    paragraph = (
        '<w:p><w:pPr><w:pStyle w:val="Etikett"/><w:spacing w:after="0"/><w:rPr/></w:pPr>'
        '<w:r><w:rPr><w:rStyle w:val="Markering"/><w:rFonts w:ascii="Arial"/></w:rPr>'
        "<w:t>Hej</w:t></w:r></w:p>"
    )
    line = (
        '<a:ln w="9525"><a:solidFill><a:schemeClr val="accent1"/></a:solidFill>'
        '<a:round/><a:headEnd type="triangle"/></a:ln>'
    )
    trip = _Round(_drawing("wps", _wsp(fill="", line=line, style=_STYLE, paragraphs=paragraph)))

    root = _stored_root(trip)
    assert _local_names(root.find(f".//{{{_WPS}}}spPr")) == ["xfrm", "prstGeom", "solidFill", "ln"]
    assert _local_names(root.find(f".//{{{_A}}}ln")) == ["solidFill", "round", "headEnd"]
    assert _local_names(root.find(f".//{{{_W}}}pPr")) == ["spacing", "jc", "rPr"]
    assert _local_names(root.find(f".//{{{_W}}}r/{{{_W}}}rPr")) == [
        "rFonts", "b", "i", "color", "sz", "u",
    ]  # fmt: skip
    assert _local_names(root.find(f".//{{{_WPS}}}wsp")) == [
        "cNvPr",
        "cNvSpPr",
        "spPr",
        "txbx",
        "bodyPr",
    ]


def test_a_dashed_arrow_outline_exports_with_the_dash_and_both_arrowheads() -> None:
    line = (
        '<a:ln w="12700"><a:solidFill><a:schemeClr val="accent2"/></a:solidFill>'
        '<a:prstDash val="dash"/><a:headEnd type="oval" w="sm" len="lg"/>'
        '<a:tailEnd type="triangle"/></a:ln>'
    )
    trip = _Round(_drawing("wps", _wsp(prst="straightConnector1", fill="", line=line)))

    (shape,) = trip.model.shapes
    assert shape.outline is not None and shape.outline.color == "ED7D31"
    assert trip.exported_model.shapes == trip.model.shapes


def test_a_label_without_any_style_gets_its_resolved_size_colour_and_alignment_explicit() -> None:
    paragraph = "<w:p><w:r><w:t>Hej</w:t></w:r></w:p>"
    trip = _Round(
        _drawing("wps", _wsp(fill=_solid('<a:schemeClr val="accent1"/>'), paragraphs=paragraph))
    )

    assert trip.exported_model.shapes[0].label == trip.model.shapes[0].label


def test_a_label_colour_from_the_font_reference_exports_as_an_explicit_run_colour() -> None:
    paragraph = "<w:p><w:r><w:t>Hej</w:t></w:r></w:p>"
    trip = _Round(_drawing("wps", _wsp(line="", style=_STYLE, paragraphs=paragraph)))

    _assert_self_contained(trip.stored)
    label = trip.model.shapes[0].label
    assert label is not None and label.paragraphs[0].runs[0].color == "FFFFFF"
    assert trip.exported_model.shapes[0].label == label


def test_every_shape_of_a_group_exports_its_own_resolved_values_in_paint_order() -> None:
    first = _wsp(
        fill=_solid('<a:schemeClr val="accent1"/>'), paragraphs="<w:p><w:r><w:t>A</w:t></w:r></w:p>"
    )
    second = _wsp(
        prst="ellipse",
        off=(950000, 0),
        ext=(900000, 900000),
        fill=_solid('<a:schemeClr val="accent2"/>'),
        line="",
        style=_STYLE,
    )
    trip = _Round(_drawing("wpg", first + second))

    _assert_self_contained(trip.stored)
    assert [s.fill for s in trip.model.shapes] == [_SOURCE_ACCENT1, "ED7D31"]
    assert trip.exported_model.shapes == trip.model.shapes


def test_a_scheme_colour_in_an_ignorable_extension_list_is_resolved_too() -> None:
    hidden = (
        "<a:extLst><a:ext uri='{909E8E84-426E-40DD-AFC4-6F175D3DCCD1}'>"
        '<a14:hiddenFill xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main">'
        '<a:solidFill><a:schemeClr val="accent1"/></a:solidFill></a14:hiddenFill>'
        "</a:ext></a:extLst>"
    )
    drawing = _drawing("wps", _wsp(fill=_solid('<a:schemeClr val="accent1"/>'), line=hidden))

    trip = _Round(drawing)

    assert "schemeClr" not in trip.stored
