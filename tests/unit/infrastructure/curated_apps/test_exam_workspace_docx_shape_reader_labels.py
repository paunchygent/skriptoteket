"""Shape reader: text box labels, runs, styles and body properties (TASK-SKRIPT-39-04-22 unit 1)."""

from __future__ import annotations

import pytest
from docx.oxml import parse_xml
from docx.oxml.ns import qn

from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_docx_shape_figure import (  # noqa: E501
    LabelParagraph,
    LabelRun,
    ShapeLabel,
)
from tests.unit.infrastructure.curated_apps.exam_workspace_docx_shape_reader_test_support import (
    NS,
    label_xml,
    make_document,
    para_xml,
    read_one,
    shape_text,
    style_xml,
    theme_xml,
    wsp_xml,
)

pytestmark = pytest.mark.unit


# --- labels ----------------------------------------------------------------------------------


def test_label_runs_marks_size_colour_alignment_and_breaks() -> None:
    content = (
        '<w:p><w:pPr><w:jc w:val="end"/></w:pPr>'
        '<w:r><w:rPr><w:b/><w:i w:val="1"/><w:u w:val="single"/><w:sz w:val="28"/>'
        '<w:color w:val="ff0000"/></w:rPr><w:t>Hej</w:t><w:tab/><w:t>du</w:t><w:br/>'
        "<w:t>nu</w:t></w:r>"
        '<w:r><w:rPr><w:b w:val="0"/><w:u w:val="none"/></w:rPr>'
        '<w:t xml:space="preserve"> x</w:t></w:r>'
        "</w:p>"
    )

    shape = read_one(wsp_xml(**label_xml(content)))

    assert shape.label == ShapeLabel(
        paragraphs=(
            LabelParagraph(
                runs=(
                    LabelRun("Hej du\nnu", 14.0, "FF0000", bold=True, italic=True, underline=True),
                    LabelRun(" x", 11.0, "000000"),
                ),
                align="right",
            ),
        ),
    )


@pytest.mark.parametrize(
    ("jc", "expected"),
    [("left", "left"), ("start", "left"), ("center", "center"), ("right", "right"),
     ("end", "right"), ("both", "justify"), (None, "left")],
)  # fmt: skip
def test_paragraph_alignment_values(jc: str | None, expected: str) -> None:
    shape = read_one(wsp_xml(**label_xml(para_xml("t", jc=jc))))

    assert shape.label is not None
    assert shape.label.paragraphs[0].align == expected


def test_paragraph_style_alignment_and_size_apply_when_the_paragraph_sets_none() -> None:
    document = make_document()
    style = document.styles.add_style("BoxText", 1)
    style.element.get_or_add_pPr().append(parse_xml(f'<w:jc {NS} w:val="center"/>'))
    style.element.get_or_add_rPr().append(parse_xml(f'<w:sz {NS} w:val="36"/>'))

    shape = read_one(
        wsp_xml(**label_xml(para_xml("t", ppr='<w:pStyle w:val="BoxText"/>'))), document=document
    )

    assert shape.label is not None
    assert shape.label.paragraphs[0].align == "center"
    assert shape.label.paragraphs[0].runs[0].size_pt == 18.0


def test_a_style_chain_supplies_bold_from_the_paragraph_style() -> None:
    document = make_document()
    style = document.styles.add_style("Strong box", 1)
    style.element.get_or_add_rPr().append(parse_xml(f"<w:b {NS}/>"))

    shape = read_one(
        wsp_xml(**label_xml(para_xml("t", ppr='<w:pStyle w:val="Strongbox"/>'))), document=document
    )

    assert shape.label is not None
    assert shape.label.paragraphs[0].runs[0].bold is True


def test_size_falls_back_to_the_document_default() -> None:
    shape = read_one(wsp_xml(**label_xml(para_xml("t"))))

    assert shape.label is not None
    assert shape.label.paragraphs[0].runs[0].size_pt == 11.0  # python-docx template: sz 22


def test_size_is_ten_points_when_nothing_sets_one() -> None:
    document = make_document()
    defaults = document.styles.element.find(qn("w:docDefaults"))
    assert defaults is not None
    document.styles.element.remove(defaults)

    shape = read_one(wsp_xml(**label_xml(para_xml("t"))), document=document)

    assert shape.label is not None
    assert shape.label.paragraphs[0].runs[0].size_pt == 10.0


def test_auto_colour_takes_the_style_font_colour_else_black() -> None:
    styled = read_one(
        wsp_xml(
            style=style_xml(ln=0, fill=0), **label_xml(para_xml("t", rpr='<w:color w:val="auto"/>'))
        ),
        document=make_document(theme_xml()),
    )
    plain = read_one(wsp_xml(**label_xml(para_xml("t"))))

    assert styled.label is not None and plain.label is not None
    assert styled.label.paragraphs[0].runs[0].color == "FFFFFF"
    assert plain.label.paragraphs[0].runs[0].color == "000000"


def test_body_properties_read_insets_and_anchor() -> None:
    body = '<wps:bodyPr lIns="1" tIns="2" rIns="3" bIns="4" anchor="b" wrap="square" rtlCol="0"/>'

    shape = read_one(wsp_xml(**label_xml(para_xml("t"), body=body)))

    assert shape.label is not None
    assert (
        shape.label.inset_left_emu,
        shape.label.inset_top_emu,
        shape.label.inset_right_emu,
        shape.label.inset_bottom_emu,
        shape.label.anchor,
    ) == (1, 2, 3, 4, "b")


def test_body_properties_default_to_word_insets_and_top_anchor() -> None:
    shape = read_one(wsp_xml(**label_xml(para_xml("t"))))

    assert shape.label is not None
    assert (
        shape.label.inset_left_emu,
        shape.label.inset_top_emu,
        shape.label.inset_right_emu,
        shape.label.inset_bottom_emu,
        shape.label.anchor,
    ) == (91440, 45720, 91440, 45720, "t")


def test_two_paragraphs_keep_their_own_alignment_and_an_empty_one_stays() -> None:
    shape = read_one(
        wsp_xml(**label_xml(para_xml("a", jc="center"), "<w:p/>", para_xml("b", jc="right")))
    )

    assert shape.label is not None
    assert [(p.align, [r.text for r in p.runs]) for p in shape.label.paragraphs] == [
        ("center", ["a"]),
        ("left", []),
        ("right", ["b"]),
    ]


def test_a_label_without_any_text_is_no_label() -> None:
    assert (
        read_one(wsp_xml(**label_xml("<w:p/>", "<w:p><w:r><w:t></w:t></w:r></w:p>"))).label is None
    )


def test_bookmarks_and_proof_marks_around_runs_are_ignored() -> None:
    content = (
        '<w:p><w:bookmarkStart w:id="0" w:name="b"/><w:proofErr w:type="spellStart"/>'
        '<w:r><w:t>ok</w:t></w:r><w:bookmarkEnd w:id="0"/></w:p>'
    )

    shape = read_one(wsp_xml(**label_xml(content)))

    assert shape_text(shape) == "ok"
