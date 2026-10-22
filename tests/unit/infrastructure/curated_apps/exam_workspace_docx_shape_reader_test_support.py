"""Word DrawingML builders shared by the shape reader tests (TASK-SKRIPT-39-04-22 unit 1).

The two fixture drawings are real Word output from ``metodprov_med_bilder.docx``; every other
case is hand-written DrawingML read through a python-docx document, so the theme, the style
chain and the document defaults are the ones python-docx exposes. Each rule family of D2 has an
unsupported case: the reader must raise ``UnsupportedShapeFigure`` and never return a partial
figure.
"""

from __future__ import annotations

from pathlib import Path
from typing import TypedDict

import pytest
from docx import Document
from docx.document import Document as DocxDocument
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import parse_xml
from docx.oxml.xmlchemy import BaseOxmlElement

from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_docx_shape_figure import (  # noqa: E501
    FigureShape,
    UnsupportedShapeFigure,
)
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_docx_shape_reader import (  # noqa: E501
    read_shape_figure,
)

FIXTURE = (
    Path(__file__).resolve().parents[3]
    / "fixtures"
    / "exam_conversion"
    / "pdf_corpus"
    / "metodprov_med_bilder.docx"
)
NS = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
    'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:wpg="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup" '
    'xmlns:wpc="http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas" '
    'xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main"'
)
_URI = {
    "wps": "http://schemas.microsoft.com/office/word/2010/wordprocessingShape",
    "wpg": "http://schemas.microsoft.com/office/word/2010/wordprocessingGroup",
    "wpc": "http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas",
}
BLACK_LINE = '<a:ln w="19050"><a:solidFill><a:srgbClr val="000000"/></a:solidFill></a:ln>'


def theme_xml(accent1: str = "4472C4", effect_third: str = "<a:effectLst/>") -> bytes:
    """A theme in the Office 2013 layout: the colour scheme and the three style lists."""
    return (
        '<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" name="T">'
        "<a:themeElements>"
        '<a:clrScheme name="T">'
        '<a:dk1><a:sysClr val="windowText" lastClr="000000"/></a:dk1>'
        '<a:lt1><a:sysClr val="window" lastClr="FFFFFF"/></a:lt1>'
        '<a:dk2><a:srgbClr val="44546A"/></a:dk2>'
        '<a:lt2><a:srgbClr val="E7E6E6"/></a:lt2>'
        f'<a:accent1><a:srgbClr val="{accent1}"/></a:accent1>'
        '<a:accent2><a:srgbClr val="ED7D31"/></a:accent2>'
        "</a:clrScheme>"
        '<a:fmtScheme name="T">'
        '<a:fillStyleLst><a:solidFill><a:schemeClr val="phClr"/></a:solidFill>'
        '<a:gradFill><a:gsLst><a:gs pos="0"><a:schemeClr val="phClr"/></a:gs>'
        "</a:gsLst></a:gradFill>"
        '<a:solidFill><a:srgbClr val="123456"/></a:solidFill></a:fillStyleLst>'
        "<a:lnStyleLst>"
        '<a:ln w="6350"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill>'
        '<a:prstDash val="solid"/></a:ln>'
        '<a:ln w="12700"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill>'
        '<a:prstDash val="solid"/></a:ln>'
        '<a:ln w="19050"><a:gradFill><a:gsLst/></a:gradFill></a:ln>'
        "</a:lnStyleLst>"
        "<a:effectStyleLst>"
        "<a:effectStyle><a:effectLst/></a:effectStyle>"
        "<a:effectStyle><a:effectLst/></a:effectStyle>"
        f"<a:effectStyle>{effect_third}</a:effectStyle>"
        "</a:effectStyleLst>"
        "</a:fmtScheme>"
        "</a:themeElements></a:theme>"
    ).encode()


def make_document(theme: bytes | None = None, *, drop_theme: bool = False) -> DocxDocument:
    document = Document()
    if drop_theme:
        rel_id = next(
            key for key, rel in document.part.rels.items() if rel.reltype == RELATIONSHIP_TYPE.THEME
        )
        document.part.rels.pop(rel_id)
    elif theme is not None:
        document.part.part_related_by(RELATIONSHIP_TYPE.THEME)._blob = theme
    return document


def make_drawing(
    graphic: str, *, uri: str = "wps", cx: int = 1000, cy: int = 500, anchor: bool = False
) -> BaseOxmlElement:
    container = "wp:anchor" if anchor else "wp:inline"
    data_uri = _URI.get(uri, uri)
    return parse_xml(
        f"<w:drawing {NS}><{container}>"
        f'<wp:extent cx="{cx}" cy="{cy}"/><wp:docPr id="1" name="d"/>'
        f'<a:graphic><a:graphicData uri="{data_uri}">{graphic}</a:graphicData></a:graphic>'
        f"</{container}></w:drawing>"
    )


def wsp_xml(
    *,
    prst: str = "rect",
    off: tuple[int, int] = (0, 0),
    ext: tuple[int, int] = (1000, 500),
    xfrm_attrs: str = "",
    geometry: str | None = None,
    guides: str = "",
    fill: str = "<a:noFill/>",
    line: str = BLACK_LINE,
    sp_extra: str = "",
    style: str = "",
    txbx: str = "",
    body: str = "<wps:bodyPr/>",
    connector: bool = False,
) -> str:
    geometry = geometry or f'<a:prstGeom prst="{prst}"><a:avLst>{guides}</a:avLst></a:prstGeom>'
    non_visual = "<wps:cNvCnPr/>" if connector else "<wps:cNvSpPr/>"
    return (
        f'<wps:wsp><wps:cNvPr id="2" name="s"/>{non_visual}<wps:spPr>'
        f'<a:xfrm {xfrm_attrs}><a:off x="{off[0]}" y="{off[1]}"/>'
        f'<a:ext cx="{ext[0]}" cy="{ext[1]}"/></a:xfrm>'
        f"{geometry}{fill}{line}{sp_extra}</wps:spPr>{style}{txbx}{body}</wps:wsp>"
    )


def group_xml(
    children: str,
    *,
    tag: str = "wpg:wgp",
    off: tuple[int, int] = (0, 0),
    ext: tuple[int, int] = (1000, 500),
    ch_off: tuple[int, int] = (0, 0),
    ch_ext: tuple[int, int] = (1000, 500),
    xfrm_attrs: str = "",
) -> str:
    return (
        f'<{tag}><wpg:cNvPr id="3" name="g"/><wpg:cNvGrpSpPr/><wpg:grpSpPr>'
        f'<a:xfrm {xfrm_attrs}><a:off x="{off[0]}" y="{off[1]}"/>'
        f'<a:ext cx="{ext[0]}" cy="{ext[1]}"/>'
        f'<a:chOff x="{ch_off[0]}" y="{ch_off[1]}"/>'
        f'<a:chExt cx="{ch_ext[0]}" cy="{ch_ext[1]}"/></a:xfrm></wpg:grpSpPr>'
        f"{children}</{tag}>"
    )


def canvas_xml(children: str, *, background: str = "<wpc:bg><a:noFill/></wpc:bg>") -> str:
    return (
        f"<wpc:wpc>{background}<wpc:whole><a:ln><a:noFill/></a:ln></wpc:whole>{children}</wpc:wpc>"
    )


class LabelXml(TypedDict):
    txbx: str
    body: str


def label_xml(*paragraphs: str, body: str = "<wps:bodyPr/>") -> LabelXml:
    content = "".join(paragraphs)
    return {"txbx": f"<wps:txbx><w:txbxContent>{content}</w:txbxContent></wps:txbx>", "body": body}


def para_xml(text: str, *, jc: str | None = None, rpr: str = "", ppr: str = "") -> str:
    jc_xml = f'<w:jc w:val="{jc}"/>' if jc else ""
    ppr_xml = f"<w:pPr>{ppr}{jc_xml}</w:pPr>" if ppr or jc else ""
    rpr_xml = f"<w:rPr>{rpr}</w:rPr>" if rpr else ""
    return f"<w:p>{ppr_xml}<w:r>{rpr_xml}<w:t>{text}</w:t></w:r></w:p>"


def read_one(shape_xml: str, *, document: DocxDocument | None = None) -> FigureShape:
    figure = read_shape_figure(make_drawing(shape_xml), document or make_document())
    assert len(figure.shapes) == 1
    return figure.shapes[0]


def rejected(drawing: BaseOxmlElement, needle: str, document: DocxDocument | None = None) -> None:
    with pytest.raises(UnsupportedShapeFigure) as caught:
        read_shape_figure(drawing, document or make_document())
    assert needle in caught.value.reason


def rejected_shape(shape_xml: str, needle: str, document: DocxDocument | None = None) -> None:
    rejected(make_drawing(shape_xml), needle, document)


def shape_text(shape: FigureShape) -> str | None:
    if shape.label is None:
        return None
    return "".join(run.text for p in shape.label.paragraphs for run in p.runs)


def scheme_fill(inner: str, scheme: str = "accent1") -> str:
    return f'<a:solidFill><a:schemeClr val="{scheme}">{inner}</a:schemeClr></a:solidFill>'


_OFFICE_STYLE = (
    "<wps:style>"
    '<a:lnRef idx="{ln}"><a:schemeClr val="accent1"><a:shade val="50000"/></a:schemeClr></a:lnRef>'
    '<a:fillRef idx="{fill}"><a:schemeClr val="accent1"/></a:fillRef>'
    '<a:effectRef idx="{effect}"><a:scrgbClr r="0" g="0" b="0"/></a:effectRef>'
    '<a:fontRef idx="minor"><a:schemeClr val="lt1"/></a:fontRef>'
    "</wps:style>"
)


def style_xml(ln: int = 2, fill: int = 1, effect: int = 0) -> str:
    return _OFFICE_STYLE.format(ln=ln, fill=fill, effect=effect)


__all__ = [
    "FIXTURE",
    "NS",
    "BLACK_LINE",
    "LabelXml",
    "canvas_xml",
    "group_xml",
    "label_xml",
    "make_document",
    "make_drawing",
    "para_xml",
    "read_one",
    "rejected",
    "rejected_shape",
    "scheme_fill",
    "shape_text",
    "style_xml",
    "theme_xml",
    "wsp_xml",
]
