"""Answer-form tests for the exam workspace DOCX writer (W4, W6-W9).

The student copy's writing area, gap rules, tick forms and parent instructions are read
back from the produced bytes, pinning the printed answer-form geometry independently of
the content/house-template suite.
"""

import io

import pytest
from docx import Document
from docx.enum.text import WD_LINE_SPACING

from skriptoteket.domain.curated_apps.exam_workspace.native_exam_body import (
    NativeParagraph,
    NativeTextSegment,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_vocabulary import (
    NativeAnswerKeyOrigin,
)
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub import (
    exam_workspace_docx_writer,
)
from tests.unit.domain.curated_apps.exam_workspace.native_exam_test_builders import (
    make_document,
    make_ordering_item,
    make_single_choice_item,
)
from tests.unit.domain.curated_apps.exam_workspace.native_matching_test_support import (
    matching_item,
)
from tests.unit.domain.curated_apps.exam_workspace.recorded_native_exam_support import (
    parent_over,
    recorded_document,
    recorded_items,
)
from tests.unit.infrastructure.curated_apps.exam_workspace_docx_writer_test_support import (
    NBSP,
    STEP_BOX,
    blank_run_count,
    border_attributes,
    border_edges,
    document_text,
    free_text_item,
    heading_texts,
    lone_gap_item,
    only_paragraph,
    ruled_paragraphs,
    run_border,
    sentence_gap_item,
)

ExamWorkspaceDocxWriter = exam_workspace_docx_writer.ExamWorkspaceDocxWriter


class TestWritingArea:
    @pytest.mark.parametrize("count", [5, 1])
    def test_free_text_prints_exactly_its_writing_lines(self, count: int) -> None:
        document = make_document(free_text_item(writing_lines=count))
        parsed = Document(io.BytesIO(ExamWorkspaceDocxWriter().build_docx_bytes(document)))
        paragraphs = list(parsed.paragraphs)

        stem_at = next(
            index
            for index, paragraph in enumerate(paragraphs)
            if paragraph.text == "Resonera om renässansen. (4 p)"
        )
        # A free-text question ends with its writing area, so the stem's tail is it.
        ruled = paragraphs[stem_at + 1 :]
        assert len(ruled) == count
        for paragraph in ruled:
            assert paragraph.text == ""
            assert border_edges(paragraph) == ("bottom", "between")
            expected = {"val": "single", "sz": "4", "space": "0", "color": "555555"}
            assert border_attributes(paragraph, "bottom") == expected
            assert border_attributes(paragraph, "between") == expected
            spacing = paragraph.paragraph_format
            assert spacing.line_spacing.mm == pytest.approx(10, abs=0.01)
            assert spacing.line_spacing_rule is WD_LINE_SPACING.EXACTLY


class TestGapRules:
    def test_a_lone_gap_is_one_ruled_line_and_no_gap_marker(self) -> None:
        parsed = Document(
            io.BytesIO(ExamWorkspaceDocxWriter().build_docx_bytes(make_document(lone_gap_item())))
        )

        assert len(ruled_paragraphs(list(parsed.paragraphs))) == 1
        assert "[___]" not in document_text(parsed)

    @pytest.mark.parametrize(
        ("accepted", "origin", "expected_count"),
        [
            (("ab", "fotosyntes"), NativeAnswerKeyOrigin.SOURCE_PROVIDED, 41),
            (("nominalfras",), NativeAnswerKeyOrigin.SOURCE_PROVIDED, 45),
            ((), NativeAnswerKeyOrigin.ABSENT, 26),
        ],
    )
    def test_a_sentence_gap_is_an_underlined_blank_sized_for_its_answer(
        self,
        accepted: tuple[str, ...],
        origin: NativeAnswerKeyOrigin,
        expected_count: int,
    ) -> None:
        document = make_document(sentence_gap_item(accepted, origin=origin))
        parsed = Document(io.BytesIO(ExamWorkspaceDocxWriter().build_docx_bytes(document)))

        paragraph = only_paragraph(
            list(parsed.paragraphs), "Svar: " + NBSP * expected_count + " (2 p)"
        )
        assert paragraph.runs[0].text == "Svar: "
        assert blank_run_count(paragraph) == expected_count
        assert paragraph.paragraph_format.line_spacing.mm == pytest.approx(10, abs=0.01)
        assert paragraph.paragraph_format.line_spacing_rule is WD_LINE_SPACING.EXACTLY
        for value in accepted:
            assert value not in document_text(parsed)


class TestForms:
    def test_single_choice_prints_tick_options(self) -> None:
        document = make_document(make_single_choice_item())
        parsed = Document(io.BytesIO(ExamWorkspaceDocxWriter().build_docx_bytes(document)))
        texts = [paragraph.text for paragraph in parsed.paragraphs]

        assert f"{STEP_BOX} A) vägen" in texts
        assert f"{STEP_BOX} B) eleverna" in texts
        assert "Facit" not in document_text(parsed)

    def test_ordering_prints_number_boxes_in_print_order(self) -> None:
        document = make_document(make_ordering_item())
        parsed = Document(io.BytesIO(ExamWorkspaceDocxWriter().build_docx_bytes(document)))

        boxed = [
            paragraph
            for paragraph in parsed.paragraphs
            if paragraph.runs and run_border(paragraph.runs[0]) is not None
        ]
        assert [paragraph.runs[-1].text for paragraph in boxed] == [
            "Avslutning",
            "Motargument",
            "Argument",
            "Tes",
        ]
        assert STEP_BOX not in document_text(parsed)
        assert run_border(boxed[0].runs[0]) == {
            "val": "single",
            "sz": "8",
            "space": "0",
            "color": "111111",
        }

    def test_matching_prints_a_grid_with_tick_cells(self) -> None:
        document = make_document(matching_item())
        parsed = Document(io.BytesIO(ExamWorkspaceDocxWriter().build_docx_bytes(document)))

        assert len(parsed.tables) == 1
        table = parsed.tables[0]
        assert len(table.rows) == 4
        assert len(table.columns) == 4
        assert [cell.text for cell in table.rows[0].cells] == ["", "däggdjur", "fåglar", "fiskar"]
        assert [cell.text for cell in table.rows[1].cells] == [
            "hund",
            STEP_BOX,
            STEP_BOX,
            STEP_BOX,
        ]


class TestMarkedAnswerForms:
    def test_single_choice_option_keeps_a_bold_word(self) -> None:
        item = make_single_choice_item()
        marked = item.choices[0].model_copy(
            update={
                "body": (
                    NativeParagraph(
                        heading_level=None,
                        segments=(
                            NativeTextSegment(text="stark", marks=("bold",)),
                            NativeTextSegment(text=" vind", marks=()),
                        ),
                    ),
                )
            }
        )
        item = item.model_copy(update={"choices": (marked, item.choices[1])})
        parsed = Document(
            io.BytesIO(ExamWorkspaceDocxWriter().build_docx_bytes(make_document(item)))
        )
        paragraph = only_paragraph(list(parsed.paragraphs), f"{STEP_BOX} A) stark vind")
        by_text = {run.text: run for run in paragraph.runs}

        assert by_text["stark"].bold is True
        assert by_text[" vind"].bold is None

    def test_matching_grid_cell_keeps_an_underlined_word(self) -> None:
        item = matching_item()
        marked = item.match_right[0].model_copy(
            update={
                "body": (
                    NativeParagraph(
                        heading_level=None,
                        segments=(
                            NativeTextSegment(text="dägg", marks=("underline",)),
                            NativeTextSegment(text="djur", marks=()),
                        ),
                    ),
                )
            }
        )
        item = item.model_copy(update={"match_right": (marked, *item.match_right[1:])})
        parsed = Document(
            io.BytesIO(ExamWorkspaceDocxWriter().build_docx_bytes(make_document(item)))
        )
        cell = parsed.tables[0].rows[0].cells[1]
        assert cell.text == "däggdjur"
        by_text = {run.text: run for paragraph in cell.paragraphs for run in paragraph.runs}

        assert by_text["dägg"].underline is True
        assert by_text["djur"].underline is None


class TestParents:
    def test_parent_instruction_is_written_once_before_its_first_child_heading(self) -> None:
        items = recorded_items()
        first, second, third = items[0], items[1], items[2]
        assert len({first.kind, second.kind, third.kind}) == 1
        parent = parent_over("parent_001", "Shared instruction for the run.", 1, 2, 3)

        content = ExamWorkspaceDocxWriter().build_docx_bytes(
            recorded_document((parent,), items=items)
        )

        parsed = Document(io.BytesIO(content))
        texts = [paragraph.text for paragraph in parsed.paragraphs]
        assert texts.count("Shared instruction for the run.") == 1
        headings = heading_texts(parsed, "Heading 3")
        instruction_at = texts.index("Shared instruction for the run.")
        first_heading_at = texts.index(headings[0])
        assert instruction_at == first_heading_at - 1
        assert len(headings) == len(items)

    def test_document_without_parents_writes_no_extra_paragraphs(self) -> None:
        items = recorded_items()
        plain = ExamWorkspaceDocxWriter().build_docx_bytes(recorded_document(items=items))
        parent = parent_over("parent_001", "Shared instruction for the run.", 1, 2)
        with_parent = ExamWorkspaceDocxWriter().build_docx_bytes(
            recorded_document((parent,), items=items)
        )

        plain_count = len(Document(io.BytesIO(plain)).paragraphs)
        assert len(Document(io.BytesIO(with_parent)).paragraphs) == plain_count + 1
