"""Group rotation and flips in the Word shape reader (TASK-SKRIPT-39-04-22 unit 1, D1, D2).

Every case is an asymmetric group of two supported shapes, so a wrong rotation direction, a flip
applied after the rotation or a lost label moves a hand-computed box centre. All numbers are in
EMU. The common group is 1000 x 400 at the origin with its child space equal to its own space;
the rectangle "A" (200 x 100, centre (200, 100), rotated 30 degrees) and the ellipse "B"
(300 x 100, centre (750, 300)) sit at different positions inside it. Word sizes ``wp:extent`` to
the rotated bounding box, so the rotated figure is 400 x 1000 and the group box is centred in it.

DrawingML order for a group: map the child into the group box (scale), mirror inside the box
(``flipH``, ``flipV``), then rotate clockwise about the box centre. A shape's own flips mirror it
before its own rotation, so a mirrored group negates the child rotation and toggles its flip.
"""

from __future__ import annotations

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.xmlchemy import BaseOxmlElement

from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_docx_shape_figure import (  # noqa: E501
    FigureShape,
    ShapeFigure,
    UnsupportedShapeFigure,
)
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_docx_shape_reader import (  # noqa: E501
    read_shape_figure,
)

pytestmark = pytest.mark.unit

_NS = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:wpg="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup"'
)
_GROUP_URI = "http://schemas.microsoft.com/office/word/2010/wordprocessingGroup"
_ROT_90 = 'rot="5400000"'
_FLIP_H = 'flipH="1"'
_FLIP_V = 'flipV="1"'


def _shape(
    prst: str, text: str, *, off: tuple[int, int], ext: tuple[int, int], rot_deg: int = 0
) -> str:
    rot = f' rot="{rot_deg * 60000}"' if rot_deg else ""
    return (
        '<wps:wsp><wps:cNvPr id="2" name="s"/><wps:cNvSpPr/><wps:spPr>'
        f'<a:xfrm{rot}><a:off x="{off[0]}" y="{off[1]}"/><a:ext cx="{ext[0]}" cy="{ext[1]}"/>'
        f'</a:xfrm><a:prstGeom prst="{prst}"><a:avLst/></a:prstGeom><a:noFill/>'
        '<a:ln w="19050"><a:solidFill><a:srgbClr val="000000"/></a:solidFill></a:ln></wps:spPr>'
        f"<wps:txbx><w:txbxContent><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:txbxContent>"
        "</wps:txbx><wps:bodyPr/></wps:wsp>"
    )


def _group(
    children: str,
    *,
    tag: str = "wpg:wgp",
    attrs: str = "",
    off: tuple[int, int] = (0, 0),
    ext: tuple[int, int] = (1000, 400),
    ch_off: tuple[int, int] = (0, 0),
    ch_ext: tuple[int, int] = (1000, 400),
) -> str:
    return (
        f'<{tag}><wpg:cNvPr id="3" name="g"/><wpg:cNvGrpSpPr/><wpg:grpSpPr>'
        f'<a:xfrm {attrs}><a:off x="{off[0]}" y="{off[1]}"/><a:ext cx="{ext[0]}" cy="{ext[1]}"/>'
        f'<a:chOff x="{ch_off[0]}" y="{ch_off[1]}"/><a:chExt cx="{ch_ext[0]}" cy="{ch_ext[1]}"/>'
        f"</a:xfrm></wpg:grpSpPr>{children}</{tag}>"
    )


def _drawing(group: str, *, cx: int, cy: int) -> BaseOxmlElement:
    return parse_xml(
        f'<w:drawing {_NS}><wp:inline><wp:extent cx="{cx}" cy="{cy}"/>'
        '<wp:docPr id="1" name="d"/>'
        f'<a:graphic><a:graphicData uri="{_GROUP_URI}">{group}</a:graphicData></a:graphic>'
        "</wp:inline></w:drawing>"
    )


def _read(group: str, *, cx: int, cy: int) -> ShapeFigure:
    return read_shape_figure(_drawing(group, cx=cx, cy=cy), Document())


def _asymmetric(attrs: str) -> str:
    rectangle = _shape("rect", "A", off=(100, 50), ext=(200, 100), rot_deg=30)
    ellipse = _shape("ellipse", "B", off=(600, 250), ext=(300, 100))
    return _group(rectangle + ellipse, attrs=attrs)


def _label_text(shape: FigureShape) -> str:
    assert shape.label is not None
    return "".join(run.text for paragraph in shape.label.paragraphs for run in paragraph.runs)


def _assert_box(
    shape: FigureShape,
    *,
    centre: tuple[float, float],
    size: tuple[float, float],
    rotation_deg: float,
    flip_h: bool = False,
    flip_v: bool = False,
) -> None:
    assert (shape.x_emu + shape.width_emu / 2, shape.y_emu + shape.height_emu / 2) == (
        pytest.approx(centre)
    )
    assert (shape.width_emu, shape.height_emu) == pytest.approx(size)
    assert shape.rotation_deg == pytest.approx(rotation_deg)
    assert (shape.flip_h, shape.flip_v) == (flip_h, flip_v)


def test_a_group_rotated_a_quarter_turn_moves_and_turns_each_shape_about_the_group_centre() -> None:
    figure = _read(_asymmetric(_ROT_90), cx=400, cy=1000)

    rectangle, ellipse = figure.shapes
    assert (_label_text(rectangle), _label_text(ellipse)) == ("A", "B")
    _assert_box(rectangle, centre=(300, 200), size=(200, 100), rotation_deg=120)
    _assert_box(ellipse, centre=(100, 750), size=(300, 100), rotation_deg=90)


def test_a_group_flipped_horizontally_mirrors_each_shape_and_negates_its_rotation() -> None:
    figure = _read(_asymmetric(_FLIP_H), cx=1000, cy=400)

    rectangle, ellipse = figure.shapes
    assert (_label_text(rectangle), _label_text(ellipse)) == ("A", "B")
    _assert_box(rectangle, centre=(800, 100), size=(200, 100), rotation_deg=330, flip_h=True)
    _assert_box(ellipse, centre=(250, 300), size=(300, 100), rotation_deg=0, flip_h=True)


def test_a_group_flipped_vertically_mirrors_each_shape_and_negates_its_rotation() -> None:
    figure = _read(_asymmetric(_FLIP_V), cx=1000, cy=400)

    rectangle, ellipse = figure.shapes
    assert (_label_text(rectangle), _label_text(ellipse)) == ("A", "B")
    _assert_box(rectangle, centre=(200, 300), size=(200, 100), rotation_deg=330, flip_v=True)
    _assert_box(ellipse, centre=(750, 100), size=(300, 100), rotation_deg=0, flip_v=True)


def test_a_group_mirrors_inside_its_box_before_it_rotates() -> None:
    figure = _read(_asymmetric(f"{_ROT_90} {_FLIP_H}"), cx=400, cy=1000)

    rectangle, ellipse = figure.shapes
    _assert_box(rectangle, centre=(300, 800), size=(200, 100), rotation_deg=60, flip_h=True)
    _assert_box(ellipse, centre=(100, 250), size=(300, 100), rotation_deg=90, flip_h=True)


def test_a_group_flipped_both_ways_toggles_both_flips_and_keeps_the_rotation() -> None:
    figure = _read(_asymmetric(f"{_FLIP_H} {_FLIP_V}"), cx=1000, cy=400)

    rectangle, ellipse = figure.shapes
    _assert_box(
        rectangle, centre=(800, 300), size=(200, 100), rotation_deg=30, flip_h=True, flip_v=True
    )
    _assert_box(
        ellipse, centre=(250, 100), size=(300, 100), rotation_deg=0, flip_h=True, flip_v=True
    )


def test_a_group_nested_in_a_rotated_group_composes_scale_offset_and_rotation() -> None:
    inner_child = _shape("rect", "N", off=(50, 20), ext=(60, 40))
    inner = _group(
        inner_child,
        tag="wpg:grpSp",
        off=(500, 100),
        ext=(400, 200),
        ch_off=(0, 0),
        ch_ext=(200, 100),
    )
    outer = _group(inner, attrs=_ROT_90)

    figure = _read(outer, cx=400, cy=1000)

    (shape,) = figure.shapes
    assert _label_text(shape) == "N"
    _assert_box(shape, centre=(220, 660), size=(120, 80), rotation_deg=90)


def test_a_rotated_group_nested_in_a_rotated_group_adds_both_rotations() -> None:
    inner_child = _shape("rect", "N", off=(50, 20), ext=(60, 40))
    inner = _group(
        inner_child,
        tag="wpg:grpSp",
        attrs=_ROT_90,
        off=(500, 100),
        ext=(400, 200),
        ch_off=(0, 0),
        ch_ext=(200, 100),
    )
    outer = _group(inner, attrs=_ROT_90)

    figure = _read(outer, cx=400, cy=1000)

    (shape,) = figure.shapes
    assert _label_text(shape) == "N"
    _assert_box(shape, centre=(240, 720), size=(120, 80), rotation_deg=180)


def test_a_quarter_turned_shape_in_an_unevenly_scaled_group_swaps_its_scale_factors() -> None:
    child = _shape("rect", "A", off=(100, 50), ext=(200, 50), rot_deg=90)
    group = _group(child, off=(0, 0), ext=(2000, 500), ch_off=(0, 0), ch_ext=(1000, 500))

    (shape,) = _read(group, cx=2000, cy=500).shapes

    # Turned a quarter, the 200 wide side is vertical and keeps the y factor 1; the 50 high side
    # is horizontal and takes the x factor 2.
    _assert_box(shape, centre=(400, 75), size=(200, 100), rotation_deg=90)


def test_an_inner_rotation_that_a_child_cancels_leaves_the_uneven_scale_a_plain_scale() -> None:
    child = _shape("rect", "A", off=(50, 50), ext=(100, 100), rot_deg=315)
    inner = _group(
        child, tag="wpg:grpSp", attrs='rot="2700000"', off=(100, 100), ext=(400, 200),
        ch_off=(0, 0), ch_ext=(400, 200),
    )  # fmt: skip
    outer = _group(inner, ext=(2000, 500), ch_ext=(1000, 500))

    (shape,) = _read(outer, cx=2000, cy=500).shapes

    # Inner centre (300, 200); child centre (100, 100) +(100, 100) = (200, 200) is (-100, 0) from
    # it, turned 45 degrees to (-70.71, -70.71); the outer group scales x by 2.
    _assert_box(
        shape, centre=(2 * (300 - 70.7106781), 200 - 70.7106781), size=(200, 100), rotation_deg=0
    )


@pytest.mark.parametrize("angle", [45, 30])
def test_a_shape_turned_off_the_axes_in_an_unevenly_scaled_group_is_a_shear(angle: int) -> None:
    child = _shape("rect", "A", off=(100, 50), ext=(200, 100), rot_deg=angle)
    group = _group(child, off=(0, 0), ext=(2000, 500), ch_off=(0, 0), ch_ext=(1000, 500))

    with pytest.raises(UnsupportedShapeFigure) as caught:
        _read(group, cx=2000, cy=500)

    assert "scaled unevenly" in caught.value.reason


def test_a_rotated_group_scaled_unevenly_scales_first_so_its_shapes_stay_rectangles() -> None:
    child = _shape("rect", "A", off=(100, 50), ext=(200, 100))
    group = _group(
        child, attrs='rot="2700000"', off=(0, 0), ext=(2000, 500), ch_off=(0, 0), ch_ext=(1000, 500)
    )
    bounding_box = 2500 / 2**0.5  # a 2000 x 500 box turned 45 degrees
    k = 2**-0.5

    (shape,) = _read(group, cx=round(bounding_box), cy=round(bounding_box)).shapes

    # Child centre (200, 100) scales to (400, 100), (-600, -150) from the group centre; a
    # 45 degree turn maps (dx, dy) to ((dx - dy) k, (dx + dy) k).
    half = round(bounding_box) / 2
    _assert_box(
        shape, centre=(half + (-600 + 150) * k, half + (-600 - 150) * k), size=(400, 100),
        rotation_deg=45,
    )  # fmt: skip


def test_an_inner_group_rotation_under_an_unevenly_scaled_group_is_a_shear() -> None:
    inner = _group(
        _shape("rect", "A", off=(50, 50), ext=(100, 100)),
        tag="wpg:grpSp", attrs='rot="2700000"', off=(100, 100), ext=(400, 200),
        ch_off=(0, 0), ch_ext=(400, 200),
    )  # fmt: skip
    outer = _group(inner, ext=(2000, 500), ch_ext=(1000, 500))

    with pytest.raises(UnsupportedShapeFigure) as caught:
        _read(outer, cx=2000, cy=500)

    assert "scaled unevenly" in caught.value.reason
