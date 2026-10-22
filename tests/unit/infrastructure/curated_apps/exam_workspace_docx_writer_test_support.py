"""Shared builders and DOCX inspection helpers for the exam workspace DOCX writer tests.

The writer's student-copy tests were split by responsibility: content and house template,
answer forms, shared pictures and stored Word shapes. Every load-bearing document builder and
every paragraph-inspection helper the split suites share lives here, so all four read the
produced bytes back the same way.
"""

import io
from uuid import uuid4

from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from PIL import Image

from skriptoteket.domain.curated_apps.exam_workspace.native_exam_body import (
    NativeGapSegment,
    NativeParagraph,
    NativeTextSegment,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_criteria import (
    NativeExamCriterion,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_document import (
    NativeExamDocument,
    NativeExamDocumentOrigin,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_item import (
    NativeExamItem,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_values import (
    NativeAnswerKey,
    NativeChoice,
    NativeGap,
    NativeItemReview,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_vocabulary import (
    NativeAnswerKeyOrigin,
    NativeExamItemKind,
    NativeItemReviewState,
    NativeParseOrigin,
)
from tests.unit.domain.curated_apps.exam_workspace.native_exam_test_builders import (
    make_free_text_item,
    plain_body,
)

STEP_BOX = "☐"
NBSP = "\u00a0"


def review_complete() -> NativeItemReview:
    return NativeItemReview(
        state=NativeItemReviewState.REVIEW_COMPLETE,
        parse_origin=NativeParseOrigin.TEACHER_CREATED,
    )


def png(color: tuple[int, int, int, int], size: tuple[int, int]) -> bytes:
    buffer = io.BytesIO()
    Image.new("RGBA", size, color).save(buffer, format="PNG")
    return buffer.getvalue()


def only_paragraph(paragraphs: list[Paragraph], text: str) -> Paragraph:
    matches = [paragraph for paragraph in paragraphs if paragraph.text == text]
    assert len(matches) == 1, f"expected one paragraph {text!r}, found {len(matches)}"
    return matches[0]


def run_shape(paragraph: Paragraph) -> list[str]:
    shape: list[str] = []
    for run in paragraph.runs:
        if run._element.find(qn("w:drawing")) is not None:
            shape.append("image")
        elif run.text:
            shape.append(run.text)
    return shape


def drawing_payload(parsed: DocxDocument, paragraph: Paragraph) -> bytes:
    drawing = next(
        element
        for run in paragraph.runs
        if (element := run._element.find(qn("w:drawing"))) is not None
    )
    blip = next(drawing.iter(qn("a:blip")))
    payload = parsed.part.related_parts[blip.attrib[qn("r:embed")]].blob
    assert isinstance(payload, bytes)
    return payload


def document_text(parsed: DocxDocument) -> str:
    parts = [paragraph.text for paragraph in parsed.paragraphs]
    for table in parsed.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def heading_texts(parsed: DocxDocument, heading: str) -> list[str]:
    return [
        paragraph.text
        for paragraph in parsed.paragraphs
        if paragraph.style is not None and paragraph.style.name == heading
    ]


def border_edges(paragraph: Paragraph) -> tuple[str, ...] | None:
    pPr = paragraph._p.find(qn("w:pPr"))
    if pPr is None:
        return None
    borders = pPr.find(qn("w:pBdr"))
    if borders is None:
        return None
    return tuple(child.tag.split("}")[1] for child in borders)


def border_attributes(paragraph: Paragraph, edge: str) -> dict[str, str | None]:
    pPr = paragraph._p.find(qn("w:pPr"))
    assert pPr is not None
    borders = pPr.find(qn("w:pBdr"))
    assert borders is not None
    element = borders.find(qn(f"w:{edge}"))
    assert element is not None
    return {name: element.get(qn(f"w:{name}")) for name in ("val", "sz", "space", "color")}


def run_border(run: Run) -> dict[str, str | None] | None:
    """The run's own border attributes, or None when the run carries no border."""
    rPr = run._element.find(qn("w:rPr"))
    if rPr is None:
        return None
    border = rPr.find(qn("w:bdr"))
    if border is None:
        return None
    return {name: border.get(qn(f"w:{name}")) for name in ("val", "sz", "space", "color")}


def ruled_paragraphs(paragraphs: list[Paragraph]) -> list[Paragraph]:
    return [paragraph for paragraph in paragraphs if border_edges(paragraph) is not None]


def blank_run_count(paragraph: Paragraph) -> int:
    blanks = [
        run for run in paragraph.runs if run.font.underline and run.text and set(run.text) == {NBSP}
    ]
    assert len(blanks) == 1, f"expected one underlined blank run, found {len(blanks)}"
    return len(blanks[0].text)


def free_text_item(*, writing_lines: int, item_id: str = "item_001") -> NativeExamItem:
    payload = make_free_text_item(item_id=item_id).model_dump()
    payload.update(writing_lines=writing_lines)
    return NativeExamItem.model_validate(payload)


def gap_fill_item(
    body: tuple[NativeParagraph, ...],
    gaps: tuple[NativeGap, ...],
    *,
    item_id: str = "item_001",
    sequence: int = 1,
    origin: NativeAnswerKeyOrigin = NativeAnswerKeyOrigin.SOURCE_PROVIDED,
) -> NativeExamItem:
    return NativeExamItem(
        item_id=item_id,
        sequence=sequence,
        kind=NativeExamItemKind.GAP_FILL,
        body=body,
        points=2,
        gaps=gaps,
        answer_key=NativeAnswerKey(origin=origin),
        review=review_complete(),
    )


def lone_gap_item() -> NativeExamItem:
    return gap_fill_item(
        (NativeParagraph(heading_level=None, segments=(NativeGapSegment(gap_id="gap_001"),)),),
        (NativeGap(gap_id="gap_001", accepted_values=("NF",)),),
    )


def sentence_gap_item(
    accepted: tuple[str, ...],
    *,
    origin: NativeAnswerKeyOrigin = NativeAnswerKeyOrigin.SOURCE_PROVIDED,
) -> NativeExamItem:
    return gap_fill_item(
        (
            NativeParagraph(
                heading_level=None,
                segments=(
                    NativeTextSegment(text="Svar: ", marks=()),
                    NativeGapSegment(gap_id="gap_001"),
                ),
            ),
        ),
        (NativeGap(gap_id="gap_001", accepted_values=accepted),),
        origin=origin,
    )


def all_kinds_document() -> NativeExamDocument:
    items = (
        NativeExamItem(
            item_id="item_001",
            sequence=1,
            kind=NativeExamItemKind.FREE_TEXT,
            writing_lines=8,
            body=plain_body("Resonera om bisatser."),
            points=4,
            criteria=(NativeExamCriterion(text="Kriteriet nämner bisatser", points=4),),
            answer_key=NativeAnswerKey(origin=NativeAnswerKeyOrigin.NOT_APPLICABLE),
            review=review_complete(),
        ),
        NativeExamItem(
            item_id="item_002",
            sequence=2,
            kind=NativeExamItemKind.SINGLE_CHOICE,
            body=plain_body("Välj rätt alternativ."),
            points=1,
            choices=(
                NativeChoice(choice_id="choice_001", body=plain_body("Nominalfras")),
                NativeChoice(choice_id="choice_002", body=plain_body("Verbfras")),
            ),
            answer_key=NativeAnswerKey(
                origin=NativeAnswerKeyOrigin.SOURCE_PROVIDED,
                correct_choice_ids=("choice_002",),
            ),
            review=review_complete(),
        ),
        NativeExamItem(
            item_id="item_003",
            sequence=3,
            kind=NativeExamItemKind.MULTIPLE_RESPONSE,
            body=plain_body("Markera alla fraser."),
            points=2,
            choices=(
                NativeChoice(choice_id="choice_001", body=plain_body("Prepositionsfras")),
                NativeChoice(choice_id="choice_002", body=plain_body("Adjektivfras")),
                NativeChoice(choice_id="choice_003", body=plain_body("Subjunktion")),
            ),
            answer_key=NativeAnswerKey(
                origin=NativeAnswerKeyOrigin.TEACHER_AUTHORED,
                correct_choice_ids=("choice_001", "choice_002"),
            ),
            review=review_complete(),
        ),
        NativeExamItem(
            item_id="item_004",
            sequence=4,
            kind=NativeExamItemKind.GAP_FILL,
            body=(
                NativeParagraph(
                    heading_level=None,
                    segments=(
                        NativeTextSegment(text="a) Frastyp: ", marks=()),
                        NativeGapSegment(gap_id="gap_001"),
                    ),
                ),
            ),
            points=2.5,
            gaps=(NativeGap(gap_id="gap_001", accepted_values=("NF", "nominalfras")),),
            answer_key=NativeAnswerKey(origin=NativeAnswerKeyOrigin.SOURCE_PROVIDED),
            review=review_complete(),
        ),
    )
    return NativeExamDocument(
        document_id=uuid4(),
        revision=1,
        title="Grammatikprov",
        instructions=(
            NativeParagraph(
                heading_level=None, segments=(NativeTextSegment(text="Svenska 2", marks=()),)
            ),
            NativeParagraph(
                heading_level=None, segments=(NativeTextSegment(text="Total: 9,5 poäng", marks=()),)
            ),
        ),
        items=items,
        origin=NativeExamDocumentOrigin(kind="created"),
    )
