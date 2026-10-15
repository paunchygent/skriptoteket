"""Picture extraction from DOCX body blocks: places, media types and repeats (D2)."""

import io

from docx import Document
from docx.document import Document as DocxDocument
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from PIL import Image

from skriptoteket.domain.curated_apps.exam_workspace.docx_chunks import asset_id_for
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_docx_chunker import (  # noqa: E501
    PythonDocxExamChunker,
)


def _png() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (120, 60), (40, 90, 160)).save(buffer, format="PNG")
    return buffer.getvalue()


def _jpeg() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (40, 20), (10, 20, 30)).save(buffer, format="JPEG")
    return buffer.getvalue()


def _save(document: DocxDocument) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _chunks(document: DocxDocument):
    return PythonDocxExamChunker().chunk(content=_save(document))


def test_a_picture_inside_a_paragraph_keeps_its_place_in_plain_and_marked_text() -> None:
    payload = _png()
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Före").italic = True
    paragraph.add_run(" ")
    paragraph.add_run().add_picture(io.BytesIO(payload))
    paragraph.add_run(" efter")

    (chunk,) = _chunks(document)

    assert chunk.text == "Före  efter"
    assert chunk.marked_text == "<i>Före</i>  efter"
    (image,) = chunk.images
    assert image.asset_id == asset_id_for(payload)
    assert image.offset == 5
    assert image.marked_offset == 12
    assert image.media_type == "image/png"
    assert (image.width, image.height) == (120, 60)


def test_a_picture_after_a_tab_and_a_break_counts_their_source_characters() -> None:
    document = Document()
    run = document.add_paragraph().add_run("A")
    run.add_tab()
    run.add_break()
    run.add_picture(io.BytesIO(_png()))

    (chunk,) = _chunks(document)

    assert chunk.text == "A\t\n"
    (image,) = chunk.images
    assert image.offset == 3
    assert image.marked_offset == 3


def test_hyperlink_wording_stays_in_plain_text_before_the_picture() -> None:
    document = Document()
    paragraph = document.add_paragraph()
    paragraph._p.append(
        parse_xml(
            f'<w:hyperlink {nsdecls("w")} w:anchor="mal"><w:r><w:t>Länk</w:t></w:r></w:hyperlink>'
        )
    )
    paragraph.add_run(" ").add_picture(io.BytesIO(_png()))

    (chunk,) = _chunks(document)

    assert chunk.text == "Länk "
    assert chunk.marked_text == "Länk "
    (image,) = chunk.images
    assert image.offset == 5
    assert image.marked_offset == 5


def test_the_same_picture_twice_in_one_run_is_two_pictures() -> None:
    payload = _png()
    document = Document()
    run = document.add_paragraph().add_run()
    run.add_picture(io.BytesIO(payload))
    run.add_picture(io.BytesIO(payload))

    (chunk,) = _chunks(document)

    assert chunk.text == ""
    assert [image.asset_id for image in chunk.images] == [asset_id_for(payload)] * 2
    assert [image.offset for image in chunk.images] == [0, 0]


def test_a_picture_in_a_table_cell_keeps_its_place_in_the_joined_text() -> None:
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Vänster"
    table.cell(0, 1).paragraphs[0].add_run().add_picture(io.BytesIO(_jpeg()))

    (chunk,) = _chunks(document)

    assert chunk.kind == "table"
    assert chunk.text == "Vänster | "
    (image,) = chunk.images
    assert image.media_type == "image/jpeg"
    assert image.offset == len("Vänster | ")
    assert image.marked_offset == len("Vänster | ")


def test_a_picture_alone_in_its_paragraph_is_a_drawing_at_offset_zero() -> None:
    document = Document()
    document.add_picture(io.BytesIO(_png()))

    (chunk,) = _chunks(document)

    assert chunk.kind == "drawing"
    assert chunk.text == ""
    (image,) = chunk.images
    assert image.offset == 0
    assert image.marked_offset == 0


def test_a_drawn_part_in_a_format_without_a_media_type_is_not_an_asset() -> None:
    document = Document()
    document.add_picture(io.BytesIO(_png()))
    paragraph = document.element.body.findall(qn("w:p"))[-1]
    blip = next(paragraph.iter(qn("a:blip")))
    emf = Part(PackURI("/word/media/figur.emf"), "image/x-emf", b"\x01\x00\x00\x00", None)
    blip.set(qn("r:embed"), document.part.relate_to(emf, RELATIONSHIP_TYPE.IMAGE))

    (chunk,) = _chunks(document)

    assert chunk.kind == "drawing"
    assert chunk.images == ()
