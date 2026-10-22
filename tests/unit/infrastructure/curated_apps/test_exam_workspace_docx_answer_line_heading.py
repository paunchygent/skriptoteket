"""A size-derived answer-line heading never returns to the imported body (W10, R2).

A Normal-style DOCX answer-line paragraph with a larger font is marked by the real
chunker as a size-derived heading. Heading restoration must still exclude it as
answer space, while a real larger-font heading is still restored.
"""

import io
from uuid import uuid4

from docx import Document
from docx.shared import Pt

from skriptoteket.application.curated_apps.handlers.exam_workspace_docx_build_admission import (
    admit_batch,
)
from skriptoteket.application.curated_apps.handlers.exam_workspace_docx_build_state import (
    DocxBuildState,
)
from skriptoteket.domain.curated_apps.exam_workspace.docx_ai_build import assemble_build_result
from skriptoteket.domain.curated_apps.exam_workspace.docx_ai_contracts import (
    DocxBuildResult,
    ProducedItem,
    ProducedParagraph,
    ProducedSegment,
)
from skriptoteket.domain.curated_apps.exam_workspace.docx_ai_item import aligned_native_item
from skriptoteket.domain.curated_apps.exam_workspace.docx_answer_line_gate import (
    ANSWER_LINE_ONLY_REASON,
)
from skriptoteket.domain.curated_apps.exam_workspace.docx_chunks import DocxChunk
from skriptoteket.domain.curated_apps.exam_workspace.docx_own_instructions import (
    LEFT_OUT_REASON,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_vocabulary import (
    NativeExamItemKind,
    NativeItemReviewState,
)
from skriptoteket.domain.curated_apps.exam_workspace.shared_body import body_plain_text
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_docx_chunker import (  # noqa: E501
    PythonDocxExamChunker,
)


def _chunked(prompt: str, second: str) -> tuple[DocxChunk, ...]:
    source = Document()
    prompt_run = source.add_paragraph().add_run(prompt)
    prompt_run.font.size = Pt(11)
    second_run = source.add_paragraph().add_run(second)
    second_run.font.size = Pt(22)
    second_run.underline = True
    stream = io.BytesIO()
    source.save(stream)
    return PythonDocxExamChunker().chunk(content=stream.getvalue())


def _paragraph(text: str) -> ProducedParagraph:
    return ProducedParagraph(
        segments=(ProducedSegment(kind="text", text=text, accepted_values=()),)
    )


def _produced(*texts: str) -> ProducedItem:
    return ProducedItem(
        source_anchors=("body:0001", "body:0002"),
        outcome="converted",
        review_reason=None,
        kind=NativeExamItemKind.FREE_TEXT,
        printed_label=None,
        key_source=None,
        facit_text=None,
        body=tuple(_paragraph(text) for text in texts),
        points=1,
        choices=(),
        match_left=(),
        match_right=(),
        match_pairs=(),
        criteria=(),
    )


def test_a_larger_font_answer_line_marked_as_heading_stays_out_of_the_body() -> None:
    chunks = _chunked("Förklara.", "_____")

    assert chunks[1].paragraphs[0].heading_level is not None

    aligned = aligned_native_item(
        _produced("Förklara.", "_____"), chunks=chunks, language="sv", foreign_chunks=()
    )

    assert body_plain_text(aligned.item.body) == "Förklara."


def test_a_larger_font_answer_line_stays_out_through_final_assembly_and_admission() -> None:
    chunks = _chunked("Förklara.", "_____")

    assert chunks[1].paragraphs[0].heading_level is not None

    produced = _produced_owned(("body:0001", "body:0002"), "Förklara.", "_____")
    assembled = _assembled(chunks, produced)

    assert [body_plain_text(item.body) for item in assembled.document.items] == ["Förklara."]

    admitted = _admitted(chunks, produced)

    assert [body_plain_text(item.body) for item in admitted.document.items] == ["Förklara."]


def test_a_larger_font_real_heading_is_still_restored() -> None:
    chunks = _chunked("Förklara.", "Del A")

    assert chunks[1].paragraphs[0].heading_level is not None

    aligned = aligned_native_item(
        _produced("Förklara."), chunks=chunks, language="sv", foreign_chunks=()
    )

    assert body_plain_text(aligned.item.body) == "Förklara.\nDel A"


def _explicitly_styled(second: str, *, underline: bool) -> tuple[DocxChunk, ...]:
    """A Normal 11pt prompt plus a Heading 1 second paragraph, as the R2 probe builds it."""
    source = Document()
    prompt_run = source.add_paragraph().add_run("Förklara.")
    prompt_run.font.size = Pt(11)
    line = source.add_paragraph(style="Heading 1")
    second_run = line.add_run(second)
    second_run.font.size = Pt(22)
    second_run.underline = underline
    stream = io.BytesIO()
    source.save(stream)
    return PythonDocxExamChunker().chunk(content=stream.getvalue())


def _produced_owned(anchors: tuple[str, ...], *texts: str) -> ProducedItem:
    return ProducedItem(
        source_anchors=anchors,
        outcome="converted",
        review_reason=None,
        kind=NativeExamItemKind.FREE_TEXT,
        printed_label=None,
        key_source=None,
        facit_text=None,
        body=tuple(_paragraph(text) for text in texts),
        points=1,
        choices=(),
        match_left=(),
        match_right=(),
        match_pairs=(),
        criteria=(),
    )


def _assembled(chunks: tuple[DocxChunk, ...], produced: ProducedItem) -> DocxBuildResult:
    aligned = aligned_native_item(produced, chunks=chunks, language="sv", foreign_chunks=())
    assert body_plain_text(aligned.item.body) == "Förklara."
    return assemble_build_result(
        origin_kind="docx_import",
        document_id=uuid4(),
        filename="synthetic.docx",
        source_sha256="0" * 64,
        chunks=chunks,
        accepted=((aligned, produced.source_anchors),),
        review_reasons={},
        questions=(produced.source_anchors,),
    )


def _admitted(chunks: tuple[DocxChunk, ...], produced: ProducedItem) -> DocxBuildResult:
    state = DocxBuildState(
        origin_kind="docx_import",
        job_id=uuid4(),
        document_id=uuid4(),
        filename="synthetic.docx",
        source_sha256="0" * 64,
        chunks=chunks,
        questions=(produced.source_anchors,),
    )
    admit_batch({"items": [produced.model_dump(mode="json")]}, state)
    return state.result()


def test_a_heading_styled_answer_line_never_returns_through_final_assembly() -> None:
    chunks = _explicitly_styled("_____", underline=True)

    assert [chunk.kind for chunk in chunks] == ["paragraph", "heading"]

    produced = _produced_owned(("body:0001", "body:0002"), "Förklara.", "_____")
    result = _assembled(chunks, produced)

    assert [body_plain_text(item.body) for item in result.document.items] == ["Förklara."]
    assert all(
        item.review.state is NativeItemReviewState.REVIEW_COMPLETE for item in result.document.items
    )


def test_a_heading_styled_answer_line_never_returns_through_application_admission() -> None:
    chunks = _explicitly_styled("_____", underline=True)

    assert [chunk.kind for chunk in chunks] == ["paragraph", "heading"]

    produced = _produced_owned(("body:0001", "body:0002"), "Förklara.", "_____")
    final = _admitted(chunks, produced)

    assert [body_plain_text(item.body) for item in final.document.items] == ["Förklara."]
    assert [item.review.state for item in final.document.items] == [
        NativeItemReviewState.REVIEW_COMPLETE
    ]


def test_a_real_heading_keeps_its_section_behaviour_through_final_assembly() -> None:
    chunks = _explicitly_styled("Del A", underline=False)

    assert [chunk.kind for chunk in chunks] == ["paragraph", "heading"]

    produced = _produced_owned(("body:0001", "body:0002"), "Förklara.")
    result = _assembled(chunks, produced)

    assert [body_plain_text(item.body) for item in result.document.items] == ["Del A\nFörklara."]
    assert result.document.parents == ()


def test_a_real_heading_keeps_its_section_behaviour_through_application_admission() -> None:
    chunks = _explicitly_styled("Del A", underline=False)

    assert [chunk.kind for chunk in chunks] == ["paragraph", "heading"]

    produced = _produced_owned(("body:0001", "body:0002"), "Förklara.")
    final = _admitted(chunks, produced)

    assert [body_plain_text(item.body) for item in final.document.items] == ["Del A\nFörklara."]
    # A listed heading that section placement restores keeps no missing-text reason
    # (TASK-SKRIPT-39-04-19): the admission event records it, the placed item shows it.
    assert all(ANSWER_LINE_ONLY_REASON not in item.review.reasons for item in final.document.items)
    assert all(LEFT_OUT_REASON not in item.review.reasons for item in final.document.items)
    assert all(
        "Del A" not in reason for item in final.document.items for reason in item.review.reasons
    )
    assert [item.review.state for item in final.document.items] == [
        NativeItemReviewState.REVIEW_COMPLETE
    ]
    assert final.document.parents == ()
