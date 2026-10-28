"""Main's new source-paragraph builders leave no printed answer line behind (W10, M1).

Each case builds its source through the real chunker and asserts BOTH the direct
``assemble_build_result`` path and the application ``admit_batch`` + ``result`` path.
The two builders under proof are ``source_gap_paragraphs.gap_paragraphs_from_source``
(GAP_FILL items) and ``source_misplaced_paragraph.rebuild_misplaced`` (a reply
paragraph repeating another question's source). The gap replies below drop the
source's own lead wording, so the asserted gap text can only come from the
source rebuild, never from the reply's projection. A standalone gap paragraph
(W6) and a picture paragraph in the same source survive both paths unchanged.
"""

import io
from uuid import uuid4

from docx import Document
from docx.document import Document as DocxDocument
from PIL import Image
from pydantic import JsonValue

from skriptoteket.application.curated_apps.handlers.exam_workspace_docx_build_admission import (
    admit_batch,
)
from skriptoteket.application.curated_apps.handlers.exam_workspace_docx_build_state import (
    DocxBuildState,
)
from skriptoteket.domain.curated_apps.exam_workspace.docx_ai_build import assemble_build_result
from skriptoteket.domain.curated_apps.exam_workspace.docx_ai_contracts import (
    DocxBuildResult,
    ProducedItem,
    ProducedParagraph,
    ProducedSegment,
)
from skriptoteket.domain.curated_apps.exam_workspace.docx_ai_item import aligned_native_item
from skriptoteket.domain.curated_apps.exam_workspace.docx_chunks import DocxChunk
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_vocabulary import (
    NativeExamItemKind,
    NativeItemReviewState,
)
from skriptoteket.domain.curated_apps.exam_workspace.shared_body import (
    body_asset_ids,
    body_plain_text,
)
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_docx_chunker import (  # noqa: E501
    PythonDocxExamChunker,
)


def _chunked(document: DocxDocument) -> tuple[DocxChunk, ...]:
    stream = io.BytesIO()
    document.save(stream)
    return PythonDocxExamChunker().chunk(content=stream.getvalue()).chunks


def _text(text: str) -> ProducedSegment:
    return ProducedSegment(kind="text", text=text, accepted_values=())


def _gap(*accepted: str) -> ProducedSegment:
    return ProducedSegment(kind="gap", text=None, accepted_values=tuple(accepted))


def _gap_item(anchors: tuple[str, ...], *paragraphs: ProducedParagraph) -> ProducedItem:
    return ProducedItem(
        source_anchors=anchors,
        outcome="converted",
        review_reason=None,
        kind=NativeExamItemKind.GAP_FILL,
        printed_label=None,
        key_source="copied",
        facit_text="svar",
        body=paragraphs,
        points=1,
        choices=(),
        match_left=(),
        match_right=(),
        match_pairs=(),
        criteria=(),
    )


def _free_item(anchors: tuple[str, ...], *texts: str) -> ProducedItem:
    paragraphs = tuple(ProducedParagraph(segments=(_text(text),)) for text in texts)
    return ProducedItem(
        source_anchors=anchors,
        outcome="converted",
        review_reason=None,
        kind=NativeExamItemKind.FREE_TEXT,
        printed_label=None,
        key_source=None,
        facit_text=None,
        body=paragraphs,
        points=1,
        choices=(),
        match_left=(),
        match_right=(),
        match_pairs=(),
        criteria=(),
    )


def _assemble(
    chunks: tuple[DocxChunk, ...],
    items: tuple[ProducedItem, ...],
    questions: tuple[tuple[str, ...], ...],
) -> DocxBuildResult:
    aligned = tuple(
        aligned_native_item(
            produced,
            chunks=[chunk for chunk in chunks if chunk.anchor in produced.source_anchors],
            language="sv",
            foreign_chunks=(),
        )
        for produced in items
    )
    return assemble_build_result(
        origin_kind="docx_import",
        document_id=uuid4(),
        filename="synthetic.docx",
        source_sha256="0" * 64,
        chunks=chunks,
        accepted=tuple((entry, produced.source_anchors) for entry, produced in zip(aligned, items)),
        review_reasons={},
        instructions=(),
        questions=questions,
    )


def _assemble_with_foreign(
    chunks: tuple[DocxChunk, ...],
    items: tuple[ProducedItem, ...],
    questions: tuple[tuple[str, ...], ...],
) -> DocxBuildResult:
    owned = {anchor for question in questions for anchor in question}
    aligned = tuple(
        aligned_native_item(
            produced,
            chunks=[chunk for chunk in chunks if chunk.anchor in produced.source_anchors],
            language="sv",
            foreign_chunks=tuple(
                chunk
                for chunk in chunks
                if chunk.anchor in owned and chunk.anchor not in produced.source_anchors
            ),
        )
        for produced in items
    )
    return assemble_build_result(
        origin_kind="docx_import",
        document_id=uuid4(),
        filename="synthetic.docx",
        source_sha256="0" * 64,
        chunks=chunks,
        accepted=tuple((entry, produced.source_anchors) for entry, produced in zip(aligned, items)),
        review_reasons={},
        instructions=(),
        questions=questions,
    )


def _admit(
    chunks: tuple[DocxChunk, ...],
    items: tuple[ProducedItem, ...],
    questions: tuple[tuple[str, ...], ...],
) -> DocxBuildResult:
    state = DocxBuildState(
        origin_kind="docx_import",
        job_id=uuid4(),
        document_id=uuid4(),
        filename="synthetic.docx",
        source_sha256="0" * 64,
        chunks=chunks,
        questions=questions,
    )
    payload: dict[str, JsonValue] = {"items": [item.model_dump(mode="json") for item in items]}
    admit_batch(payload, state)
    return state.result()


def _gap_source_document() -> DocxDocument:
    document = Document()
    document.add_paragraph("Fyll i: ____ år")
    document.add_paragraph("_____")
    document.add_paragraph("Motivera.")
    return document


def test_gap_source_paragraphs_drop_the_printed_answer_line_on_both_paths() -> None:
    chunks = _chunked(_gap_source_document())
    anchors = tuple(chunk.anchor for chunk in chunks)
    assert [chunk.kind for chunk in chunks] == ["paragraph", "paragraph", "paragraph"]
    # The reply drops the source's own "Fyll i: " lead, so the asserted gap text
    # can only come from the source rebuild, never from the reply's projection.
    produced = _gap_item(
        anchors,
        ProducedParagraph(segments=(_gap("2000"), _text(" år"))),
        ProducedParagraph(segments=(_text("Motivera."),)),
    )
    expected = "Fyll i: [___] år\nMotivera."

    assembled = _assemble(chunks, (produced,), (anchors,))
    (item,) = assembled.document.items
    assert body_plain_text(item.body) == expected
    assert "_____" not in body_plain_text(item.body)
    assert item.gaps != ()
    assert item.review.state is NativeItemReviewState.REVIEW_COMPLETE
    assert item.review.reasons == ()

    admitted = _admit(chunks, (produced,), (anchors,))
    (item,) = admitted.document.items
    assert body_plain_text(item.body) == expected
    assert "_____" not in body_plain_text(item.body)
    assert item.gaps != ()
    assert item.review.state is NativeItemReviewState.REVIEW_COMPLETE
    assert item.review.reasons == ()


def test_a_standalone_gap_paragraph_survives_on_both_paths() -> None:
    document = Document()
    document.add_paragraph("_____")
    chunks = _chunked(document)
    anchors = (chunks[0].anchor,)
    produced = _gap_item(anchors, ProducedParagraph(segments=(_gap("2000"),)))

    assembled = _assemble(chunks, (produced,), (anchors,))
    (item,) = assembled.document.items
    assert body_plain_text(item.body) == "[___]"
    assert item.gaps != ()
    assert item.review.state is NativeItemReviewState.REVIEW_COMPLETE
    assert item.review.reasons == ()

    admitted = _admit(chunks, (produced,), (anchors,))
    (item,) = admitted.document.items
    assert body_plain_text(item.body) == "[___]"
    assert item.gaps != ()
    assert item.review.state is NativeItemReviewState.REVIEW_COMPLETE
    assert item.review.reasons == ()


def test_a_rebuilt_misplaced_paragraph_drops_the_printed_answer_line_on_both_paths() -> None:
    document = Document()
    document.add_paragraph("1. Explain the water cycle.")
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "Förklara."
    cell.add_paragraph("_____")
    chunks = _chunked(document)
    assert chunks[1].kind == "table"
    first_anchors = (chunks[0].anchor,)
    second_anchors = (chunks[1].anchor,)
    questions = (first_anchors, second_anchors)

    first = _free_item(first_anchors, "Explain the water cycle.")
    second = _free_item(second_anchors, "Explain the water cycle.")

    assembled = _assemble_with_foreign(chunks, (first, second), questions)
    assert [body_plain_text(item.body) for item in assembled.document.items] == [
        "Explain the water cycle.",
        "Förklara.",
    ]
    assert "_____" not in body_plain_text(assembled.document.items[1].body)
    assert assembled.document.items[1].review.state is NativeItemReviewState.REVIEW_COMPLETE
    assert assembled.document.items[1].review.reasons == ()

    admitted = _admit(chunks, (first, second), questions)
    assert [body_plain_text(item.body) for item in admitted.document.items] == [
        "Explain the water cycle.",
        "Förklara.",
    ]
    assert "_____" not in body_plain_text(admitted.document.items[1].body)
    assert admitted.document.items[1].review.state is NativeItemReviewState.REVIEW_COMPLETE
    assert admitted.document.items[1].review.reasons == ()
    assert [event for event in admitted.events if event.code == "own_chunks_left_out"] == []


def test_a_picture_paragraph_in_gap_source_keeps_its_asset_on_both_paths() -> None:
    buffer = io.BytesIO()
    Image.new("RGB", (120, 60), (40, 90, 160)).save(buffer, format="PNG")
    payload = buffer.getvalue()
    document = _gap_source_document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Före ")
    paragraph.add_run().add_picture(io.BytesIO(payload))
    paragraph.add_run(" efter")
    chunks = _chunked(document)
    anchors = tuple(chunk.anchor for chunk in chunks)
    asset_id = chunks[3].images[0].asset_id
    produced = _gap_item(
        anchors,
        ProducedParagraph(segments=(_gap("2000"), _text(" år"))),
        ProducedParagraph(segments=(_text("Motivera."),)),
        ProducedParagraph(
            segments=(
                _text("Före "),
                ProducedSegment(kind="asset", text=None, asset_id=asset_id, accepted_values=()),
                _text(" efter"),
            )
        ),
    )
    expected = "Fyll i: [___] år\nMotivera.\nFöre    efter"

    assembled = _assemble(chunks, (produced,), (anchors,))
    (item,) = assembled.document.items
    assert body_plain_text(item.body) == expected
    assert "_____" not in body_plain_text(item.body)
    assert body_asset_ids(item.body) == (asset_id,)
    assert item.review.state is NativeItemReviewState.REVIEW_COMPLETE
    assert item.review.reasons == ()

    admitted = _admit(chunks, (produced,), (anchors,))
    (item,) = admitted.document.items
    assert body_plain_text(item.body) == expected
    assert "_____" not in body_plain_text(item.body)
    assert body_asset_ids(item.body) == (asset_id,)
    assert item.review.state is NativeItemReviewState.REVIEW_COMPLETE
    assert item.review.reasons == ()
