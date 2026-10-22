"""Shape reader: real Word groups, containers and preset geometry (TASK-SKRIPT-39-04-22 unit 1).

The real cases read the two fixture drawings of ``metodprov_med_bilder.docx``; the others place
hand-written shapes in a single shape, group and canvas container.
"""

from __future__ import annotations

from typing import get_args

import pytest
from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.oxml.xmlchemy import BaseOxmlElement

from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_docx_shape_figure import (  # noqa: E501
    FigureShape,
    LabelAnchor,
    LabelParagraph,
    LabelRun,
    PresetGeometry,
    ShapeLabel,
    ShapeOutline,
)
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_docx_shape_reader import (  # noqa: E501
    read_shape_figure,
)
from tests.unit.infrastructure.curated_apps.exam_workspace_docx_shape_reader_test_support import (
    FIXTURE,
    canvas_xml,
    group_xml,
    make_document,
    make_drawing,
    read_one,
    rejected,
    shape_text,
    wsp_xml,
)

pytestmark = pytest.mark.unit


def _fixture_drawing(doc_pr_id: str) -> tuple[DocxDocument, BaseOxmlElement]:
    document = Document(str(FIXTURE))
    for drawing in document.element.body.iter(qn("w:drawing")):
        doc_pr = drawing.find(f".//{qn('wp:docPr')}")
        if doc_pr is not None and doc_pr.get("id") == doc_pr_id:
            return document, drawing
    raise AssertionError(f"drawing {doc_pr_id} missing from the fixture")


def _centered(text: str, anchor: LabelAnchor) -> ShapeLabel:
    return ShapeLabel(
        paragraphs=(LabelParagraph((LabelRun(text, 10.0, "000000"),), "center"),),
        inset_left_emu=36000,
        inset_top_emu=36000,
        inset_right_emu=36000,
        inset_bottom_emu=36000,
        anchor=anchor,
    )


_BLACK_19050 = ShapeOutline(color="000000", width_emu=19050)


# --- real Word output ------------------------------------------------------------------------


def test_population_group_reads_the_two_ellipses_exactly() -> None:
    document, drawing = _fixture_drawing("9001")

    figure = read_shape_figure(drawing, document)

    assert (figure.width_emu, figure.height_emu) == (4320000, 1655999)
    assert figure.shapes == (
        FigureShape(
            preset="ellipse",
            x_emu=0,
            y_emu=0,
            width_emu=4320000,
            height_emu=1655999,
            fill=None,
            outline=_BLACK_19050,
            label=_centered("Svenska gymnasieelever (populationen)", "t"),
        ),
        FigureShape(
            preset="ellipse",
            x_emu=2232000,
            y_emu=684000,
            width_emu=1800000,
            height_emu=827999,
            fill="DDEBF7",
            outline=_BLACK_19050,
            label=_centered("Elever på ett program (urvalet)", "ctr"),
        ),
    )
    assert not figure.plain_lines


def test_selection_group_reads_boxes_and_arrows_in_paint_order() -> None:
    document, drawing = _fixture_drawing("9002")

    figure = read_shape_figure(drawing, document)

    assert (figure.width_emu, figure.height_emu) == (5040000, 864000)
    expected = [
        ("rect", 0, 1296000, 864000, None, "Alla elever på skolan"),
        ("rightArrow", 1404000, 432000, 288000, "404040", None),
        ("rect", 1908000, 1296000, 864000, None, "Elever som är i biblioteket"),
        ("rightArrow", 3311999, 432000, 288000, "404040", None),
        ("rect", 3744000, 1296000, 864000, None, "Svarar på frågan"),
    ]
    assert [
        (s.preset, s.x_emu, s.width_emu, s.height_emu, s.fill, shape_text(s)) for s in figure.shapes
    ] == expected
    assert [s.y_emu for s in figure.shapes] == [0, 288000, 0, 288000, 0]
    assert all(s.outline == _BLACK_19050 for s in figure.shapes)
    boxes = [s for s in figure.shapes if s.preset == "rect"]
    assert all(s.label == _centered(shape_text(s) or "", "ctr") for s in boxes)
    assert all(s.label is None for s in figure.shapes if s.preset == "rightArrow")


def test_a_picture_drawing_of_the_fixture_is_outside_d2() -> None:
    document, drawing = _fixture_drawing("9206")

    rejected(drawing, "picture", document)


# --- containers ------------------------------------------------------------------------------


def test_single_shape_with_ext_equal_to_the_extent_fills_the_box_whatever_its_offset() -> None:
    drawing = make_drawing(wsp_xml(off=(777, 888), ext=(4000, 2000)), cx=4000, cy=2000)

    figure = read_shape_figure(drawing, make_document())

    assert (figure.width_emu, figure.height_emu) == (4000, 2000)
    shape = figure.shapes[0]
    assert (shape.x_emu, shape.y_emu, shape.width_emu, shape.height_emu) == (0, 0, 4000, 2000)


def test_rotated_single_shape_keeps_its_own_ext_centred_in_the_extent_box() -> None:
    # Word sizes wp:extent to the rotated bounding box: a 1000000 square turned 45 degrees.
    drawing = make_drawing(
        wsp_xml(xfrm_attrs='rot="2700000"', off=(777, 888), ext=(1000000, 1000000)),
        cx=1414214,
        cy=1414214,
    )

    figure = read_shape_figure(drawing, make_document())

    assert (figure.width_emu, figure.height_emu) == (1414214, 1414214)
    shape = figure.shapes[0]
    assert (shape.x_emu, shape.y_emu, shape.width_emu, shape.height_emu) == (
        207107,
        207107,
        1000000,
        1000000,
    )
    assert shape.rotation_deg == 45.0


def test_single_shape_ext_centres_independently_in_x_and_y() -> None:
    drawing = make_drawing(wsp_xml(ext=(3000, 500)), cx=4000, cy=2000)

    shape = read_shape_figure(drawing, make_document()).shapes[0]

    assert (shape.x_emu, shape.y_emu, shape.width_emu, shape.height_emu) == (500, 750, 3000, 500)


def test_an_anchored_drawing_reads_like_an_inline_one() -> None:
    figure = read_shape_figure(
        make_drawing(wsp_xml(), anchor=True, cx=900, cy=300), make_document()
    )

    assert (figure.width_emu, figure.height_emu) == (900, 300)


def test_group_children_scale_from_child_space_to_the_figure() -> None:
    child = wsp_xml(off=(1500, 2500), ext=(1000, 500))
    group = group_xml(child, ext=(4000, 1000), ch_off=(1000, 2000), ch_ext=(2000, 500))

    figure = read_shape_figure(make_drawing(group, uri="wpg", cx=4000, cy=1000), make_document())

    shape = figure.shapes[0]
    assert (shape.x_emu, shape.y_emu, shape.width_emu, shape.height_emu) == (1000, 1000, 2000, 1000)


def test_a_top_level_group_maps_onto_the_wp_extent_not_its_own_offset() -> None:
    child = wsp_xml(off=(0, 0), ext=(100, 100))
    group = group_xml(child, off=(5000, 6000), ext=(200, 200), ch_off=(0, 0), ch_ext=(100, 100))

    figure = read_shape_figure(make_drawing(group, uri="wpg", cx=200, cy=200), make_document())

    shape = figure.shapes[0]
    assert (shape.x_emu, shape.y_emu, shape.width_emu, shape.height_emu) == (0, 0, 200, 200)


def test_nested_group_composes_inside_the_parent_mapping() -> None:
    inner = group_xml(
        wsp_xml(off=(10, 10), ext=(20, 20)),
        tag="wpg:grpSp",
        off=(100, 200),
        ext=(60, 60),
        ch_off=(0, 0),
        ch_ext=(30, 30),
    )
    outer = group_xml(inner, ext=(2000, 1000), ch_off=(0, 0), ch_ext=(1000, 500))

    figure = read_shape_figure(make_drawing(outer, uri="wpg", cx=2000, cy=1000), make_document())

    # child (10,10)-(30,30) -> nested space x2: (20,20) size 40 -> shifted to (120,220)
    # -> outer scale x2: (240, 440) size 80.
    shape = figure.shapes[0]
    assert (shape.x_emu, shape.y_emu, shape.width_emu, shape.height_emu) == (240, 440, 80, 80)


def test_group_with_shapes_and_nested_group_keeps_paint_order() -> None:
    first = wsp_xml(prst="rect", off=(0, 0), ext=(10, 10))
    nested = group_xml(
        wsp_xml(prst="ellipse", off=(0, 0), ext=(10, 10)),
        tag="wpg:grpSp",
        ext=(10, 10),
        ch_ext=(10, 10),
    )
    last = wsp_xml(prst="triangle", off=(0, 0), ext=(10, 10))
    group = group_xml(first + nested + last)

    figure = read_shape_figure(make_drawing(group, uri="wpg"), make_document())

    assert [s.preset for s in figure.shapes] == ["rect", "ellipse", "triangle"]


def test_canvas_children_keep_their_emu_coordinates() -> None:
    canvas = canvas_xml(wsp_xml(off=(300, 400), ext=(500, 600)))

    figure = read_shape_figure(make_drawing(canvas, uri="wpc", cx=2000, cy=2000), make_document())

    shape = figure.shapes[0]
    assert (shape.x_emu, shape.y_emu, shape.width_emu, shape.height_emu) == (300, 400, 500, 600)


def test_canvas_with_a_group_applies_the_group_mapping() -> None:
    inner = group_xml(
        wsp_xml(off=(0, 0), ext=(10, 10)),
        tag="wpg:grpSp",
        off=(100, 100),
        ext=(50, 50),
        ch_ext=(10, 10),
    )

    figure = read_shape_figure(make_drawing(canvas_xml(inner), uri="wpc"), make_document())

    shape = figure.shapes[0]
    assert (shape.x_emu, shape.y_emu, shape.width_emu, shape.height_emu) == (100, 100, 50, 50)


# --- geometry --------------------------------------------------------------------------------


@pytest.mark.parametrize("preset", sorted(get_args(PresetGeometry)))
def test_every_supported_preset_is_read(preset: str) -> None:
    connector = preset in {"line", "straightConnector1"}

    shape = read_one(wsp_xml(prst=preset, connector=connector))

    assert shape.preset == preset


def test_rotation_flips_and_adjustments_are_read() -> None:
    shape = read_one(
        wsp_xml(
            prst="roundRect",
            xfrm_attrs='rot="5400000" flipH="1" flipV="true"',
            guides='<a:gd name="adj" fmla="val 25000"/>',
        )
    )

    assert shape.rotation_deg == 90.0
    assert shape.flip_h is True
    assert shape.flip_v is True
    assert shape.adjustments == (("adj", 25000),)


def test_rotation_is_normalised_into_a_turn() -> None:
    assert read_one(wsp_xml(xfrm_attrs='rot="-5400000"')).rotation_deg == 270.0
    assert read_one(wsp_xml(xfrm_attrs='rot="21600000"')).rotation_deg == 0.0


def test_rotated_shape_in_an_evenly_scaled_group_keeps_its_rotation() -> None:
    group = group_xml(
        wsp_xml(xfrm_attrs='rot="2700000"', ext=(10, 10)), ext=(2000, 1000), ch_ext=(2000, 1000)
    )

    figure = read_shape_figure(make_drawing(group, uri="wpg", cx=2000, cy=1000), make_document())

    assert figure.shapes[0].rotation_deg == 45.0
