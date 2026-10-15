"""A source part markdown cannot carry reaches the labeler as an image of that part (D14, R-F5).

The production pipeline converts exams built in Word: a table with a merged cell, a picture
or a text box in a paragraph of its own, and a picture in a format no image can be made of.
The Decisions first pass is the real HTTP adapter; only the network is replaced.
"""

import base64
import io
import json
from collections.abc import Mapping
from datetime import UTC, datetime
from uuid import uuid4

import httpx
import pytest
from docx import Document
from docx.document import Document as DocxDocument
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.oxml.text.paragraph import CT_P
from PIL import Image
from pydantic import BaseModel, JsonValue

from skriptoteket.application.curated_apps.handlers.exam_workspace_docx_run import (
    DocxConversionJob,
    DocxConversionOutcome,
)
from skriptoteket.application.curated_apps.handlers.exam_workspace_docx_serial_uow import (
    SerializedUnitOfWork,
)
from skriptoteket.config import Settings
from skriptoteket.di.docx_conversion import build_docx_conversion_pipeline
from skriptoteket.domain.curated_apps.exam_conversion.digiexam_answer_key_llm_contracts import (
    StructuredLLMProviderProfile,
    StructuredLLMRequest,
    StructuredLLMResponse,
    StructuredLLMUsage,
)
from skriptoteket.domain.curated_apps.exam_workspace.docx_chunks import asset_id_for
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_export_rules import (
    native_export_blockers,
)
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_docx_chunker import (  # noqa: E501
    PythonDocxExamChunker,
)
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_labeler_markdown import (  # noqa: E501
    PythonDocxLabelerMarkdown,
)
from skriptoteket.protocols.item_type_labeler import ItemTypeLabel
from tests.fixtures.time_fixtures import FixedClock
from tests.unit.application.curated_apps.handlers.docx_pipeline_replay import (
    GrantingLeases,
    NoSleep,
    NoTransaction,
    RecordingProgress,
)

pytestmark = pytest.mark.unit

_TABLE_STEM = "Para ihop varje ord med sin ordklass."
_PICTURE_STEM = "Vad visar bilden?"


def _png() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (120, 60), (40, 90, 160)).save(buffer, format="PNG")
    return buffer.getvalue()


def _emf(document: DocxDocument, paragraph: CT_P) -> None:
    """Point the paragraph's picture at an EMF part, which no browser can draw."""
    blip = next(paragraph.iter(qn("a:blip")))
    emf = Part(PackURI("/word/media/figur.emf"), "image/x-emf", b"\x01\x00\x00\x00", None)
    blip.set(qn("r:embed"), document.part.relate_to(emf, RELATIONSHIP_TYPE.IMAGE))


def _save(document: DocxDocument) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _exam() -> bytes:
    """body:0001 title, body:0002-0003 a question with a merged-cell table, body:0004-0005 a
    question whose stem holds an EMF picture."""
    document = Document()
    document.add_paragraph("Ordklasser", style="Title")
    document.add_paragraph(f"1. {_TABLE_STEM}")
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Ord och ordklass"
    for row, (word, kind) in enumerate((("springa", "verb"), ("hus", "substantiv")), start=1):
        table.cell(row, 0).text = word
        table.cell(row, 1).text = kind
    document.add_paragraph(f"2. {_PICTURE_STEM} ").add_run().add_picture(io.BytesIO(_png()))
    picture = document.element.body.findall(qn("w:p"))[-1]
    document.add_paragraph("Svara med en mening.")
    _emf(document, picture)
    return _save(document)


_TEXT_BOX = (
    f"<mc:AlternateContent {nsdecls('w')} "
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
    '<mc:Choice Requires="wps"><w:drawing><w:txbxContent>'
    "<w:p><w:r><w:t>Ja eller nej</w:t></w:r></w:p></w:txbxContent></w:drawing></mc:Choice>"
    "<mc:Fallback><w:pict><w:txbxContent>"
    "<w:p><w:r><w:t>Ja eller nej</w:t></w:r></w:p></w:txbxContent></w:pict></mc:Fallback>"
    "</mc:AlternateContent>"
)


def _drawn_exam() -> bytes:
    """Each drawn block sits in a paragraph with no text of its own: body:0003 a PNG,
    body:0006 a text box, body:0008 an EMF picture."""
    document = Document()
    document.add_paragraph("Bildfrågor", style="Title")
    document.add_paragraph(f"1. {_PICTURE_STEM}")
    document.add_picture(io.BytesIO(_png()))
    document.add_paragraph("Svara med en mening.")
    document.add_paragraph("2. Vad står i rutan?")
    document.add_paragraph().add_run()._r.append(parse_xml(_TEXT_BOX))
    document.add_paragraph("3. Beskriv figuren.")
    document.add_picture(io.BytesIO(_png()))
    _emf(document, document.element.body.findall(qn("w:p"))[-1])
    return _save(document)


def _between_exam() -> bytes:
    """body:0003 a picture alone in its paragraph, between two questions."""
    document = Document()
    document.add_paragraph("Bildfrågor", style="Title")
    document.add_paragraph(f"1. {_PICTURE_STEM}")
    document.add_picture(io.BytesIO(_png()))
    document.add_paragraph("2. Beskriv vädret i dag.")
    return _save(document)


def _question(question_id: str, *anchors: int) -> dict[str, JsonValue]:
    return {
        "id": question_id,
        "anchors": [f"body:{anchor:04d}" for anchor in anchors],
        "subpart_anchors": [],
        "answer_key_anchors": [],
    }


_DRAWN_GROUPING: dict[str, JsonValue] = {
    "title_anchors": ["body:0001"],
    "questions": [_question("q1", 2, 3, 4), _question("q2", 5, 6), _question("q3", 7, 8)],
    "shared_contexts": [],
    "other_anchors": [],
}


_BETWEEN_GROUPING: dict[str, JsonValue] = {
    "title_anchors": ["body:0001"],
    "questions": [_question("q1", 2, 3), _question("q2", 4)],
    "shared_contexts": [],
    "other_anchors": [],
}


def _converted(text: str, *anchors: int, asset_ids: tuple[str, ...] = ()) -> dict[str, JsonValue]:
    """A converted free-text item in the shape the converter's schema returns."""
    body: list[JsonValue] = [
        {
            "segments": [
                {
                    "kind": "text",
                    "text": text[3:],
                    "accepted_values": [],
                    "spelling_candidates": [],
                }
            ]
        }
    ]
    body.extend(
        {
            "segments": [
                {
                    "kind": "asset",
                    "text": None,
                    "accepted_values": [],
                    "spelling_candidates": [],
                    "asset_id": asset_id,
                }
            ]
        }
        for asset_id in asset_ids
    )
    return {
        "source_anchors": [f"body:{anchor:04d}" for anchor in anchors],
        "outcome": "converted",
        "conflict_kind": None,
        "review_reason": None,
        "kind": "free_text",
        "key_source": None,
        "facit_text": None,
        "printed_label": text[0],
        "body": body,
        "points": 1,
        "choices": [],
        "match_left": [],
        "match_right": [],
        "match_pairs": [],
        "criteria": [],
    }


_GROUPING: dict[str, JsonValue] = {
    "title_anchors": ["body:0001"],
    "questions": [
        {
            "id": "q1",
            "anchors": ["body:0002", "body:0003"],
            "subpart_anchors": [],
            "answer_key_anchors": [],
        },
        {
            "id": "q2",
            "anchors": ["body:0004", "body:0005"],
            "subpart_anchors": [],
            "answer_key_anchors": [],
        },
    ],
    "shared_contexts": [],
    "other_anchors": [],
}


class _Provider:
    """Answer grouping with the exam's questions and every converter batch with the given
    title and items whose first source anchor the batch holds."""

    def __init__(
        self,
        grouping: Mapping[str, JsonValue],
        items: tuple[dict[str, JsonValue], ...] = (),
        title: dict[str, JsonValue] | None = None,
    ) -> None:
        self.grouping = grouping
        self.items = items
        self.title = title
        self.requests: list[StructuredLLMRequest] = []

    async def complete_structured(
        self, *, request: StructuredLLMRequest, profile: StructuredLLMProviderProfile
    ) -> StructuredLLMResponse:
        self.requests.append(request)
        content: Mapping[str, JsonValue] = (
            self.grouping
            if request.item_id.startswith("docx_grouping")
            else {
                "title": self.title if self.title and self._holds(request, self.title) else None,
                "instructions": [],
                "answer_key_anchors": [],
                "items": [item for item in self.items if self._holds(request, item)],
            }
        )
        return StructuredLLMResponse(
            content=dict(content), finish_reason="completed", usage=StructuredLLMUsage()
        )

    @staticmethod
    def _holds(request: StructuredLLMRequest, item: dict[str, JsonValue]) -> bool:
        anchors = item["source_anchors"]
        assert isinstance(anchors, list)
        return f'"{anchors[0]}"' in request.user_payload


def _free_text(name: str) -> dict[str, JsonValue]:
    return {
        "type": "choice",
        "name": name,
        "choice": "free_text",
        "probabilities": [
            {"value": label.value, "probability": 1.0 if label is ItemTypeLabel.FREE_TEXT else 0.0}
            for label in ItemTypeLabel
        ],
        "confidence": 1.0,
    }


class _Part(BaseModel):
    type: str
    text: str = ""
    image_url: str = ""


class _Message(BaseModel):
    content: list[_Part]


class _Question(BaseModel):
    name: str


class _DecisionsBody(BaseModel):
    """The fields of a Decisions request body these tests read."""

    input: list[_Message] | str  # text only when no source image goes in
    questions: list[_Question]


async def _convert(
    content: bytes,
    grouping: Mapping[str, JsonValue],
    items: tuple[dict[str, JsonValue], ...] = (),
    title: dict[str, JsonValue] | None = None,
) -> tuple[list[_DecisionsBody], _Provider, DocxConversionOutcome]:
    decisions: list[_DecisionsBody] = []

    def answer(request: httpx.Request) -> httpx.Response:
        body = _DecisionsBody.model_validate_json(request.content)
        decisions.append(body)
        answers = [_free_text(question.name) for question in body.questions]
        return httpx.Response(200, json={"answers": answers})

    provider = _Provider(grouping, items, title)
    pipeline = build_docx_conversion_pipeline(
        client=httpx.AsyncClient(transport=httpx.MockTransport(answer)),
        provider=provider,
        sleeper=NoSleep(),
        leases=GrantingLeases(),
        uow=SerializedUnitOfWork(NoTransaction()),
        clock=FixedClock(datetime(2026, 10, 11, tzinfo=UTC)),
        settings=Settings.model_construct(
            LLM_ANSWER_KEY_ENABLED=True, OPENAI_LLM_ANSWER_KEY_API_KEY="test-key"
        ),
        progress=RecordingProgress(),
    )
    outcome = await pipeline.convert(
        job=DocxConversionJob(
            job_id=uuid4(),
            worker_id="worker",
            document_id=uuid4(),
            filename="prov.docx",
            source_sha256="0" * 64,
            source_docx=content,
            chunks=PythonDocxExamChunker().chunk(content=content),
        )
    )
    return decisions, provider, outcome


def _image(part: _Part) -> Image.Image:
    assert part.image_url.startswith("data:image/png;base64,")
    return Image.open(
        io.BytesIO(base64.b64decode(part.image_url.removeprefix("data:image/png;base64,")))
    )


@pytest.mark.asyncio
async def test_a_merged_cell_reaches_the_labeler_as_an_image_and_a_lost_picture_is_explicit():
    decisions, provider, outcome = await _convert(_exam(), _GROUPING)

    (labeling,) = decisions
    assert [q.name for q in labeling.questions] == ["kind_q1"], "q2 must not be labeled"
    assert not isinstance(labeling.input, str), "the source images must go in"
    (message,) = labeling.input
    parts = message.content
    assert [part.type for part in parts] == ["input_text", "input_image"]
    assert parts[0].text.endswith(f"### Question q1\n1. {_TABLE_STEM}")
    assert "springa" not in parts[0].text, "the table must not also go in as text"
    image = _image(parts[1])
    assert image.width > 0 and image.height > 0

    events = [e for e in outcome.result.events if e.code == "labeler_source_image_unavailable"]
    assert [e.anchors for e in events] == [("body:0004",)]
    converted = [r.item_id for r in provider.requests if r.item_id.startswith("docx_")]
    assert "docx_wi-01" not in converted, "a named missing picture must not reach conversion"
    assert any(
        "figur.emf" in reason
        for item in outcome.result.document.items
        for reason in item.review.reasons
    )
    assert any(event.code == "source_picture_missing" for event in outcome.result.events)
    assert outcome.update.pending == ()


def test_a_block_with_only_a_drawing_is_projected_at_its_source_position():
    source = PythonDocxLabelerMarkdown().source(content=_drawn_exam())

    assert [chunk.anchor for chunk in source.chunks] == [f"body:{i:04d}" for i in range(1, 9)]
    assert sorted(source.images) == ["body:0003", "body:0006"]
    assert list(source.unavailable) == ["body:0008"]


@pytest.mark.asyncio
async def test_a_picture_or_text_box_alone_in_its_paragraph_reaches_the_labeler_and_converter():
    decisions, provider, outcome = await _convert(_drawn_exam(), _DRAWN_GROUPING)

    (grouping,) = [r for r in provider.requests if r.item_id.startswith("docx_grouping")]
    for anchor in ("body:0003", "body:0006", "body:0008"):
        assert f"[{anchor}]" in grouping.user_payload, "grouping must see every drawn block"

    (labeling,) = decisions
    assert [q.name for q in labeling.questions] == ["kind_q1", "kind_q2"]
    assert not isinstance(labeling.input, str), "the source images must go in"
    (message,) = labeling.input
    parts = message.content
    assert [part.type for part in parts] == ["input_text", "input_image"] * 2
    assert parts[0].text.endswith(f"### Question q1\n1. {_PICTURE_STEM}")
    assert parts[2].text.startswith("\nSvara med en mening."), "the picture stands before it"
    assert parts[2].text.endswith("### Question q2\n2. Vad står i rutan?")
    assert all("Ja eller nej" not in part.text for part in parts), "a text box is drawn, not typed"
    assert _image(parts[1]).width > 0 and _image(parts[3]).width > 0

    events = [e for e in outcome.result.events if e.code == "labeler_source_image_unavailable"]
    assert [e.anchors for e in events] == [("body:0008",)]
    batches = {
        r.item_id: r.user_payload for r in provider.requests if r.item_id.startswith("docx_w")
    }
    assert not any(item_id.startswith("docx_wi") for item_id in batches), (
        "the missing EMF is review-only"
    )
    missing_question = next(
        item
        for item in outcome.result.document.items
        if any("figur.emf" in reason for reason in item.review.reasons)
    )
    assert missing_question.review.state.value == "review_required"
    assert missing_question.retry_source is not None
    assert missing_question.retry_source.has_missing_pictures()
    assert any(
        blocker.item_id == missing_question.item_id and blocker.reason == "unbuilt_question"
        for blocker in native_export_blockers(outcome.result.document)
    )
    labeled = "".join(
        payload for item_id, payload in batches.items() if item_id.startswith("docx_w1")
    )
    assert '"body:0003"' in labeled and '"body:0006"' in labeled
    assert outcome.update.pending == ()


@pytest.mark.asyncio
@pytest.mark.parametrize("carried", [True, False], ids=["carrier-placed", "carrier-missing"])
async def test_a_picture_between_two_questions_places_the_carrier_or_falls_back_to_source(
    carried: bool,
):
    asset_id = asset_id_for(_png())
    first = _converted(
        f"1. {_PICTURE_STEM}",
        *((2, 3) if carried else (2,)),
        asset_ids=(asset_id,) if carried else (),
    )
    items = (first, _converted("2. Beskriv vädret i dag.", 4))

    title: dict[str, JsonValue] = {"text": "Bildfrågor", "source_anchors": ["body:0001"]}

    _, provider, outcome = await _convert(_between_exam(), _BETWEEN_GROUPING, items, title)

    (grouping,) = [r for r in provider.requests if r.item_id.startswith("docx_grouping")]
    assert f'[body:0003] (drawing) <image asset_id="{asset_id}"/>' in grouping.user_payload
    (holding,) = [
        r
        for r in provider.requests
        if r.item_id.startswith("docx_w") and '"body:0003"' in r.user_payload
    ]
    payload = json.loads(holding.user_payload)
    assert {
        "anchor": "body:0003",
        "kind": "drawing",
        "text": f'<image asset_id="{asset_id}"/>',
    } in payload["chunks"]

    document = outcome.result.document
    assert document.title == "Bildfrågor"
    placed = [(item.source_anchor, item.review.state.value) for item in document.items]
    if carried:
        assert placed == [("body:0002", "review_complete"), ("body:0004", "review_complete")]
    else:
        assert placed == [
            ("body:0002", "review_required"),
            ("body:0004", "review_complete"),
        ], "a picture the reply did not place keeps the question's own source, picture included"
    assert document.items[0].body_asset_ids() == (asset_id,)
    assert outcome.update.pending == ()
