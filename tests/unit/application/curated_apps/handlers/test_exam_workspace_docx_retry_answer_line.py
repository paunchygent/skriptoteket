"""Single-question retry passes the same W10 answer-line gate as the build (R5.1).

Every item ``RetryExamWorkspaceItemHandler.handle`` returns has passed
``docx_answer_line_gate.gated_item``: no printed answer line reaches a retried
question, and an answer-line-only retry still goes to teacher review with
``ANSWER_LINE_ONLY_REASON``. Each case drives the real handler the way the
fixtures and the rereview probe build it, with a scripted exchange reply.
"""

from __future__ import annotations

import hashlib
import io
from uuid import uuid4

import pytest
from docx import Document
from docx.document import Document as DocxDocument
from PIL import Image
from pydantic import JsonValue

from skriptoteket.domain.curated_apps.exam_workspace.docx_answer_line_gate import (
    ANSWER_LINE_ONLY_REASON,
)
from skriptoteket.domain.curated_apps.exam_workspace.docx_chunks import DocxChunk
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_document import (
    NativeExamAsset,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_retry_records import (
    NativeExamRetryMissingImage,
    NativeExamRetrySource,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_retry_source import (
    question_retry_source,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_vocabulary import (
    NativeExamItemKind,
    NativeItemReviewState,
)
from skriptoteket.domain.curated_apps.exam_workspace.shared_body import (
    body_asset_ids,
    body_plain_text,
)
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_docx_chunker import (  # noqa: E501
    PythonDocxExamChunker,
)
from tests.unit.application.curated_apps.handlers.exam_workspace_docx_retry_fixtures import (
    Env,
)
from tests.unit.application.curated_apps.handlers.exam_workspace_docx_retry_fixtures import (
    document as make_document,
)
from tests.unit.application.curated_apps.handlers.exam_workspace_docx_retry_fixtures import (
    item as make_item,
)
from tests.unit.application.curated_apps.handlers.exam_workspace_docx_retry_fixtures import (
    source as make_source,
)

pytestmark = [pytest.mark.unit, pytest.mark.asyncio]

_DISABLED_REASON = "AI-tolkning är avstängd. Kontrollera källtexten manuellt."


@pytest.fixture()
def env() -> Env:
    return Env()


def _chunked(document: DocxDocument) -> tuple[DocxChunk, ...]:
    stream = io.BytesIO()
    document.save(stream)
    return PythonDocxExamChunker().chunk(content=stream.getvalue()).chunks


def _question_document() -> DocxDocument:
    document = Document()
    document.add_paragraph("1. Beskriv försöket.")
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "Förklara."
    cell.add_paragraph("_____")
    document.add_paragraph("Motivera slutsatsen.")
    return document


def _text_segment(text: str) -> dict[str, JsonValue]:
    return {"kind": "text", "text": text, "accepted_values": [], "asset_id": None}


def _reply(
    anchors: tuple[str, ...],
    *texts: str,
    kind: str = "free_text",
    body: JsonValue | None = None,
) -> dict[str, JsonValue]:
    keyed = kind != "free_text"
    return {
        "title": None,
        "instructions": [],
        "answer_key_anchors": [],
        "items": [
            {
                "source_anchors": list(anchors),
                "outcome": "converted",
                "conflict_kind": None,
                "review_reason": None,
                "kind": kind,
                "key_source": "copied" if keyed else None,
                "facit_text": "svar" if keyed else None,
                "printed_label": None,
                "body": (
                    body
                    if body is not None
                    else [{"segments": [_text_segment(text)]} for text in texts]
                ),
                "points": 1,
                "choices": [],
                "match_left": [],
                "match_right": [],
                "match_pairs": [],
                "criteria": [],
            }
        ],
    }


def _owned_context_source(
    chunks: tuple[DocxChunk, ...],
) -> tuple[tuple[str, ...], tuple[str, ...], NativeExamRetrySource]:
    own = (chunks[0].anchor, chunks[2].anchor)
    owned = (chunks[1].anchor,)
    source = question_retry_source(
        own_anchors=own,
        owned_context_anchors=owned,
        key_anchors=(),
        read_anchors=frozenset(),
        chunks=chunks,
        label_kind=NativeExamItemKind.FREE_TEXT,
    )
    return own, owned, source


async def test_retry_drops_the_owned_context_answer_line_like_the_build(env: Env) -> None:
    chunks = _chunked(_question_document())
    assert [chunk.kind for chunk in chunks] == ["paragraph", "table", "paragraph"]
    own, _owned, saved_source = _owned_context_source(chunks)
    saved = make_item(
        "item_001", 1, anchor=own[0], text="1. Beskriv försöket.", source=saved_source
    )
    document = make_document(lineage_id=uuid4(), items=(saved,))
    await env.seed(document)
    env.exchange.reply = _reply(own, "1. Beskriv försöket.", "Motivera slutsatsen.")

    response = await env.retry(document)

    item = response.item
    assert body_plain_text(item.body) == "1. Beskriv försöket.\nFörklara.\nMotivera slutsatsen."
    assert item.review.state is NativeItemReviewState.REVIEW_COMPLETE
    assert item.item_id == saved.item_id
    assert item.source_anchor == saved.source_anchor
    assert item.retry_source == saved.retry_source


async def test_retry_with_only_an_answer_line_goes_to_review(env: Env) -> None:
    source_document = Document()
    source_document.add_paragraph("_____")
    chunks = _chunked(source_document)
    saved_source = question_retry_source(
        own_anchors=(chunks[0].anchor,),
        owned_context_anchors=(),
        key_anchors=(),
        read_anchors=frozenset(),
        chunks=chunks,
        label_kind=NativeExamItemKind.FREE_TEXT,
    )
    saved = make_item("item_001", 1, anchor=chunks[0].anchor, text="_____", source=saved_source)
    document = make_document(lineage_id=uuid4(), items=(saved,))
    await env.seed(document)
    env.exchange.reply = _reply((chunks[0].anchor,), "_____")

    response = await env.retry(document)

    item = response.item
    assert body_plain_text(item.body) == ""
    assert item.review.state is NativeItemReviewState.REVIEW_REQUIRED
    assert item.review.reasons.count(ANSWER_LINE_ONLY_REASON) == 1


async def test_disabled_ai_keeps_the_source_without_the_answer_line() -> None:
    disabled = Env(enabled=False)
    chunks = _chunked(_question_document())
    own, _owned, saved_source = _owned_context_source(chunks)
    saved = make_item(
        "item_001", 1, anchor=own[0], text="1. Beskriv försöket.", source=saved_source
    )
    document = make_document(lineage_id=uuid4(), items=(saved,))
    await disabled.seed(document)

    response = await disabled.retry(document)

    item = response.item
    assert "_____" not in body_plain_text(item.body)
    assert "Förklara." in body_plain_text(item.body)
    assert item.review.state is NativeItemReviewState.REVIEW_REQUIRED
    assert item.review.reasons == (_DISABLED_REASON,)
    assert disabled.exchange.requests == []


async def test_missing_picture_review_keeps_its_reason_without_the_answer_line(
    env: Env,
) -> None:
    chunks = _chunked(_question_document())
    own, _owned, base = _owned_context_source(chunks)
    missing = NativeExamRetryMissingImage(
        reference="/word/media/image3.emf",
        media_type="image/x-emf",
        offset=0,
        marked_offset=0,
    )
    saved_source = base.model_copy(
        update={
            "chunks": tuple(
                record.model_copy(update={"missing_images": (missing,)})
                if record.anchor == own[0]
                else record
                for record in base.chunks
            )
        }
    )
    saved = make_item(
        "item_001", 1, anchor=own[0], text="1. Beskriv försöket.", source=saved_source
    )
    document = make_document(lineage_id=uuid4(), items=(saved,))
    await env.seed(document)

    response = await env.retry(document)

    item = response.item
    assert "_____" not in body_plain_text(item.body)
    assert item.review.state is NativeItemReviewState.REVIEW_REQUIRED
    assert any("image3.emf" in reason for reason in item.review.reasons)
    assert env.exchange.requests == []


async def test_retry_keeps_a_gap_paragraph(env: Env) -> None:
    document = make_document(
        lineage_id=uuid4(),
        items=(
            make_item(
                "item_001",
                1,
                anchor="body:0001",
                text="Fyll i: ____ år",
                source=make_source(text="Fyll i: ____ år"),
            ),
        ),
    )
    await env.seed(document)
    env.exchange.reply = _reply(
        ("body:0001",),
        kind="gap_fill",
        body=[
            {
                "segments": [
                    _text_segment("Fyll i: "),
                    {
                        "kind": "gap",
                        "text": None,
                        "accepted_values": ["2000"],
                        "asset_id": None,
                    },
                    _text_segment(" år"),
                ]
            }
        ],
    )

    response = await env.retry(document, kind=NativeExamItemKind.GAP_FILL)

    item = response.item
    assert body_plain_text(item.body) == "Fyll i: [___] år"
    assert item.review.state is NativeItemReviewState.REVIEW_COMPLETE


async def test_retry_keeps_a_source_picture(env: Env) -> None:
    buffer = io.BytesIO()
    Image.new("RGB", (120, 60), (40, 90, 160)).save(buffer, format="PNG")
    payload = buffer.getvalue()
    source_document = Document()
    paragraph = source_document.add_paragraph()
    paragraph.add_run("Före ")
    paragraph.add_run().add_picture(io.BytesIO(payload))
    paragraph.add_run(" efter")
    chunks = _chunked(source_document)
    asset_id = chunks[0].images[0].asset_id
    saved_source = question_retry_source(
        own_anchors=(chunks[0].anchor,),
        owned_context_anchors=(),
        key_anchors=(),
        read_anchors=frozenset(),
        chunks=chunks,
        label_kind=NativeExamItemKind.FREE_TEXT,
    )
    asset = NativeExamAsset(
        asset_id=asset_id,
        sha256=hashlib.sha256(payload).hexdigest(),
        media_type="image/png",
        byte_length=len(payload),
    )
    saved = make_item(
        "item_001", 1, anchor=chunks[0].anchor, text="Före efter", source=saved_source
    )
    document = make_document(lineage_id=uuid4(), items=(saved,), assets=(asset,))
    await env.seed(document, assets_by_id={asset_id: payload})
    env.exchange.reply = _reply(
        (chunks[0].anchor,),
        body=[
            {
                "segments": [
                    _text_segment("Före "),
                    {
                        "kind": "asset",
                        "text": None,
                        "accepted_values": [],
                        "asset_id": asset_id,
                    },
                    _text_segment(" efter"),
                ]
            }
        ],
    )

    response = await env.retry(document)

    item = response.item
    assert body_asset_ids(item.body) == (asset_id,)
    assert item.review.state is NativeItemReviewState.REVIEW_COMPLETE


async def test_retry_without_an_answer_line_returns_the_conversion_unchanged(env: Env) -> None:
    document = make_document(
        lineage_id=uuid4(),
        items=(make_item("item_001", 1, anchor="body:0001", source=make_source()),),
    )
    await env.seed(document)
    env.exchange.reply = _reply(("body:0001",), "Förklara vattnets kretslopp kortfattat.")

    response = await env.retry(document)

    item = response.item
    assert body_plain_text(item.body) == "Förklara vattnets kretslopp kortfattat."
    assert item.review.state is NativeItemReviewState.REVIEW_COMPLETE
    assert item.review.reasons == ()
