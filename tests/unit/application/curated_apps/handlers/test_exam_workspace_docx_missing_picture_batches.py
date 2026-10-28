"""An unavailable picture gates only the families that depend on it (review F4).

One type batch packs several families of one item kind. A source picture the chunker could
not turn into an image must hold only the families that claim or read its chunk; every
other family in the same batch still converts. The Decisions first pass is the real HTTP
adapter; only the network is replaced.
"""

from __future__ import annotations

import io
import time
from collections.abc import Mapping
from uuid import uuid4

import pytest
from docx import Document
from docx.document import Document as DocxDocument
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml.ns import qn
from docx.oxml.text.paragraph import CT_P
from PIL import Image
from pydantic import BaseModel, JsonValue

from skriptoteket.application.curated_apps.handlers.exam_workspace_docx_pipeline_batches import (
    BatchPart,
    source_parts,
)
from skriptoteket.application.curated_apps.handlers.exam_workspace_docx_run import (
    DocxConversionJob,
    DocxConversionOutcome,
    DocxConversionRun,
)
from skriptoteket.application.curated_apps.handlers.exam_workspace_docx_serial_uow import (
    SerializedUnitOfWork,
)
from skriptoteket.config import Settings
from skriptoteket.di.docx_conversion import build_docx_conversion_pipeline
from skriptoteket.domain.curated_apps.exam_conversion.digiexam_answer_key_llm_contracts import (
    StructuredLLMProviderProfile,
    StructuredLLMRequest,
    StructuredLLMResponse,
    StructuredLLMUsage,
)
from skriptoteket.domain.curated_apps.exam_workspace.docx_batches import (
    ConversionBatch,
    QuestionFamily,
)
from skriptoteket.domain.curated_apps.exam_workspace.docx_chunks import (
    DocxChunk,
    DocxMissingImage,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_vocabulary import (
    NativeExamItemKind,
    NativeParseOrigin,
)
from skriptoteket.domain.llm_door import AiDoorCaller
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_docx_chunker import (  # noqa: E501
    PythonDocxExamChunker,
)
from skriptoteket.protocols.exam_answer_key import StructuredLLMTextObserver
from skriptoteket.protocols.item_type_labeler import ItemTypeLabel
from skriptoteket.protocols.llm_door import DoorOperation, DoorReply
from tests.unit.application.curated_apps.handlers.docx_pipeline_replay import (
    NoSleep,
    NoTransaction,
    RecordingProgress,
)
from tests.unit.llm_door_fakes import DoorSent, RecordingDoor, fixture_caller, json_reply

pytestmark = pytest.mark.unit


def _png() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (120, 60), (40, 90, 160)).save(buffer, format="PNG")
    return buffer.getvalue()


def _emf(document: DocxDocument, paragraph: CT_P) -> None:
    """Point the paragraph's picture at an EMF part, which no browser can draw."""
    blip = next(paragraph.iter(qn("a:blip")))
    emf = Part(PackURI("/word/media/figur.emf"), "image/x-emf", b"\x01\x00\x00\x00", None)
    blip.set(qn("r:embed"), document.part.relate_to(emf, RELATIONSHIP_TYPE.IMAGE))


def _exam() -> bytes:
    """body:0001 title, body:0002 an unplaced EMF picture, body:0003-0004 two questions."""
    document = Document()
    document.add_paragraph("Frågeprov", style="Title")
    document.add_picture(io.BytesIO(_png()))
    _emf(document, document.element.body.findall(qn("w:p"))[-1])
    document.add_paragraph("1. Beskriv vädret.")
    document.add_paragraph("2. Beskriv solen.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _question(question_id: str, *anchors: int) -> dict[str, JsonValue]:
    return {
        "id": question_id,
        "anchors": [f"body:{anchor:04d}" for anchor in anchors],
        "subpart_anchors": [],
        "answer_key_anchors": [],
    }


_GROUPING: dict[str, JsonValue] = {
    "title_anchors": ["body:0001"],
    "questions": [_question("q1", 3), _question("q2", 4)],
    "shared_contexts": [],
    "header_anchors": [],
}


def _converted(text: str, *anchors: int) -> dict[str, JsonValue]:
    """A converted free-text item in the shape the converter's schema returns."""
    return {
        "source_anchors": [f"body:{anchor:04d}" for anchor in anchors],
        "outcome": "converted",
        "conflict_kind": None,
        "review_reason": None,
        "kind": "free_text",
        "key_source": None,
        "facit_text": None,
        "printed_label": text[0],
        "body": [
            {
                "segments": [
                    {
                        "kind": "text",
                        "text": text[3:],
                        "accepted_values": [],
                    }
                ]
            }
        ],
        "points": 1,
        "choices": [],
        "match_left": [],
        "match_right": [],
        "match_pairs": [],
        "criteria": [],
    }


class _Provider:
    """Answer grouping with the exam's questions and every converter batch with the given
    title and items whose first source anchor the batch holds."""

    def __init__(
        self,
        grouping: Mapping[str, JsonValue],
        items: tuple[dict[str, JsonValue], ...] = (),
    ) -> None:
        self.grouping = grouping
        self.items = items
        self.requests: list[StructuredLLMRequest] = []

    async def complete_structured(
        self,
        *,
        request: StructuredLLMRequest,
        profile: StructuredLLMProviderProfile,
        caller: AiDoorCaller,
        on_text: StructuredLLMTextObserver | None = None,
    ) -> StructuredLLMResponse:
        self.requests.append(request)
        content: Mapping[str, JsonValue] = (
            self.grouping
            if request.item_id.startswith("docx_grouping")
            else {
                "title": None,
                "instructions": [],
                "answer_key_anchors": [],
                "items": [item for item in self.items if self._holds(request, item)],
            }
        )
        return StructuredLLMResponse(
            content=dict(content), finish_reason="completed", usage=StructuredLLMUsage()
        )

    @staticmethod
    def _holds(request: StructuredLLMRequest, item: dict[str, JsonValue]) -> bool:
        anchors = item["source_anchors"]
        assert isinstance(anchors, list)
        return f'"{anchors[0]}"' in request.user_payload


def _free_text(name: str) -> dict[str, JsonValue]:
    return {
        "type": "choice",
        "name": name,
        "choice": "free_text",
        "probabilities": [
            {"value": label.value, "probability": 1.0 if label is ItemTypeLabel.FREE_TEXT else 0.0}
            for label in ItemTypeLabel
        ],
        "confidence": 1.0,
    }


class _Part(BaseModel):
    type: str
    text: str = ""
    image_url: str = ""


class _Message(BaseModel):
    content: list[_Part]


class _Question(BaseModel):
    name: str


class _DecisionsBody(BaseModel):
    """The fields of a Decisions request body these tests read."""

    input: list[_Message] | str  # text only when no source image goes in
    questions: list[_Question]


async def _convert(
    content: bytes,
    grouping: Mapping[str, JsonValue],
    items: tuple[dict[str, JsonValue], ...] = (),
    *,
    families_per_batch: int = 1,
) -> tuple[list[_DecisionsBody], _Provider, DocxConversionOutcome]:
    decisions: list[_DecisionsBody] = []

    def answer(sent: DoorSent) -> DoorReply:
        assert sent.operation is DoorOperation.DECISIONS
        body = _DecisionsBody.model_validate(sent.payload)
        decisions.append(body)
        answers: list[JsonValue] = [_free_text(question.name) for question in body.questions]
        return json_reply({"answers": answers})

    provider = _Provider(grouping, items)
    pipeline = build_docx_conversion_pipeline(
        door=RecordingDoor(answer),
        provider=provider,
        sleeper=NoSleep(),
        uow=SerializedUnitOfWork(NoTransaction()),
        settings=Settings.model_construct(
            LLM_ANSWER_KEY_ENABLED=True,
            LLM_DOCX_BATCH_FAMILIES=families_per_batch,
        ),
        progress=RecordingProgress(),
    )
    outcome = await pipeline.convert(
        job=DocxConversionJob(
            origin_kind="docx_import",
            job_id=uuid4(),
            worker_id="worker",
            document_id=uuid4(),
            filename="prov.docx",
            source_sha256="0" * 64,
            source_docx=content,
            chunks=PythonDocxExamChunker().chunk(content=content).chunks,
            source_notes=(),
            course_key=None,
            caller=fixture_caller(),
        )
    )
    return decisions, provider, outcome


@pytest.mark.asyncio
async def test_a_missing_picture_holds_only_its_family_while_the_other_family_converts():
    decisions, provider, outcome = await _convert(
        _exam(),
        _GROUPING,
        (_converted("2. Beskriv solen.", 4),),
        families_per_batch=2,
    )

    assert [question.name for body in decisions for question in body.questions] == [
        "kind_q1",
        "kind_q2",
    ], "both questions are labelable: the picture is not part of either question's source"
    batches = {request.item_id: request for request in provider.requests}
    assert "docx_w1-01" in batches, "the clean family's batch must still reach the converter"
    payload = batches["docx_w1-01"].user_payload
    assert '"body:0004"' in payload
    assert '"body:0002"' not in payload, "the batch must not claim the unavailable picture"

    items = {item.source_anchor: item for item in outcome.result.document.items}
    converted = items["body:0004"]
    assert converted.review.parse_origin is NativeParseOrigin.LLM_PARSED
    assert not converted.review.reasons
    held = items["body:0003"]
    assert held.review.parse_origin is NativeParseOrigin.SOURCE_FALLBACK
    assert any("figur.emf" in reason for reason in held.review.reasons)
    events = [event for event in outcome.result.events if event.code == "source_picture_missing"]
    assert [(event.batch, event.anchors) for event in events] == [("w1-01", ("body:0002",))]
    assert outcome.update.pending == ()


def _chunk(anchor: str, *, missing: bool = False) -> DocxChunk:
    images = (
        (DocxMissingImage(reference="/word/media/figur.emf", media_type="image/x-emf"),)
        if missing
        else ()
    )
    return DocxChunk(
        anchor=anchor,
        kind="paragraph",
        text="Text",
        marked_text="Text",
        missing_images=images,
    )


def _run(chunks: tuple[DocxChunk, ...], families: tuple[QuestionFamily, ...]) -> DocxConversionRun:
    run = DocxConversionRun(
        DocxConversionJob(
            origin_kind="docx_import",
            job_id=uuid4(),
            worker_id="worker",
            document_id=uuid4(),
            filename="prov.docx",
            source_sha256="0" * 64,
            source_docx=b"",
            chunks=chunks,
            source_notes=(),
            course_key=None,
            caller=fixture_caller(),
        ),
        time.monotonic(),
    )
    run.families = families
    return run


def test_a_shared_missing_picture_gates_its_claimant_and_reader_but_not_a_clean_family():
    chunks = (
        _chunk("body:0001", missing=True),
        _chunk("body:0002"),
        _chunk("body:0003"),
        _chunk("body:0004"),
    )
    claimant = QuestionFamily(
        "q1", 0, ("body:0002",), frozenset({"body:0001", "body:0002"}), frozenset()
    )
    reader = QuestionFamily(
        "q2", 1, ("body:0003",), frozenset({"body:0003"}), frozenset({"body:0001"})
    )
    clean = QuestionFamily("q3", 2, ("body:0004",), frozenset({"body:0004"}), frozenset())
    batch = ConversionBatch(
        batch_id="w1-01",
        kinds=frozenset({NativeExamItemKind.FREE_TEXT}),
        question_ids=("q1", "q2", "q3"),
        claims=frozenset({"body:0001", "body:0002", "body:0003", "body:0004"}),
        reads=frozenset(),
    )

    held, converted = source_parts(_run(chunks, (claimant, reader, clean)), batch)

    assert held.batch.question_ids == ("q1", "q2")
    assert held.batch.claims == frozenset({"body:0001", "body:0002", "body:0003"})
    assert held.batch.reads == frozenset()
    assert [chunk.anchor for chunk in held.missing] == ["body:0001"]
    assert "figur.emf" in held.missing_reason
    assert converted.batch.question_ids == ("q3",)
    assert converted.batch.claims == frozenset({"body:0004"})
    assert converted.missing == ()


def test_a_batch_stays_whole_when_no_family_or_every_family_is_gated():
    family = QuestionFamily(
        "q1", 0, ("body:0002",), frozenset({"body:0001", "body:0002"}), frozenset()
    )
    batch = ConversionBatch(
        batch_id="w1-01",
        kinds=frozenset({NativeExamItemKind.FREE_TEXT}),
        question_ids=("q1",),
        claims=frozenset({"body:0001", "body:0002"}),
        reads=frozenset(),
    )

    (converting,) = source_parts(_run((_chunk("body:0001"), _chunk("body:0002")), (family,)), batch)
    assert converting == BatchPart(batch)

    (holding,) = source_parts(
        _run((_chunk("body:0001", missing=True), _chunk("body:0002")), (family,)), batch
    )
    assert holding.batch == batch
    assert [chunk.anchor for chunk in holding.missing] == ["body:0001"]
