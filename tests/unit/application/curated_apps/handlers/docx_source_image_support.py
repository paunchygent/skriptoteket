"""Exams, groupings, the scripted provider and the conversion run of the source-image tests."""

import base64
import io
from collections.abc import Mapping
from uuid import uuid4

from docx import Document
from docx.document import Document as DocxDocument
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.oxml.text.paragraph import CT_P
from PIL import Image
from pydantic import BaseModel, JsonValue

from skriptoteket.application.curated_apps.handlers.exam_workspace_docx_run import (
    DocxConversionJob,
    DocxConversionOutcome,
)
from skriptoteket.application.curated_apps.handlers.exam_workspace_docx_serial_uow import (
    SerializedUnitOfWork,
)
from skriptoteket.config import Settings
from skriptoteket.di.docx_conversion import build_docx_conversion_pipeline
from skriptoteket.domain.curated_apps.exam_conversion.digiexam_answer_key_llm_contracts import (
    StructuredLLMProviderProfile,
    StructuredLLMRequest,
    StructuredLLMResponse,
    StructuredLLMUsage,
)
from skriptoteket.domain.llm_door import AiDoorCaller
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_docx_chunker import (  # noqa: E501
    PythonDocxExamChunker,
)
from skriptoteket.protocols.exam_answer_key import StructuredLLMTextObserver
from skriptoteket.protocols.exam_workspace import DocxExamChunkerProtocol
from skriptoteket.protocols.item_type_labeler import ItemTypeLabel
from skriptoteket.protocols.llm_door import DoorOperation, DoorReply
from tests.unit.application.curated_apps.handlers.docx_pipeline_replay import (
    NoSleep,
    NoTransaction,
    RecordingProgress,
)
from tests.unit.llm_door_fakes import DoorSent, RecordingDoor, fixture_caller, json_reply

TABLE_STEM = "Para ihop varje ord med sin ordklass."
PICTURE_STEM = "Vad visar bilden?"


def picture_png() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (120, 60), (40, 90, 160)).save(buffer, format="PNG")
    return buffer.getvalue()


def add_emf(document: DocxDocument, paragraph: CT_P) -> None:
    """Point the paragraph's picture at an EMF part, which no browser can draw."""
    blip = next(paragraph.iter(qn("a:blip")))
    emf = Part(PackURI("/word/media/figur.emf"), "image/x-emf", b"\x01\x00\x00\x00", None)
    blip.set(qn("r:embed"), document.part.relate_to(emf, RELATIONSHIP_TYPE.IMAGE))


def save_docx(document: DocxDocument) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def table_exam() -> bytes:
    """body:0001 title, body:0002-0003 a question with a merged-cell table, body:0004-0005 a
    question whose stem holds an EMF picture."""
    document = Document()
    document.add_paragraph("Ordklasser", style="Title")
    document.add_paragraph(f"1. {TABLE_STEM}")
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Ord och ordklass"
    for row, (word, kind) in enumerate((("springa", "verb"), ("hus", "substantiv")), start=1):
        table.cell(row, 0).text = word
        table.cell(row, 1).text = kind
    document.add_paragraph(f"2. {PICTURE_STEM} ").add_run().add_picture(io.BytesIO(picture_png()))
    picture = document.element.body.findall(qn("w:p"))[-1]
    document.add_paragraph("Svara med en mening.")
    add_emf(document, picture)
    return save_docx(document)


TEXT_BOX = (
    f"<mc:AlternateContent {nsdecls('w')} "
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
    '<mc:Choice Requires="wps"><w:drawing><w:txbxContent>'
    "<w:p><w:r><w:t>Ja eller nej</w:t></w:r></w:p></w:txbxContent></w:drawing></mc:Choice>"
    "<mc:Fallback><w:pict><w:txbxContent>"
    "<w:p><w:r><w:t>Ja eller nej</w:t></w:r></w:p></w:txbxContent></w:pict></mc:Fallback>"
    "</mc:AlternateContent>"
)


def drawn_exam() -> bytes:
    """Each drawn block sits in a paragraph with no text of its own: body:0003 a PNG,
    body:0006 a text box, body:0008 an EMF picture."""
    document = Document()
    document.add_paragraph("Bildfrågor", style="Title")
    document.add_paragraph(f"1. {PICTURE_STEM}")
    document.add_picture(io.BytesIO(picture_png()))
    document.add_paragraph("Svara med en mening.")
    document.add_paragraph("2. Vad står i rutan?")
    document.add_paragraph().add_run()._r.append(parse_xml(TEXT_BOX))
    document.add_paragraph("3. Beskriv figuren.")
    document.add_picture(io.BytesIO(picture_png()))
    add_emf(document, document.element.body.findall(qn("w:p"))[-1])
    return save_docx(document)


def between_exam() -> bytes:
    """body:0003 a picture alone in its paragraph, between two questions."""
    document = Document()
    document.add_paragraph("Bildfrågor", style="Title")
    document.add_paragraph(f"1. {PICTURE_STEM}")
    document.add_picture(io.BytesIO(picture_png()))
    document.add_paragraph("2. Beskriv vädret i dag.")
    return save_docx(document)


def grouped_question(question_id: str, *anchors: int) -> dict[str, JsonValue]:
    return {
        "id": question_id,
        "anchors": [f"body:{anchor:04d}" for anchor in anchors],
        "subpart_anchors": [],
        "answer_key_anchors": [],
    }


DRAWN_GROUPING: dict[str, JsonValue] = {
    "title_anchors": ["body:0001"],
    "questions": [
        grouped_question("q1", 2, 3, 4),
        grouped_question("q2", 5, 6),
        grouped_question("q3", 7, 8),
    ],
    "shared_contexts": [],
    "header_anchors": [],
}


BETWEEN_GROUPING: dict[str, JsonValue] = {
    "title_anchors": ["body:0001"],
    "questions": [grouped_question("q1", 2, 3), grouped_question("q2", 4)],
    "shared_contexts": [],
    "header_anchors": [],
}


def converted(text: str, *anchors: int, asset_ids: tuple[str, ...] = ()) -> dict[str, JsonValue]:
    """A converted free-text item in the shape the converter's schema returns."""
    body: list[JsonValue] = [
        {
            "segments": [
                {
                    "kind": "text",
                    "text": text[3:],
                    "accepted_values": [],
                }
            ]
        }
    ]
    body.extend(
        {
            "segments": [
                {
                    "kind": "asset",
                    "text": None,
                    "accepted_values": [],
                    "asset_id": asset_id,
                }
            ]
        }
        for asset_id in asset_ids
    )
    return {
        "source_anchors": [f"body:{anchor:04d}" for anchor in anchors],
        "outcome": "converted",
        "conflict_kind": None,
        "review_reason": None,
        "kind": "free_text",
        "key_source": None,
        "facit_text": None,
        "printed_label": text[0],
        "body": body,
        "points": 1,
        "choices": [],
        "match_left": [],
        "match_right": [],
        "match_pairs": [],
        "criteria": [],
    }


GROUPING: dict[str, JsonValue] = {
    "title_anchors": ["body:0001"],
    "questions": [
        {
            "id": "q1",
            "anchors": ["body:0002", "body:0003"],
            "subpart_anchors": [],
            "answer_key_anchors": [],
        },
        {
            "id": "q2",
            "anchors": ["body:0004", "body:0005"],
            "subpart_anchors": [],
            "answer_key_anchors": [],
        },
    ],
    "shared_contexts": [],
    "header_anchors": [],
}


class ScriptedProvider:
    """Answer grouping with the exam's questions and every converter batch with the given
    title and items whose first source anchor the batch holds."""

    def __init__(
        self,
        grouping: Mapping[str, JsonValue],
        items: tuple[dict[str, JsonValue], ...] = (),
        title: dict[str, JsonValue] | None = None,
    ) -> None:
        self.grouping = grouping
        self.items = items
        self.title = title
        self.requests: list[StructuredLLMRequest] = []

    async def complete_structured(
        self,
        *,
        request: StructuredLLMRequest,
        profile: StructuredLLMProviderProfile,
        caller: AiDoorCaller,
        on_text: StructuredLLMTextObserver | None = None,
    ) -> StructuredLLMResponse:
        self.requests.append(request)
        content: Mapping[str, JsonValue] = (
            self.grouping
            if request.item_id.startswith("docx_grouping")
            else {
                "title": self.title if self.title and self._holds(request, self.title) else None,
                "instructions": [],
                "answer_key_anchors": [],
                "items": [item for item in self.items if self._holds(request, item)],
            }
        )
        return StructuredLLMResponse(
            content=dict(content), finish_reason="completed", usage=StructuredLLMUsage()
        )

    @staticmethod
    def _holds(request: StructuredLLMRequest, item: dict[str, JsonValue]) -> bool:
        anchors = item["source_anchors"]
        assert isinstance(anchors, list)
        return f'"{anchors[0]}"' in request.user_payload


def free_text_answer(name: str) -> dict[str, JsonValue]:
    return {
        "type": "choice",
        "name": name,
        "choice": "free_text",
        "probabilities": [
            {"value": label.value, "probability": 1.0 if label is ItemTypeLabel.FREE_TEXT else 0.0}
            for label in ItemTypeLabel
        ],
        "confidence": 1.0,
    }


class ImagePart(BaseModel):
    type: str
    text: str = ""
    image_url: str = ""


class Message(BaseModel):
    content: list[ImagePart]


class LabeledQuestion(BaseModel):
    name: str


class DecisionsBody(BaseModel):
    """The fields of a Decisions request body these tests read."""

    input: list[Message] | str  # text only when no source image goes in
    questions: list[LabeledQuestion]


async def convert(
    content: bytes,
    grouping: Mapping[str, JsonValue],
    items: tuple[dict[str, JsonValue], ...] = (),
    title: dict[str, JsonValue] | None = None,
    chunker: DocxExamChunkerProtocol | None = None,
    progress: RecordingProgress | None = None,
) -> tuple[list[DecisionsBody], ScriptedProvider, DocxConversionOutcome]:
    decisions: list[DecisionsBody] = []
    progress = progress or RecordingProgress()

    def answer(sent: DoorSent) -> DoorReply:
        assert sent.operation is DoorOperation.DECISIONS
        body = DecisionsBody.model_validate(sent.payload)
        decisions.append(body)
        progress.trace.append("door:decisions")
        answers: list[JsonValue] = [free_text_answer(question.name) for question in body.questions]
        return json_reply({"answers": answers})

    provider = ScriptedProvider(grouping, items, title)
    pipeline = build_docx_conversion_pipeline(
        door=RecordingDoor(answer),
        provider=provider,
        sleeper=NoSleep(),
        uow=SerializedUnitOfWork(NoTransaction()),
        settings=Settings.model_construct(LLM_ANSWER_KEY_ENABLED=True),
        progress=progress,
    )
    outcome = await pipeline.convert(
        job=DocxConversionJob(
            origin_kind="docx_import",
            job_id=uuid4(),
            worker_id="worker",
            document_id=uuid4(),
            filename="prov.docx",
            source_sha256="0" * 64,
            source_docx=content,
            chunks=(chunker or PythonDocxExamChunker()).chunk(content=content).chunks,
            source_notes=(),
            course_key=None,
            caller=fixture_caller(),
        )
    )
    return decisions, provider, outcome


def part_image(part: ImagePart) -> Image.Image:
    assert part.image_url.startswith("data:image/png;base64,")
    return Image.open(
        io.BytesIO(base64.b64decode(part.image_url.removeprefix("data:image/png;base64,")))
    )
