"""Empty exam texts save, block export when they are a question text, and print nothing as a
part instruction (TASK-SKRIPT-40-04-06 D3, D4)."""

import io

import pytest
from docx import Document
from pydantic import ValidationError

from skriptoteket.domain.curated_apps.exam_conversion.examnet_qti_item_admissibility import (
    blocking_errors,
)
from skriptoteket.domain.curated_apps.exam_workspace.examnet_pdf_export import (
    native_exam_to_pdf_document,
)
from skriptoteket.domain.curated_apps.exam_workspace.examnet_qti_export import (
    native_exam_to_examnet_qti_items,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_body import (
    NativeAssetSegment,
    NativeGapSegment,
    NativeParagraph,
    NativeTextSegment,
    is_empty_body,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_document import (
    NativeExamAsset,
    NativeExamDocument,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_export_rules import (
    native_export_blockers,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_item import NativeExamItem
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_parent import NativeExamParent
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_values import (
    NativeChoice,
    NativeGap,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_vocabulary import (
    NativeExamItemKind,
)
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_docx_writer import (  # noqa: E501
    ExamWorkspaceDocxWriter,
)
from tests.unit.domain.curated_apps.exam_workspace.native_exam_test_builders import (
    make_document,
    make_free_text_item,
    make_gap_fill_item,
    make_paragraph,
    make_single_choice_item,
    plain_body,
)

_EMPTY_PARAGRAPH = NativeParagraph(segments=(), heading_level=None)
_ASSET_ID = "asset_0123456789abcdef"
_ASSET = NativeExamAsset(
    asset_id=_ASSET_ID, sha256="a" * 64, media_type="image/png", byte_length=10
)


def _changed(
    item: NativeExamItem,
    *,
    body: tuple[NativeParagraph, ...] | None = None,
    kind: NativeExamItemKind | None = None,
    gaps: tuple[NativeGap, ...] | None = None,
    choices: tuple[NativeChoice, ...] | None = None,
) -> NativeExamItem:
    update: dict[str, object] = {
        field: value
        for field, value in (("body", body), ("kind", kind), ("gaps", gaps), ("choices", choices))
        if value is not None
    }
    return NativeExamItem.model_validate({**item.model_dump(), **update})


def _reasons(item: NativeExamItem) -> list[str]:
    return [blocker.reason for blocker in native_export_blockers(make_document(item))]


def test_an_empty_paragraph_between_texts_round_trips() -> None:
    body = (make_paragraph("Före"), _EMPTY_PARAGRAPH, make_paragraph("Efter"))
    item = _changed(make_free_text_item(), body=body)

    dumped = item.model_dump()

    assert dumped["body"][1]["segments"] == ()
    assert NativeExamItem.model_validate(dumped).body == body


def test_a_text_run_still_cannot_be_empty() -> None:
    with pytest.raises(ValidationError):
        NativeTextSegment(text="", marks=())


def test_a_question_with_no_text_validates() -> None:
    assert _changed(make_free_text_item(), body=()).body == ()


def test_a_choice_with_no_text_validates() -> None:
    item = make_single_choice_item()
    choices = tuple(choice.model_copy(update={"body": ()}) for choice in item.choices)

    assert all(choice.body == () for choice in _changed(item, choices=choices).choices)


def test_a_parent_with_no_instruction_validates() -> None:
    assert NativeExamParent(parent_id="parent_001", place=0, instruction=(), item_ids=("item_001",))


def test_a_gap_fill_question_with_no_gap_and_no_text_validates() -> None:
    item = _changed(make_gap_fill_item(), gaps=(), body=())

    assert item.gaps == () and item.body == ()


def test_a_placed_gap_without_a_record_is_still_refused() -> None:
    with pytest.raises(ValidationError, match="place every declared gap"):
        _changed(make_gap_fill_item(), gaps=())


def test_a_gap_record_the_body_does_not_place_is_still_refused() -> None:
    with pytest.raises(ValidationError, match="place every declared gap"):
        _changed(make_gap_fill_item(), body=plain_body("Ange frastyp: "))


@pytest.mark.parametrize("kind", [NativeExamItemKind.FREE_TEXT, NativeExamItemKind.SINGLE_CHOICE])
def test_a_gap_segment_in_a_question_of_another_kind_is_still_refused(
    kind: NativeExamItemKind,
) -> None:
    template = (
        make_free_text_item()
        if kind is NativeExamItemKind.FREE_TEXT
        else (make_single_choice_item())
    )
    body = (NativeParagraph(segments=(NativeGapSegment(gap_id="gap_001"),), heading_level=None),)

    with pytest.raises(ValidationError, match="only gap_fill bodies"):
        _changed(template, body=body)


@pytest.mark.parametrize(
    "body",
    [
        (),
        (_EMPTY_PARAGRAPH,),
        (
            NativeParagraph(
                segments=(NativeTextSegment(text=" \n\t", marks=("bold",)),), heading_level=None
            ),
        ),
        (make_paragraph("  "), _EMPTY_PARAGRAPH, make_paragraph("\n")),
    ],
)
def test_a_question_without_visible_text_is_an_empty_question_text_blocker(
    body: tuple[NativeParagraph, ...],
) -> None:
    assert is_empty_body(body)
    assert _reasons(_changed(make_free_text_item(), body=body)) == ["empty_question_text"]


def test_a_question_with_only_a_picture_is_not_an_empty_question_text() -> None:
    body = (
        NativeParagraph(segments=(NativeAssetSegment(asset_id=_ASSET_ID),), heading_level=None),
    )
    base = make_document(make_free_text_item()).model_dump()
    item = _changed(make_free_text_item(), body=body)
    document = NativeExamDocument.model_validate({**base, "items": (item,), "assets": (_ASSET,)})

    assert not is_empty_body(body)
    assert native_export_blockers(document) == ()


def test_a_question_with_only_a_gap_is_not_an_empty_question_text() -> None:
    item = make_gap_fill_item()
    body = (NativeParagraph(segments=(NativeGapSegment(gap_id="gap_001"),), heading_level=None),)

    assert "empty_question_text" not in _reasons(_changed(item, body=body))


def test_a_question_with_a_single_letter_is_not_an_empty_question_text() -> None:
    assert _reasons(_changed(make_free_text_item(), body=plain_body("x"))) == []


@pytest.mark.parametrize("kind", [NativeExamItemKind.GAP_FILL, NativeExamItemKind.INLINE_CHOICE])
def test_a_gap_question_with_no_gap_is_a_no_gaps_blocker(kind: NativeExamItemKind) -> None:
    item = _changed(make_gap_fill_item(), kind=kind, gaps=(), body=plain_body("Ingen lucka"))

    assert _reasons(item) == ["no_gaps"]


def test_a_gap_question_with_no_gap_and_no_text_reports_both_blockers() -> None:
    item = _changed(make_gap_fill_item(), gaps=(), body=())

    assert sorted(_reasons(item)) == ["empty_question_text", "no_gaps"]


def test_a_free_text_question_is_never_a_no_gaps_blocker() -> None:
    assert "no_gaps" not in _reasons(make_free_text_item())


def _two_children_under_an_empty_instruction() -> NativeExamDocument:
    document = make_document(
        make_free_text_item(sequence=1, item_id="item_001"),
        make_free_text_item(sequence=2, item_id="item_002"),
    )
    parent = NativeExamParent(
        parent_id="parent_001",
        place=document.items[0].sequence - 1,
        instruction=(),
        item_ids=("item_001", "item_002"),
    )
    return NativeExamDocument.model_validate({**document.model_dump(), "parents": (parent,)})


def test_pdf_prints_no_parent_instruction_for_an_empty_instruction() -> None:
    html = native_exam_to_pdf_document(_two_children_under_an_empty_instruction()).html

    assert 'class="parent-instruction"' not in html
    assert html.count("Resonera om renässansen.") == 2


def test_docx_writes_no_instruction_paragraph_for_an_empty_instruction() -> None:
    document = _two_children_under_an_empty_instruction()
    bare = make_document(*document.items)

    def paragraphs(source: NativeExamDocument) -> list[str]:
        data = ExamWorkspaceDocxWriter().build_docx_bytes(source)
        return [paragraph.text for paragraph in Document(io.BytesIO(data)).paragraphs]

    assert paragraphs(document) == paragraphs(bare)


def test_qti_exports_the_children_of_an_empty_instruction_as_standalone_items() -> None:
    document = _two_children_under_an_empty_instruction()

    items = native_exam_to_examnet_qti_items(document)

    assert [item.item_id for item in items] == ["item_001", "item_002"]
    assert all(not item.children for item in items)
    assert blocking_errors(items) == ()


def test_qti_still_groups_the_children_of_a_non_empty_instruction() -> None:
    document = _two_children_under_an_empty_instruction()
    parent = document.parents[0].model_copy(update={"instruction": (make_paragraph("Gemensamt."),)})
    document = NativeExamDocument.model_validate({**document.model_dump(), "parents": (parent,)})

    items = native_exam_to_examnet_qti_items(document)

    assert [item.item_id for item in items] == ["parent_001"]
    assert blocking_errors(items) == ()


def test_an_empty_exam_instruction_prints_nothing() -> None:
    document = make_document(make_free_text_item())
    assert document.instructions == ()

    html = native_exam_to_pdf_document(document).html

    assert 'class="document-instruction"' not in html
    (qti_item,) = native_exam_to_examnet_qti_items(document)
    assert qti_item.prompt_lines == ("Resonera om renässansen.",)


def _with_exam_instructions(*paragraphs: NativeParagraph) -> NativeExamDocument:
    document = make_document(make_free_text_item())
    return NativeExamDocument.model_validate({**document.model_dump(), "instructions": paragraphs})


def test_a_whitespace_only_exam_instruction_prints_nothing_in_any_export() -> None:
    blank = _with_exam_instructions(make_paragraph("   "), _EMPTY_PARAGRAPH)
    bare = make_document(make_free_text_item())

    html = native_exam_to_pdf_document(blank).html
    assert 'class="document-instruction"' not in html
    assert "<p>   </p>" not in html and "<p></p>" not in html

    def paragraphs(source: NativeExamDocument) -> list[str]:
        data = ExamWorkspaceDocxWriter().build_docx_bytes(source)
        return [paragraph.text for paragraph in Document(io.BytesIO(data)).paragraphs]

    assert paragraphs(blank) == paragraphs(bare)
    (qti_item,) = native_exam_to_examnet_qti_items(blank)
    assert qti_item.prompt_lines == ("Resonera om renässansen.",)


def test_an_exam_instruction_with_visible_text_still_prints_its_blank_lines() -> None:
    document = _with_exam_instructions(make_paragraph("Läs noga."), make_paragraph("  "))

    html = native_exam_to_pdf_document(document).html

    assert '<div class="document-instruction"><p>Läs noga.</p><p>  </p></div>' in html
