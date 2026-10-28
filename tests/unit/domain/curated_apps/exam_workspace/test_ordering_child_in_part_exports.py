"""Exports of a part holding an ordering child worth 1 point (TASK-SKRIPT-39-06-11 D1, D3).

A part is one composite QTI item: one shared instruction, then one interaction per
child, scored apart from its siblings. The ordering child keeps the all-or-nothing
shape the Exam.net probe TASK-SKRIPT-39-06-09 proved, and its 1 point is the only
point the workspace gives it once it sits in the part. PDF and DOCX print that point.
"""

import io
import re
from uuid import uuid4
from xml.etree import ElementTree

from docx import Document

from skriptoteket.domain.curated_apps.exam_conversion.digiexam_examnet_pdf_contracts import (
    DigiExamExamNetPdfStatus,
)
from skriptoteket.domain.curated_apps.exam_conversion.examnet_qti_item_xml import (
    serialize_qti_assessment_item,
)
from skriptoteket.domain.curated_apps.exam_conversion.examnet_qti_terms import qti_tag
from skriptoteket.domain.curated_apps.exam_workspace.examnet_pdf_export import (
    native_exam_to_pdf_document,
)
from skriptoteket.domain.curated_apps.exam_workspace.examnet_qti_export import (
    native_exam_to_examnet_qti_items,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_document import (
    NativeExamDocument,
    NativeExamDocumentOrigin,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_item import NativeExamItem
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_parent import (
    NativeExamParent,
    parent_points,
)
from skriptoteket.domain.curated_apps.exam_workspace.native_exam_parent_commands import move_item
from skriptoteket.infrastructure.curated_apps.apps.conversion_hub.exam_workspace_docx_writer import (  # noqa: E501
    ExamWorkspaceDocxWriter,
)
from tests.unit.domain.curated_apps.exam_workspace.native_exam_test_builders import (
    make_gap_fill_item,
    make_ordering_item,
    make_paragraph,
)

PART_INSTRUCTION = "Del 1: Besvara båda deluppgifterna."
SAVED_PRINT_ORDER = ("choice_003", "choice_001", "choice_004", "choice_002")
KEY_ORDER = ["choice_001", "choice_002", "choice_003", "choice_004"]


def _one_point_gap_child() -> NativeExamItem:
    return NativeExamItem.model_validate(
        {**make_gap_fill_item(sequence=2, item_id="item_002").model_dump(), "points": 1}
    )


def _standalone_ordering_before_a_part() -> NativeExamDocument:
    return NativeExamDocument(
        document_id=uuid4(),
        revision=1,
        title="Delprov med ordning",
        items=(
            make_ordering_item(
                sequence=1, item_id="item_001", points=3, print_order=SAVED_PRINT_ORDER
            ),
            _one_point_gap_child(),
        ),
        parents=(
            NativeExamParent(
                parent_id="parent_001",
                place=1,
                instruction=(make_paragraph(PART_INSTRUCTION),),
                item_ids=("item_002",),
            ),
        ),
        origin=NativeExamDocumentOrigin(kind="created"),
    )


def _ordering_moved_into_the_part() -> tuple[NativeExamDocument, NativeExamDocument]:
    """The part before and after the move: the ordering question is 3 points, then 1."""
    before = _standalone_ordering_before_a_part()
    moved = move_item(before, item_id="item_001", target_parent_id="parent_001")
    return before, moved


def test_the_move_brings_the_ordering_child_into_the_part_worth_one_point() -> None:
    before, moved = _ordering_moved_into_the_part()
    part = moved.parents[0]

    assert before.item_by_id("item_001").points == 3
    assert part.item_ids == ("item_002", "item_001")
    ordering = moved.item_by_id("item_001")
    assert ordering.points == 1
    assert ordering.points_estimated is False
    assert parent_points(moved, part) == 2


def test_qti_writes_the_ordering_child_with_the_probe_structure_and_the_part_total() -> None:
    _, moved = _ordering_moved_into_the_part()
    part = moved.parents[0]

    (qti,) = native_exam_to_examnet_qti_items(moved)

    assert qti.item_id == "parent_001"
    assert qti.max_score == parent_points(moved, part) == 2
    assert [child.interaction_type.value for child in qti.children] == [
        "gap_fill",
        "ordering",
    ]
    root = ElementTree.fromstring(serialize_qti_assessment_item(qti, image_paths=()))

    (interaction,) = root.iter(qti_tag("qti-order-interaction"))
    assert interaction.attrib["response-identifier"] == "item_001_RESPONSE"

    declarations = {
        element.attrib["identifier"]: element
        for element in root.iter(qti_tag("qti-response-declaration"))
    }
    declaration = declarations["item_001_RESPONSE"]
    assert declaration.attrib["cardinality"] == "ordered"
    assert declaration.attrib["base-type"] == "identifier"
    assert list(declaration.iter(qti_tag("qti-mapping"))) == []
    values = declaration.findall(f"{qti_tag('qti-correct-response')}/{qti_tag('qti-value')}")
    assert [value.text for value in values] == KEY_ORDER

    (match,) = root.iter(qti_tag("qti-match"))
    variables = match.findall(qti_tag("qti-variable"))
    assert [variable.attrib["identifier"] for variable in variables] == ["item_001_RESPONSE"]
    correct = match.find(qti_tag("qti-correct"))
    assert correct is not None and correct.attrib["identifier"] == "item_001_RESPONSE"
    added = [
        float(value.text or "")
        for value in root.iterfind(
            f".//{qti_tag('qti-response-condition')}/{qti_tag('qti-response-if')}"
            f"/{qti_tag('qti-set-outcome-value')}/{qti_tag('qti-sum')}"
            f"/{qti_tag('qti-base-value')}"
        )
    ]
    assert added == [1.0]

    outcomes = {
        element.attrib["identifier"]: element
        for element in root.iter(qti_tag("qti-outcome-declaration"))
    }
    max_score = outcomes["MAXSCORE"].find(f"{qti_tag('qti-default-value')}/{qti_tag('qti-value')}")
    assert max_score is not None and float(max_score.text or "") == 2.0


def test_pdf_prints_the_ordering_child_at_one_point_and_the_part_adds_up() -> None:
    _, moved = _ordering_moved_into_the_part()
    part = moved.parents[0]

    pdf = native_exam_to_pdf_document(moved)

    assert pdf.status is DigiExamExamNetPdfStatus.SUCCESS, pdf.warnings
    printed = [
        float(value) for value in re.findall(r'<p class="points">[^<]*: ([\d.,]+)</p>', pdf.html)
    ]
    assert printed == [1.0, 1.0]
    assert sum(printed) == parent_points(moved, part) == 2
    assert len(re.findall(r'<p class="ordering-step">', pdf.html)) == 4
    assert pdf.html.count(PART_INSTRUCTION) == 1


def test_docx_prints_the_ordering_child_at_one_point_and_the_part_adds_up() -> None:
    _, moved = _ordering_moved_into_the_part()
    part = moved.parents[0]

    writer = ExamWorkspaceDocxWriter()
    content = writer.build_docx_bytes(moved)
    writer.validate_docx_bytes(content, expected_item_count=len(moved.items))

    texts = [paragraph.text for paragraph in Document(io.BytesIO(content)).paragraphs]
    assert "Fråga 1a" in texts
    assert "Fråga 1b" in texts
    ordering_text = next(text for text in texts if "rätt ordning" in text)
    assert ordering_text.endswith("(1 p)")
    at_one_point = [text for text in texts if text.endswith("(1 p)")]
    assert len(at_one_point) == 2
    printed_total = 2
    assert printed_total == parent_points(moved, part) == 2
